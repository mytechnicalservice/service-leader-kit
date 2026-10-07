# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl==3.1.5", "python-docx==1.1.2", "python-pptx==1.0.2"]
# ///
"""Betriebsleitung (agent `betrieb`, spec §4/§5): numbers and document outlines for kapazitaet-lage,
eskalation-topkunde and teamleiter-runde. Every number leaves as a kennzahlen.wert (D8); Claude's document skills
write the files from `gliederung` (D7). Team level only, never per person (§9.3). Nothing is sent or decided."""
from __future__ import annotations

import datetime as dt
import re
import sys
from pathlib import Path

import daten_pruefen
import kennzahlen
import vorgang
from kennzahlen import deutsch, wert
from slk_common import JsonParser, run, zahl

STANDARD = "Standarddefinition des Kits – in der Einrichtung noch nicht festgelegt"
BEISPIEL = "Beispieldaten – Muster Maschinenbau GmbH"
AUSLASTUNG_ZIEL = 85.0  # % – kit standard while kpi-ziele.md names no "Auslastung" (D9)
BAND_GRUEN, BAND_GELB = 10.0, 20.0  # percentage points distance to the target
TREND_PP = 3.0
MIN_TEAM = 3  # §9.3: a team with fewer technicians never gets a line of its own
POOL = "Weitere Teams (unter 3 Technikern)"
OHNE_TEAM = "ohne Team"
ARBEITSTAGE = 21
ALT_TAGE = 14
ERLEDIGT = {"abgeschlossen", "erledigt", "storniert", "fakturiert"}
PERSONENSCHUTZ = (f"Teams mit weniger als {MIN_TEAM} Technikern sind zusammengefasst; ist auch die Zusammenfassung "
                  "kleiner, zählen sie nur in der Gesamtzeile (keine Auswertung einzelner Personen, §9.3).")
BEFEHLE: dict = {}  # subcommand -> (function(ws, args, defs) -> dict, {option: argparse kwargs})


class BetriebFehler(Exception):
    pass


def leer(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def z(v) -> float:
    return 0.0 if leer(v) else zahl(v)


def tag(v) -> dt.date | None:
    return None if leer(v) else dt.date.fromisoformat(str(v)[:10])


def de_datum(v) -> str:
    d = tag(v)
    return d.strftime("%d.%m.%Y") if d else "–"


def monat_minus(monat: str, n: int) -> str:
    j, m = map(int, monat.split("-"))
    k = j * 12 + m - 1 - n
    return f"{k // 12:04d}-{k % 12 + 1:02d}"


def monatsende(monat: str) -> dt.date:
    return dt.date.fromisoformat(monat_minus(monat, -1) + "-01") - dt.timedelta(days=1)


def zahltext(w: dict) -> str:
    return deutsch(w["betrag"], 1 if w["einheit"] in ("%", "Tage") else 0)


def fmt(w: dict | None) -> str:
    if w is None:
        return "–"
    return f"{zahltext(w)} {'€' if w['einheit'] == 'EUR' else w['einheit']}"


def zeilen_quelle(zeilen: list[dict]) -> list[str]:
    """'07_Daten/x.csv Zeilen 2, 5–7' per file: the rows a count came from."""
    je_datei: dict[str, list[int]] = {}
    for r in zeilen:
        je_datei.setdefault(r["_datei"], []).append(int(r["_zeile"]))
    out = []
    for datei, nums in sorted(je_datei.items()):
        nums, teile, start = sorted(set(nums)), [], None
        for i, n in enumerate(nums):
            start = n if start is None else start
            if i + 1 == len(nums) or nums[i + 1] != n + 1:
                teile.append(str(start) if start == n else f"{start}–{n}")
                start = None
        out.append(f"{datei} Zeile{'n' if len(nums) > 1 else ''} {', '.join(teile)}")
    return out


def anzahl(name: str, zeilen: list[dict], formel: str, einheit: str, leer_quelle: str) -> dict:
    return wert(name, float(len(zeilen)), zeilen_quelle(zeilen) or [leer_quelle], formel, einheit)


def summe(zeilen: list[dict], spalte: str, name: str, einheit: str,
          leer_quelle: str = "07_Daten/: keine passenden Zeilen") -> dict:
    if not zeilen:
        return wert(name, 0.0, [leer_quelle], None, einheit)
    w = kennzahlen.summe(zeilen, spalte, name)
    w["einheit"] = einheit
    return w


def alle_werte(obj) -> list[dict]:
    if isinstance(obj, dict):
        if {"name", "betrag", "einheit", "quelle"} <= obj.keys():
            return [obj]
        return [w for v in obj.values() for w in alle_werte(v)]
    if isinstance(obj, list):
        return [w for v in obj for w in alle_werte(v)]
    return []


def lade(ws: Path, vorlage: str, hinweise: list[str]) -> list[dict]:
    """All validated rows of one template; none at all yields [] and a German hint."""
    try:
        zeilen = kennzahlen.lade(ws, vorlage)
    except kennzahlen.KennzahlFehler as exc:
        hinweise.append(str(exc))
        return []
    if not zeilen:
        hinweise.append(f"Keine Daten der Vorlage '{vorlage}' in 07_Daten/ – bitte den Export mit 'daten-pruefen' "
                        "übernehmen.")
    return zeilen


def ziel_frei(ws: Path, rel: str) -> str:
    """Never overwrite: a taken name gets ' (2)', ' (3)' …"""
    p, n = ws / rel, 2
    while p.exists():
        p = (ws / rel).with_name(f"{Path(rel).stem} ({n}){Path(rel).suffix}")
        n += 1
    return p.relative_to(ws).as_posix()


def briefkopf(ws: Path) -> str | None:
    for basis in (ws, kennzahlen.datenquelle(ws)["ordner"]):
        p = basis / "Unternehmen" / "vorlagen" / "briefkopf.docx"
        if p.is_file():
            return p.relative_to(ws).as_posix()
    return None


def abschnitt(titel: str, absaetze: list[str] | None = None, tabelle: tuple | None = None,
              eingabe: str | None = None) -> dict:
    a: dict = {"titel": titel, "absaetze": absaetze or []}
    if tabelle:
        a["tabelle"] = {"spalten": tabelle[0], "zeilen": tabelle[1]}
    if eingabe:
        a["eingabe"] = eingabe
    return a


def standards(ziel_standard: bool) -> list[str]:
    s = [f"Ampel: grün bis {deutsch(BAND_GRUEN)}, gelb bis {deutsch(BAND_GELB)} Prozentpunkte Abstand zum "
         "Auslastungsziel",
         f"Rückstand: alt ab {ALT_TAGE} Tagen, Reichweite mit {ARBEITSTAGE} Arbeitstagen je Monat"]
    return ([f"Ziel-Auslastung {deutsch(AUSLASTUNG_ZIEL)} %"] if ziel_standard else []) + s


def ergebnis(ws: Path, titel: str, ziel: str, abschnitte: list[dict], daten: dict, pflicht: list[str],
             hinweise: list[str], standard: list[str]) -> dict:
    """The one output shape of every subcommand: data, outline for the document skill, mandatory figures."""
    beispiel = kennzahlen.datenquelle(ws)["beispiel"]
    werte = list({w["name"]: w for w in alle_werte(daten)}.values())
    hinweise = ([BEISPIEL] if beispiel else []) + hinweise + [f"{s}: {STANDARD}" for s in standard]
    abschnitte = [*abschnitte]
    if hinweise:
        abschnitte.append(abschnitt("Hinweise", hinweise))
    abschnitte.append(abschnitt("Quellen und Berechnungen", tabelle=(
        ["Kennzahl", "Wert", "Quelle", "Formel"],
        [[w["name"], fmt(w), "; ".join(w["quelle"]), w["formel"] or "–"] for w in werte])))
    return {"ok": True, "beispiel": beispiel, "kennzeichnung": BEISPIEL if beispiel else None,
            "ziel": ziel_frei(ws, ziel), "briefkopf": briefkopf(ws),
            "gliederung": {"titel": titel, "format": "docx", "abschnitte": abschnitte},
            "pflichtangaben": list(dict.fromkeys(pflicht)), "daten": daten, "hinweise": hinweise}


def dokument_text(p: Path) -> str:
    endung = p.suffix.lower()
    if endung == ".docx":
        import docx

        d = docx.Document(str(p))
        teile = [x.text for x in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]
        for s in d.sections:
            teile += [x.text for x in s.header.paragraphs]
        return "\n".join(teile)
    if endung == ".pptx":
        from pptx import Presentation

        teile = []
        for folie in Presentation(str(p)).slides:
            for sh in folie.shapes:
                if sh.has_text_frame:
                    teile.append(sh.text_frame.text)
                if getattr(sh, "has_table", False):
                    teile += [c.text for r in sh.table.rows for c in r.cells]
        return "\n".join(teile)
    if endung in (".md", ".txt"):
        return p.read_text(encoding="utf-8", errors="ignore")
    raise BetriebFehler(f"{p.name}: Dateiformat {endung or '(ohne Endung)'} kann nicht geprüft werden (docx, pptx, md).")


def pruefe_dokument(ws: Path, datei: str, pflicht: list[str]) -> dict:
    """Every mandatory figure must appear verbatim, not as part of a longer number."""
    p = Path(datei) if Path(datei).is_absolute() else ws / datei
    if not p.is_file():
        raise BetriebFehler(f"Dokument nicht gefunden: {datei}")
    text = dokument_text(p).replace(" ", " ").replace(" ", " ")
    fehlt = [s for s in pflicht if not re.search(rf"(?<![\d.,]){re.escape(s)}(?!\d|[.,]\d)", text)]
    return {"datei": datei, "vollstaendig": not fehlt, "fehlt": fehlt}


# --- Kapazität und Rückstand (team level only) ---

def kpi_ziel(defs: dict) -> tuple[float, bool]:
    """Target utilisation in % from kpi-ziele.md (D9) and whether the kit standard is used instead."""
    kpi = defs.get("kpi-ziele") or defs.get("kpi_ziele") or {}
    for k in kpi.get("kennzahlen") or []:
        if str(k.get("name", "")).strip().casefold().startswith("auslastung") and not leer(k.get("ziel")):
            ziel = zahl(str(k["ziel"]).replace("%", "").strip())
            return (ziel * 100 if ziel <= 1.5 else ziel), False
    return AUSLASTUNG_ZIEL, True


def ampel(auslastung: float, ziel: float) -> str:
    abstand = abs(auslastung - ziel)
    return "grün" if abstand <= BAND_GRUEN else "gelb" if abstand <= BAND_GELB else "rot"


def team_linien(zeilen: list[dict]) -> dict[str, str | None]:
    """Team -> line it is reported on; small teams are pooled, a pool below MIN_TEAM only counts in the total."""
    koepfe: dict[str, float] = {}
    for r in zeilen:
        t = str(r["Team"]).strip()
        koepfe[t] = koepfe.get(t, 0.0) + z(r["Techniker_Anzahl"])
    klein = {t for t, n in koepfe.items() if n < MIN_TEAM}
    pool = sum(koepfe[t] for t in klein)
    return {t: t if t not in klein else (POOL if pool >= MIN_TEAM else None) for t in koepfe}


def kapazitaet_monat(kap: list[dict], monat: str, ziel: float) -> tuple[dict, dict]:
    zeilen = [r for r in kap if str(r["Monat"])[:7] == monat]
    linien = team_linien(zeilen)
    gruppen: dict[str, list[dict]] = {}
    for r in sorted(zeilen, key=lambda r: linien[str(r["Team"]).strip()] == POOL):
        linie = linien[str(r["Team"]).strip()]
        if linie:
            gruppen.setdefault(linie, []).append(r)

    def block(name: str, rows: list[dict]) -> dict:
        soll = summe(rows, "Soll_Stunden", f"Soll-Stunden {name} {monat}", "h")
        ist = summe(rows, "Ist_Stunden", f"Ist-Stunden {name} {monat}", "h")
        aus = None
        if soll["betrag"] > 0:
            aus = wert(f"Auslastung {name} {monat}", round(ist["betrag"] / soll["betrag"] * 100, 1),
                       list(dict.fromkeys(soll["quelle"] + ist["quelle"])), "Ist_Stunden / Soll_Stunden × 100", "%")
        return {"team": name, "techniker": summe(rows, "Techniker_Anzahl", f"Techniker {name} {monat}", "Techniker"),
                "soll": soll, "ist": ist, "auslastung": aus,
                "ampel": ampel(aus["betrag"], ziel) if aus else "keine Soll-Stunden"}

    return {"monat": monat, "teams": [block(t, rows) for t, rows in gruppen.items()],
            "gesamt": block("Gesamt", zeilen)}, linien


def trend(aktuell: dict, frueher: dict | None) -> None:
    vorher = {t["team"]: t["auslastung"] for t in frueher["teams"]} if frueher else {}
    for t in aktuell["teams"]:
        a, b = t["auslastung"], vorher.get(t["team"])
        t["auslastung_vorher"] = b
        if not (a and b):
            t["trend"] = "keine Vergleichsdaten"
            continue
        d = a["betrag"] - b["betrag"]
        t["trend"] = "steigend" if d >= TREND_PP else "fallend" if d <= -TREND_PP else "stabil"


def ist_offen(r: dict, stichtag: dt.date) -> bool:
    eingang = tag(r.get("Eingang"))
    return (eingang is not None and eingang <= stichtag and leer(r.get("Abschluss"))
            and str(r.get("Status") or "").strip().casefold() not in ERLEDIGT)


def rueckstand(auftraege: list[dict], linien: dict, soll: dict[str, dict], stichtag: dt.date) -> dict:
    offen = [r for r in auftraege if ist_offen(r, stichtag)]
    gruppen: dict[str, list[dict]] = {}
    for r in offen:
        t = OHNE_TEAM if leer(r.get("Team")) else str(r["Team"]).strip()
        linie = linien.get(t, t)
        if linie:
            gruppen.setdefault(linie, []).append(r)
    formel = (f"Anzahl Aufträge mit leerem Abschluss, Status nicht {'/'.join(sorted(ERLEDIGT))}, "
              f"Eingang bis {stichtag.strftime('%d.%m.%Y')}")
    leer_q = "07_Daten/auftraege_*.csv: keine offenen Aufträge"

    def block(name: str, rows: list[dict]) -> dict:
        alt = [r for r in rows if (stichtag - tag(r["Eingang"])).days > ALT_TAGE]
        std = summe(rows, "Stunden", f"Rückstand Stunden {name}", "h", leer_q)
        s, reich = soll.get(name), None
        if s and s["betrag"] > 0:
            reich = wert(f"Reichweite Rückstand {name}", round(std["betrag"] / (s["betrag"] / ARBEITSTAGE), 1),
                         list(dict.fromkeys(std["quelle"] + s["quelle"])),
                         f"Rückstand Stunden / (Soll-Stunden des Monats / {ARBEITSTAGE} Arbeitstage)", "Tage")
        return {"team": name, "auftraege": anzahl(f"Offene Aufträge {name}", rows, formel, "Aufträge", leer_q),
                "aelter": anzahl(f"Offene Aufträge {name} älter als {ALT_TAGE} Tage", alt,
                                 f"{formel}, Eingang mehr als {ALT_TAGE} Tage vor dem Stichtag", "Aufträge", leer_q),
                "stunden": std, "reichweite": reich}

    return {"stichtag": stichtag.isoformat(), "teams": [block(t, rows) for t, rows in gruppen.items()],
            "gesamt": block("Gesamt", offen)}


def vergleich(ws: Path, datei: str, monat: str, ist: dict, heute: str) -> dict:
    """Second source for the same month: dry run of daten-pruefen, never imported, never averaged."""
    p = (Path(datei) if Path(datei).is_absolute() else ws / datei).resolve()
    if not p.is_relative_to(ws.resolve()):
        raise BetriebFehler(f"{datei} liegt nicht im Kundendienst-Ordner.")
    rel = p.relative_to(ws.resolve()).as_posix()
    r = daten_pruefen.pruefe(ws, p, "kapazitaet", None, [], False, heute)
    if r["summe"] is None or r["periode"] != monat:
        grund = "; ".join(r["meldungen"]) or f"Datei betrifft {r['periode']}, nicht {monat}"
        return {"datei": rel, "vergleichbar": False, "grund": grund}
    andere = wert(f"Ist-Stunden gesamt {monat} laut {p.name}", r["summe"], [f"{rel} (nicht übernommen)"], None, "h")
    diff = wert(f"Differenz Ist-Stunden {monat}", round(r["summe"] - ist["betrag"], 2), andere["quelle"] + ist["quelle"],
                f"{p.name} − 07_Daten", "h")
    return {"datei": rel, "vergleichbar": True, "widerspruch": abs(diff["betrag"]) > 0.005, "andere_quelle": andere,
            "differenz": diff}


def cmd_kapazitaet_lage(ws: Path, a, defs: dict) -> dict:
    hinweise: list[str] = []
    if a.monat and not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", a.monat):
        raise BetriebFehler(f"--monat '{a.monat}' ist kein Monat (JJJJ-MM).")
    kap = lade(ws, "kapazitaet", [])
    if not kap:
        raise BetriebFehler("In 07_Daten/ liegen keine Kapazitätsdaten. Bitte zuerst den Kapazitäts-Export (je Team) "
                            "mit dem Skill 'daten-pruefen' übernehmen.")
    monate = sorted({str(r["Monat"])[:7] for r in kap})
    monat = a.monat or monate[-1]
    if monat not in monate:
        raise BetriebFehler(f"Für {daten_pruefen.monatsname(monat)} gibt es keine Kapazitätsdaten "
                            f"(vorhanden: {', '.join(monate)}).")
    ziel, ziel_standard = kpi_ziel(defs)
    lage, linien = kapazitaet_monat(kap, monat, ziel)
    vorher = monat_minus(monat, 2)
    trend(lage, kapazitaet_monat(kap, vorher, ziel)[0] if vorher in monate else None)
    g = lage["gesamt"]
    soll = {t["team"]: t["soll"] for t in [*lage["teams"], g]}
    rs = rueckstand(lade(ws, "auftraege", hinweise), linien, soll, monatsende(monat))
    if any(linie != t for t, linie in linien.items()):
        hinweise.append(PERSONENSCHUTZ)
    daten = {"monat": monat, "kapazitaet": lage, "rueckstand": rs,
             "ziel_auslastung": wert("Ziel-Auslastung", ziel, ["Standarddefinition des Kits"] if ziel_standard
                                     else ["Unternehmen/kpi-ziele.md"], None, "%")}
    rote = [t["team"] for t in lage["teams"] if t["ampel"] == "rot"]
    rg = rs["gesamt"]
    kurz = [f"Auslastung gesamt {fmt(g['auslastung'])} ({g['ampel']}): {fmt(g['ist'])} Ist von {fmt(g['soll'])} Soll.",
            f"Rückstand am {de_datum(rs['stichtag'])}: {fmt(rg['auftraege'])}, {fmt(rg['stunden'])}"
            + (f", Reichweite {fmt(rg['reichweite'])}" if rg["reichweite"] else "") + ".",
            ("Rot: " + ", ".join(rote) + ".") if rote else "Kein Team ist rot."]
    abschnitte = [
        abschnitt("Auf einen Blick", kurz),
        abschnitt(f"Auslastung je Team ({daten_pruefen.monatsname(monat)})", tabelle=(
            ["Team", "Techniker", "Soll-Stunden", "Ist-Stunden", "Auslastung", "Ampel", "Trend (2 Monate)"],
            [[t["team"], zahltext(t["techniker"]), fmt(t["soll"]), fmt(t["ist"]), fmt(t["auslastung"]), t["ampel"],
              t.get("trend", "–")] for t in [*lage["teams"], g]])),
        abschnitt(f"Rückstand je Team (Stichtag {de_datum(rs['stichtag'])})", tabelle=(
            ["Team", "Offene Aufträge", f"davon älter als {ALT_TAGE} Tage", "Stunden", "Reichweite"],
            [[t["team"], zahltext(t["auftraege"]), zahltext(t["aelter"]), fmt(t["stunden"]), fmt(t["reichweite"])]
             for t in [*rs["teams"], rg]]))]
    if a.vergleich:
        v = daten["vergleich"] = vergleich(ws, a.vergleich, monat, g["ist"], a.heute)
        if v["vergleichbar"] and v["widerspruch"]:
            abschnitte.append(abschnitt("Widerspruch zwischen Quellen", [
                f"07_Daten: {fmt(g['ist'])} Ist-Stunden; {v['datei']}: {fmt(v['andere_quelle'])}; Differenz "
                f"{fmt(v['differenz'])}.",
                "Die Werte werden nicht gemittelt. Bitte klären, welche Quelle gilt; bis dahin gilt 07_Daten."]))
        elif not v["vergleichbar"]:
            hinweise.append(f"{v['datei']} konnte nicht verglichen werden: {v['grund']}")
    abschnitte.append(abschnitt("Maßnahmen", eingabe=(
        "Für jedes rote Team: Maßnahme mit verantwortlicher Person und Termin (wird ein Vorgang) oder 'keine Maßnahme' "
        "mit Begründung – nur, was der Nutzer sagt; sonst 'offen – mit dem Nutzer klären'.")))
    pflicht = [zahltext(g["soll"]), zahltext(g["ist"])]
    pflicht += [zahltext(t["auslastung"]) for t in [*lage["teams"], g] if t["auslastung"]]
    pflicht.append(zahltext(rg["stunden"]))
    return ergebnis(ws, f"Kapazitätslage {daten_pruefen.monatsname(monat)}", f"03_Berichte/{a.heute}_kapazitaet-lage.docx",
                    abschnitte, daten, pflicht, hinweise, standards(ziel_standard))


BEFEHLE["kapazitaet-lage"] = (cmd_kapazitaet_lage, {"monat": {}, "vergleich": {}})


# --- CLI ---

def parser() -> JsonParser:
    ap = JsonParser(prog="betrieb")
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name, (_, optionen) in BEFEHLE.items():
        sp = sub.add_parser(name)
        sp.add_argument("--ws", required=True)
        sp.add_argument("--heute", default=dt.date.today().isoformat())
        sp.add_argument("--pruefe")
        for o, kw in optionen.items():
            sp.add_argument(f"--{o}", **kw)
    return ap


def _main(argv: list[str] | None) -> tuple[int, dict]:
    a = parser().parse_args(argv)
    ws = Path(a.ws)
    try:
        if not (ws / "Unternehmen" / ".kit-config").is_file():
            raise BetriebFehler(f"{ws} ist kein Kundendienst-Ordner (Unternehmen/.kit-config fehlt).")
        try:
            dt.date.fromisoformat(a.heute)
        except ValueError:
            raise BetriebFehler(f"--heute '{a.heute}' ist kein Datum (JJJJ-MM-TT).") from None
        out = BEFEHLE[a.cmd][0](ws, a, kennzahlen.definitionen(ws))
        if a.pruefe:
            out["dokument"] = pruefe_dokument(ws, a.pruefe, out["pflichtangaben"])
            if not out["dokument"]["vollstaendig"]:
                out["ok"] = False
                out["meldungen"] = ["Im Dokument fehlen Zahlen aus dem Skript: " + ", ".join(out["dokument"]["fehlt"])]
                return 1, out
        return 0, out
    except (BetriebFehler, kennzahlen.KennzahlFehler, vorgang.VorgangFehler) as exc:
        return 1, {"ok": False, "fehler": [str(exc)], "meldungen": [str(exc)]}


def main(argv: list[str] | None = None) -> int:
    return run(_main, argv)


if __name__ == "__main__":
    sys.exit(main())
