# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl==3.1.5"]
# ///
"""Serviceangebot (agent `angebot`, spec §5): Serviceprodukt-Konzept, Preislisten-Update, Portfolio-Review.

Every number leaves as a kennzahlen.wert. Documents are written by Claude's document skills (D7); the only file this
script writes is the new price list, a data file next year's update reads again (Plan 4f, default 17)."""
from __future__ import annotations

import datetime as dt
import os
import re
import sys
import tempfile
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from kennzahlen import KennzahlFehler, datenquelle, definitionen, deutsch, lade, summe, wert
from slk_common import JsonParser, csv_sicher, run, zahl

STANDARD_HINWEIS = "Standarddefinition des Kits – in der Einrichtung noch nicht festgelegt"
BEISPIEL_HINWEIS = "Beispieldaten – Muster Maschinenbau GmbH"
SPALTEN = ["Artikelnr", "Bezeichnung", "Kategorie", "Einheit", "Preis_EUR", "Gueltig_ab", "Bemerkung"]
PFLICHT = SPALTEN[:6]
GEWICHTE = {"Stundensatz": ("0.80", "0", "0.20"), "Pauschale": ("0.40", "0", "0.60"),
            "Vertrag": ("0.70", "0.15", "0.15"), "Schulung": ("0.70", "0", "0.30"),
            "Ersatzteil": ("0", "1", "0"), "Sonstiges": ("0", "0", "1")}  # default 12: Lohn, Material, allgemein
ELASTIZITAET = Decimal("8")
GRENZEN = (Decimal("-20"), Decimal("30"))
ZIELMARGE_STANDARD = 35.0  # used when kpi-ziele.md sets no DB II target (K1)
ZIEL_KPI = "DB II-Marge"
DB2 = re.compile(r"\b(db|deckungsbeitrag)\s*(ii|2)\b")
WACHSTUM_AUSBAU, ANTEIL_KLEIN, ANTEIL_GROSS = 5.0, 5.0, 10.0
ANWEISUNG = re.compile(r"@|ignorier|\bKI\b|schick|sende|anweisung", re.I)
PREIS_GLIEDERUNG = ["Ergebnis auf einen Blick (Änderung je Kategorie)", "Annahmen (Lohn, Material, allgemeine Kosten, "
                    "Markt) mit Quelle", "Formel und Rundungsregel", "Preiselastizität und Risiken",
                    "Nicht angepasste Positionen", "Entscheidung: wer gibt frei (Vorgang)", "Quellen"]


class AngebotFehler(Exception):
    pass


def monate(bis: str, anzahl: int) -> list[str]:
    try:
        j, m = (int(x) for x in bis.split("-"))
        dt.date(j, m, 1)
    except ValueError as exc:
        raise AngebotFehler(f"Monat '{bis}' ist ungültig (erwartet JJJJ-MM)") from exc
    out = []
    for _ in range(anzahl):
        out.append(f"{j:04d}-{m:02d}")
        j, m = (j, m - 1) if m > 1 else (j - 1, 12)
    return out[::-1]


def freier_name(ws: Path, ordner: str, stamm: str, endung: str) -> str:
    """Never overwrite: stamm.endung, else stamm_v2.endung, _v3 … (compared case-insensitively)."""
    d = ws / ordner
    vorhanden = {p.name.casefold() for p in d.iterdir()} if d.is_dir() else set()
    name, n = f"{stamm}{endung}", 2
    while name.casefold() in vorhanden:
        name, n = f"{stamm}_v{n}{endung}", n + 1
    return f"{ordner}/{name}"


def quelle_von(zeilen: list[dict]) -> list[str]:
    bereiche: dict[str, list[int]] = {}
    for z in zeilen:
        bereiche.setdefault(str(z["_datei"]), []).append(int(z["_zeile"]))
    return [f"{d} Zeilen {min(n)}–{max(n)}" for d, n in sorted(bereiche.items())]


def w(name, betrag, quelle, formel=None, einheit="EUR", nachkomma=0) -> dict:
    return wert(name, float(betrag), list(quelle), formel, einheit) | {"anzeige": deutsch(float(betrag), nachkomma)}


def abschnitt(name: str, defs: dict) -> dict:
    return defs.get(name) or defs.get(name.replace("-", "_")) or {"standard": True}


def ist_db2(name) -> bool:
    """K1 (Max, 2026-10-07): the target margin is the DB II KPI – 'DB II-Marge', 'DB2', 'Deckungsbeitrag II' …;
    a plain 'Marge' or a DB I KPI does not count."""
    return bool(DB2.search(re.sub(r"[^0-9a-zäöüß]+", " ", str(name).casefold())))


def zielmarge(ws: Path) -> tuple[float, str]:
    """Target margin = the `ziel` of the first kpi-ziele KPI that is the DB II margin, else ZIELMARGE_STANDARD."""
    kpi = abschnitt("kpi-ziele", definitionen(ws))
    for k in kpi.get("kennzahlen") or []:
        if ist_db2(k.get("name", "")) and k.get("ziel") is not None:
            return zahl(k["ziel"]), STANDARD_HINWEIS if kpi.get("standard") else f"Unternehmen/kpi-ziele.md: {k['name']}"
    return ZIELMARGE_STANDARD, STANDARD_HINWEIS


def entscheidungsrecht(ws: Path, thema: str) -> str:
    for r in abschnitt("ergebnisrechnung", definitionen(ws)).get("entscheidungsrechte") or []:
        if str(r.get("thema")) == thema:
            bis = r.get("allein_bis_eur")
            allein = f"allein bis {deutsch(zahl(bis))} EUR, sonst " if bis is not None else ""
            return f"{allein}{r.get('sonst') or 'nicht festgelegt'} (Unternehmen/ergebnisrechnung.md)"
    return f"nicht festgelegt – {STANDARD_HINWEIS}"


# ---------- preisliste ----------

def stufe(preis: Decimal) -> Decimal:
    return Decimal("0.50") if preis < 100 else Decimal("1") if preis < 1000 else Decimal("10")


def runden(preis: Decimal) -> Decimal:
    s = stufe(preis)
    return (preis / s).quantize(Decimal("1"), ROUND_HALF_UP) * s


def prozent(text, was: str) -> Decimal:
    try:
        p = Decimal(str(zahl(str(text).replace("%", "").strip())))
    except ValueError as exc:
        raise AngebotFehler(f"{was}: '{text}' ist keine Prozentzahl") from exc
    if not GRENZEN[0] <= p <= GRENZEN[1]:
        raise AngebotFehler(f"{was}: {text} % liegt außerhalb von −20 … +30 % – bitte prüfen")
    return p


def kategorie_aenderung(kat: str, lohn: Decimal, material: Decimal, allgemein: Decimal, markt: dict) -> Decimal:
    gl, gm, ga = (Decimal(x) for x in GEWICHTE[kat])
    roh = gl * lohn + gm * material + ga * allgemein + markt.get(kat, Decimal("0"))
    return roh.quantize(Decimal("0.01"), ROUND_HALF_UP)


def lies_preisliste(pfad: Path) -> list[dict]:
    if not pfad.is_file():
        raise AngebotFehler(f"Preisliste {pfad.name} nicht gefunden – bitte die Datei nennen")
    try:
        wb = load_workbook(pfad, read_only=True, data_only=True)
    except PermissionError as exc:
        raise AngebotFehler(f"{pfad.name} ist noch in Excel geöffnet – bitte schließen und noch einmal versuchen") from exc
    except Exception as exc:  # BadZipFile, InvalidFileException …: one German message for all
        raise AngebotFehler(f"{pfad.name} lässt sich nicht als Excel-Datei lesen ({type(exc).__name__})") from exc
    blatt = wb["Preisliste"] if "Preisliste" in wb.sheetnames else wb.worksheets[0]
    rows = list(blatt.iter_rows(values_only=True))
    wb.close()
    kopf = [str(c).strip() if c is not None else "" for c in (rows[0] if rows else [])]
    fehlt = [s for s in PFLICHT if s not in kopf]
    if fehlt:
        raise AngebotFehler(f"{pfad.name}: Spalte(n) {', '.join(fehlt)} fehlen – erwartet: {', '.join(SPALTEN)}")
    idx = {s: kopf.index(s) for s in SPALTEN if s in kopf}
    return [{s: (r[i] if i < len(r) else None) for s, i in idx.items()} | {"_zeile": nr}
            for nr, r in enumerate(rows[1:], start=2) if any(v not in (None, "") for v in r)]


def neue_preise(zeilen: list[dict], aend: dict, datei: str) -> tuple[list[dict], list[str]]:
    vorkommen: dict[str, list[int]] = {}
    for z in zeilen:
        vorkommen.setdefault(str(z.get("Artikelnr") or "").strip(), []).append(z["_zeile"])
    ergebnis, meldungen = [], []
    for z in zeilen:
        art, kat = str(z.get("Artikelnr") or "").strip(), str(z.get("Kategorie") or "").strip()
        r = {"zeile": z["_zeile"], "artikelnr": art, "bezeichnung": z.get("Bezeichnung"), "kategorie": kat,
             "alt": None, "neu": None, "aenderung_prozent": None, "status": "unverändert", "grund": None,
             "quelle": f"{datei} Zeile {z['_zeile']}", "formel": None}
        if ANWEISUNG.search(str(z.get("Bemerkung") or "")):
            meldungen.append(f"Zeile {z['_zeile']} ({art}): Bemerkung enthält eine Anweisung oder Mailadresse – "
                             "nur als Daten behandelt, bitte prüfen")
        try:
            r["alt"] = zahl(z.get("Preis_EUR"))
        except ValueError:
            r["grund"] = "Preis fehlt oder ist keine Zahl"
        if not art:
            r["grund"] = "Artikelnr fehlt"
        elif len(vorkommen[art]) > 1:
            r["grund"] = "Artikelnr mehrfach (Zeilen " + ", ".join(map(str, vorkommen[art])) + ")"
        elif kat not in GEWICHTE:
            r["grund"] = f"Kategorie '{kat}' unbekannt (erlaubt: {', '.join(GEWICHTE)})"
        if r["grund"]:
            meldungen.append(f"Zeile {z['_zeile']} ({art or 'ohne Nr.'}): {r['grund']} – nicht angepasst, bitte klären")
        else:
            p = aend[kat]
            roh = Decimal(str(r["alt"])) * (1 + p / 100)
            r |= {"neu": float(runden(roh)), "aenderung_prozent": float(p), "status": "angepasst",
                  "formel": f"{deutsch(r['alt'], 2)} × (1 + {deutsch(float(p), 2)} %) = {deutsch(float(roh), 2)}, "
                            f"gerundet auf {deutsch(float(stufe(roh)), 2)} EUR"}
        ergebnis.append(r)
    return ergebnis, meldungen


def speichere_neu(ziel: Path, wb: Workbook) -> None:
    """Saves through a temp file; refuses if the target appeared meanwhile (never overwrite)."""
    ziel.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=ziel.parent, prefix=f".{ziel.name}.", suffix=".tmp")
    os.close(fd)
    try:
        wb.save(tmp)
        if ziel.exists():
            raise AngebotFehler(f"{ziel.name} ist inzwischen vorhanden – nichts überschrieben, bitte noch einmal starten")
        os.replace(tmp, ziel)
    finally:
        Path(tmp).unlink(missing_ok=True)


def schreibe_liste(ziel: Path, zeilen, ergebnis, jahr: int, grundlage: str, annahmen: dict) -> None:
    wb = Workbook()
    pl = wb.active
    pl.title = "Preisliste"
    pl.append(SPALTEN)
    for z, r in zip(zeilen, ergebnis):
        preis = r["neu"] if r["status"] == "angepasst" else z.get("Preis_EUR")
        pl.append([csv_sicher(z.get("Artikelnr")), csv_sicher(z.get("Bezeichnung")), csv_sicher(z.get("Kategorie")),
                   csv_sicher(z.get("Einheit")), csv_sicher(preis), dt.date(jahr, 1, 1), csv_sicher(z.get("Bemerkung"))])
    ae = wb.create_sheet("Änderungen")
    ae.append(["Artikelnr", "Bezeichnung", "Kategorie", "Preis_alt_EUR", "Preis_neu_EUR", "Aenderung_Prozent",
               "Status", "Grund", "Quelle"])
    for r in ergebnis:
        ae.append([csv_sicher(r["artikelnr"]), csv_sicher(r["bezeichnung"]), csv_sicher(r["kategorie"]), r["alt"],
                   r["neu"], r["aenderung_prozent"], r["status"], r["grund"], r["quelle"]])
    hi = wb.create_sheet("Hinweis")
    hi.append(["Entwurf – gilt erst nach Freigabe (Vorgang in 01_Vorgaenge)."])
    hi.append(["Grundlage", grundlage])
    for k, v in annahmen.items():
        hi.append([k, v])
    speichere_neu(ziel, wb)


def preisliste(ws, q, jahr, lohn, material, allgemein, markt, datei, schreiben, heute) -> dict:
    l, m, a = prozent(lohn, "Lohn"), prozent(material, "Material"), prozent(allgemein, "Allgemeine Kosten")
    mk = {}
    for eintrag in markt:
        kat, _, wert_ = eintrag.partition("=")
        if kat.strip() not in GEWICHTE:
            raise AngebotFehler(f"Marktanpassung '{eintrag}': Kategorie unbekannt (erlaubt: {', '.join(GEWICHTE)})")
        mk[kat.strip()] = prozent(wert_, f"Markt {kat.strip()}")
    pfad = Path(datei) if datei else q["ordner"] / "04_Angebote" / f"preisliste_{jahr - 1}.xlsx"
    pfad = pfad if pfad.is_absolute() else ws / pfad
    grundlage = pfad.relative_to(ws).as_posix() if pfad.is_relative_to(ws) else str(pfad)
    zeilen = lies_preisliste(pfad)
    aend = {k: kategorie_aenderung(k, l, m, a, mk) for k in GEWICHTE}
    ergebnis, meldungen = neue_preise(zeilen, aend, grundlage)
    kategorien, warnungen = [], []
    for k in sorted({r["kategorie"] for r in ergebnis if r["kategorie"] in GEWICHTE}):
        gl, gm, ga = (g.replace(".", ",") for g in GEWICHTE[k])
        formel = (f"{gl} × Lohn {deutsch(float(l), 2)} % + {gm} × Material {deutsch(float(m), 2)} % + {ga} × "
                  f"allgemeine Kosten {deutsch(float(a), 2)} % + Markt {deutsch(float(mk.get(k, 0)), 2)} Prozentpunkte")
        kategorien.append(w(f"Preisänderung {k}", aend[k], ["Angaben des Nutzers (Lohn, Material, allgemeine Kosten, "
                            "Markt)", "Kostenanteile: Standard des Angebots-Skripts (Plan 4f)"], formel, "%", 2))
        if aend[k] > ELASTIZITAET:
            warnungen.append(f"{k}: +{deutsch(float(aend[k]), 2)} % – über 8 %: Preiselastizität beachten "
                             "(Kündigungen, Abwanderung zum Wettbewerb); Stufung über zwei Jahre oder Ankündigung prüfen")
        elif aend[k] < 0:
            warnungen.append(f"{k}: Preissenkung {deutsch(float(aend[k]), 2)} % – bewusst so gewollt?")
    annahmen = {"Lohn %": float(l), "Material %": float(m), "Allgemeine Kosten %": float(a),
                **{f"Markt {k} (Prozentpunkte)": float(v) for k, v in mk.items()}}
    ziel = freier_name(ws, "04_Angebote", f"preisliste_{jahr}", ".xlsx")
    if schreiben:
        schreibe_liste(ws / ziel, zeilen, ergebnis, jahr, grundlage, annahmen)
    return {"jahr": jahr, "grundlage": grundlage, "annahmen": annahmen, "kategorien": kategorien,
            "zeilen": ergebnis, "meldungen": meldungen, "warnungen": warnungen, "ziel": ziel, "geschrieben": schreiben,
            "memo_ziel": freier_name(ws, "04_Angebote", f"{heute}_preisliste-update", ".docx"),
            "budget_annahmen": [{"kategorie": k["name"].removeprefix("Preisänderung "), "aenderung_prozent": k["betrag"]}
                                for k in kategorien],
            "entscheidung_preise": entscheidungsrecht(ws, "preise"), "gliederung": PREIS_GLIEDERUNG}


# ---------- portfolio ----------

PORTFOLIO_GLIEDERUNG = ["Ergebnis auf einen Blick (Klassen je Produkt)", "Umsatz, Marge, Wachstum je Produkt",
                        "Ziel: DB II-Marge (Herkunft) und verwendete Definitionen", "Vertragsdeckung der Installed Base",
                        "Empfehlung je Produkt: halten / sanieren / ausbauen / auslaufen prüfen",
                        "Vermarktung: Zielgruppe und Botschaft für Produkte 'ausbauen'",
                        "Vorgeschlagene Vorgänge (nur nach Bestätigung)", "Datenlücken und Hinweise", "Quellen"]


def im_zeitraum(zeilen: list[dict], spalte: str, perioden: set[str]) -> list[dict]:
    return [z for z in zeilen if str(z.get(spalte) or "")[:7] in perioden]


def hat_werte(zeilen: list[dict], spalte: str) -> bool:
    return bool(zeilen) and all(z.get(spalte) is not None and str(z.get(spalte)).strip() for z in zeilen)


def keine_zeilen(name: str) -> list[str]:
    return [f"keine Datenzeilen für '{name}' im Zeitraum"]


def summe_von(zeilen: list[dict], spalte: str, name: str) -> dict:
    """kennzahlen.summe plus `anzeige`; an empty period is 0 with a "no rows" source instead of an error
    (kennzahlen.wert refuses an empty source list)."""
    if not zeilen:
        return w(name, 0, keine_zeilen(name), None)
    s = summe(zeilen, spalte, name)
    return s | {"anzeige": deutsch(s["betrag"], 0)}


def prozent_wert(name: str, zaehler: float, nenner: float, quelle: list[str], formel: str) -> dict | None:
    return w(name, round(zaehler / nenner * 100, 1), quelle, formel, "%", 1) if nenner else None


def teile_umsatz(zeilen: list[dict], name: str) -> dict:
    betrag = round(sum(zahl(z["Menge"]) * zahl(z["Stueckpreis_EUR"]) for z in zeilen), 2)
    return w(name, betrag, quelle_von(zeilen) or keine_zeilen(name), "Σ Menge × Stueckpreis_EUR")


def klasse(marge, wachstum, anteil, ziel) -> tuple[str, str]:
    if marge is None:
        return "unklar", "Kosten fehlen – Marge nicht berechenbar"
    if marge < 0:
        if anteil >= ANTEIL_GROSS:
            return "sanieren", "negative Marge bei großem Umsatzanteil"
        return "auslaufen prüfen", "negative Marge"
    if marge < ziel:
        if anteil < ANTEIL_KLEIN and (wachstum is None or wachstum <= 0):
            return "auslaufen prüfen", "Marge unter Ziel (DB II-Marge), kleiner Anteil, kein Wachstum"
        return "sanieren", "Marge unter Ziel (DB II-Marge)"
    if wachstum is not None and wachstum >= WACHSTUM_AUSBAU:
        return "ausbauen", "Marge über Ziel (DB II-Marge) und Wachstum ab 5 %"
    return "halten", "Marge über Ziel (DB II-Marge)"


def produkt(name, herkunft, u, kosten, u1, u2) -> dict:
    db = w(f"DB {name}", round(u["betrag"] - kosten["betrag"], 2), u["quelle"] + kosten["quelle"],
           "Umsatz − Kosten") if kosten else None
    return {"produkt": name, "herkunft": herkunft, "umsatz": u, "kosten": kosten, "db": db,
            "marge": prozent_wert(f"Marge {name}", db["betrag"], u["betrag"], u["quelle"], "DB ÷ Umsatz") if db else None,
            "umsatz_h1": u1, "umsatz_h2": u2,
            "wachstum": prozent_wert(f"Wachstum {name}", u2["betrag"] - u1["betrag"], u1["betrag"],
                                     u1["quelle"] + u2["quelle"], "(Umsatz 2. Halbjahr − 1. Halbjahr) ÷ 1. Halbjahr")}


def vertraege(basis: Path, stichtag: dt.date) -> tuple[list[dict], list[str]]:
    gefunden, hinweise = [], []
    for p in sorted((basis / "06_Kunden").glob("*/vertrag.md")):
        m = re.match(r"\A---\r?\n(.*?)\r?\n---", p.read_text(encoding="utf-8-sig"), re.S)
        kopf = {}
        for zeile in m.group(1).splitlines() if m else []:
            k, _, v = zeile.partition(":")
            kopf[k.strip()] = v.strip().strip('"')
        rel = p.relative_to(basis).as_posix()
        try:
            gebuehr, beginn = zahl(kopf["jahresgebuehr_eur"]), dt.date.fromisoformat(kopf["beginn"])
            ende = None if kopf.get("ende") in (None, "", "null") else dt.date.fromisoformat(kopf["ende"])
        except (KeyError, ValueError):
            hinweise.append(f"{rel}: Kopf unvollständig (jahresgebuehr_eur, beginn, ende) – nicht gezählt")
            continue
        if beginn <= stichtag and (ende is None or ende >= stichtag):
            gefunden.append({"datei": rel, "gebuehr": gebuehr})
    return gefunden, hinweise


def portfolio(ws: Path, basis: Path, bis: str, heute: str) -> dict:
    alle = monate(bis, 12)
    h1, h2 = set(alle[:6]), set(alle[6:])
    auf, teile = lade(basis, "auftraege", alle), lade(basis, "ersatzteile", alle)
    ziel, ziel_quelle = zielmarge(ws)
    produkte, hinweise = [], []
    for art in sorted({str(z.get("Auftragsart") or "").strip() for z in auf} - {""}):
        zz = [z for z in auf if str(z.get("Auftragsart") or "").strip() == art]
        kosten = summe_von(zz, "Kosten_EUR", f"Kosten {art}") if hat_werte(zz, "Kosten_EUR") else None
        produkte.append(produkt(art, "Aufträge", summe_von(zz, "Umsatz_EUR", f"Umsatz {art}"), kosten,
                                summe_von(im_zeitraum(zz, "Eingang", h1), "Umsatz_EUR", f"Umsatz {art} 1. Halbjahr"),
                                summe_von(im_zeitraum(zz, "Eingang", h2), "Umsatz_EUR", f"Umsatz {art} 2. Halbjahr")))
    if teile:
        produkte.append(produkt("Ersatzteile", "Ersatzteile", teile_umsatz(teile, "Umsatz Ersatzteile"), None,
                                teile_umsatz(im_zeitraum(teile, "Datum", h1), "Umsatz Ersatzteile 1. Halbjahr"),
                                teile_umsatz(im_zeitraum(teile, "Datum", h2), "Umsatz Ersatzteile 2. Halbjahr")))
    gesamt = round(sum(p["umsatz"]["betrag"] for p in produkte), 2)
    quellen = [q for p in produkte for q in p["umsatz"]["quelle"]]
    for p in produkte:
        p["anteil"] = prozent_wert(f"Anteil {p['produkt']}", p["umsatz"]["betrag"], gesamt, quellen,
                                   "Umsatz Produkt ÷ Umsatz Portfolio (Aufträge + Ersatzteile)")
        m, g = (p[k]["betrag"] if p[k] else None for k in ("marge", "wachstum"))
        p["klasse"], p["begruendung"] = klasse(m, g, p["anteil"]["betrag"] if p["anteil"] else 0.0, ziel)
    j, mo = (int(x) for x in bis.split("-"))
    stichtag = dt.date(j + mo // 12, mo % 12 + 1, 1) - dt.timedelta(days=1)
    vt, vt_hinweise = vertraege(basis, stichtag)
    hinweise += vt_hinweise
    if vt:
        u = w("Jahreswert laufende Wartungsverträge", round(sum(v["gebuehr"] for v in vt), 2),
              [v["datei"] for v in vt], f"Σ jahresgebuehr_eur der am {stichtag.isoformat()} laufenden Verträge")
        produkte.append({"produkt": "Wartungsverträge", "herkunft": "Verträge", "umsatz": u, "kosten": None,
                         "db": None, "marge": None, "umsatz_h1": None, "umsatz_h2": None, "wachstum": None,
                         "anteil": None, "klasse": "unklar",
                         "begruendung": "Jahreswert, keine Periodenumsätze und keine getrennten Kosten – nicht klassifiziert"})
    else:
        hinweise.append("Keine laufenden Wartungsverträge gefunden (06_Kunden/<Kunde>/vertrag.md) – "
                        "Wartungsverträge fehlen im Portfolio")
    try:
        ib = lade(basis, "installed_base", [bis])
        mit = [z for z in ib if str(z.get("Vertrag")).strip() == "ja"]
        quote = prozent_wert("Vertragsdeckung Installed Base", len(mit), len(ib), quelle_von(ib),
                             "Anlagen mit Vertrag ÷ alle Anlagen")
    except KennzahlFehler as exc:
        quote = None
        hinweise.append(f"Vertragsdeckung nicht berechenbar: {exc}")
    return {"zeitraum": f"{alle[0]} bis {alle[-1]}", "produkte": produkte,
            "umsatz_auftraege": summe_von(auf, "Umsatz_EUR", "Umsatz Aufträge gesamt"),
            "umsatz_portfolio": w("Umsatz Portfolio (Aufträge + Ersatzteile)", gesamt, quellen, "Σ Umsatz je Produkt"),
            "zielmarge": {"prozent": ziel, "quelle": ziel_quelle, "kennzahl": ZIEL_KPI,
                          "anzeige": f"Ziel: {ZIEL_KPI} {deutsch(ziel, 1)} %"},
            "vertragsquote": quote,
            "margendefinition": "Marge je Produkt = (Umsatz − Kosten_EUR) ÷ Umsatz auf Auftragsebene, nicht DB I/II "
                                "aus ergebnisrechnung.md; verglichen mit dem Ziel der DB II-Marge (kpi-ziele.md)",
            "hinweise": hinweise, "ziel": freier_name(ws, "03_Berichte", f"{heute}_portfolio-review", ".docx"),
            "gliederung": PORTFOLIO_GLIEDERUNG}


# ---------- CLI ----------

def parser() -> JsonParser:
    ap = JsonParser(prog="angebot")
    sub = ap.add_subparsers(dest="cmd", required=True)

    def add(name: str):
        sp = sub.add_parser(name)
        sp.add_argument("--ws", required=True)
        sp.add_argument("--heute", default=dt.date.today().isoformat())
        return sp

    p = add("preisliste")
    p.add_argument("--jahr", type=int, required=True)
    for n in ("lohn", "material", "allgemein"):
        p.add_argument(f"--{n}", required=True)
    p.add_argument("--markt", action="append", default=[])
    p.add_argument("--datei")
    p.add_argument("--schreiben", action="store_true")
    r = add("portfolio")
    r.add_argument("--bis", required=True)
    return ap


def fehler(text: str) -> dict:
    return {"ok": False, "fehler": [text], "meldungen": [text]}


def ausfuehren(a, ws: Path, q: dict) -> dict:
    if a.cmd == "preisliste":
        return preisliste(ws, q, a.jahr, a.lohn, a.material, a.allgemein, a.markt, a.datei, a.schreiben, a.heute)
    if a.cmd == "portfolio":
        return portfolio(ws, q["ordner"], a.bis, a.heute)
    raise AngebotFehler(f"Unbekannter Befehl {a.cmd}")


def _main(argv: list[str] | None) -> tuple[int, dict]:
    a = parser().parse_args(argv)
    ws = Path(a.ws).expanduser().resolve()
    if not (ws / "Unternehmen").is_dir():
        return 1, fehler(f"{ws} ist kein Kundendienst-Ordner (Unternehmen/ fehlt)")
    try:
        dt.date.fromisoformat(a.heute)
    except ValueError:
        return 1, fehler(f"--heute '{a.heute}' ist kein Datum (JJJJ-MM-TT)")
    try:
        q = datenquelle(ws)
        out = ausfuehren(a, ws, q)
    except (AngebotFehler, KennzahlFehler) as exc:
        return 1, fehler(str(exc))
    if q["beispiel"]:
        out["hinweis_beispiel"] = BEISPIEL_HINWEIS
    return 0, {"ok": True, **out}


def main(argv: list[str] | None = None) -> int:
    return run(_main, argv)


if __name__ == "__main__":
    sys.exit(main())
