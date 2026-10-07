# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl==3.1.5"]
# ///
"""Vertrieb: Großangebot, Key-Account-Review, Verlängerungs-Radar, Installed-Base-Potenziale (spec §5, §8).
Every number leaves this script as a kennzahlen.wert with source and formula; the skills only present them."""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

from openpyxl import load_workbook

from daten_pruefen import monatsname
from kennzahlen import KennzahlFehler, datenquelle, definitionen, deutsch, lade, summe, wert
from slk_common import JsonParser, run, zahl
from vorgang import VorgangFehler, all_cases

BEISPIEL_HINWEIS = "Beispieldaten – Muster Maschinenbau GmbH"
STANDARD = "Standardannahme des Kits – in der Einrichtung noch nicht festgelegt"

# Kit standards (Plan 4d, domain defaults 1–22). Every output names the ones it used.
HORIZONT_MONATE = 6
DRINGEND_MONATE = 3
VORLAUF_TAGE = 60
TOP_N = 5
RETROFIT_ALTER = {"MM-400": 12, "MM-600": 12, "MM-800": 15}
RETROFIT_ALTER_SONST = 12
LAUFZEIT_JAHRE = 3
BESUCHE_JE_JAHR = 2
STUNDEN_JE_BESUCH = 8
RABATT_STAFFEL = ((5, 8.0), (2, 5.0), (1, 0.0))  # (ab Anlagen, Prozent); the first match wins
GUELTIG_TAGE = 30
BEZEICHNUNG = {"auftraege": "Aufträge", "ersatzteile": "Ersatzteilverkäufe"}
FRONT = re.compile(r"\A---\n(.*?)\n---\n", re.S)
LEER = {"", "null", "~"}


class VertriebFehler(Exception):
    pass


class StufeFehlt(VertriebFehler):
    """Several contract levels match and no --stufe was given (Max's correction K1): the skill asks the user."""

    def __init__(self, text: str, stufen: list[dict]):
        super().__init__(text)
        self.stufen = stufen


# --- small helpers ------------------------------------------------------------------------------------------------

def text(v) -> str:
    return "" if v is None else str(v).strip()


def gleich(a, b) -> bool:
    return text(a).casefold() == text(b).casefold()


def betrag(v) -> float:
    return 0.0 if text(v).casefold() in LEER else zahl(v)


def datum(v) -> dt.date | None:
    return None if text(v).casefold() in LEER else dt.date.fromisoformat(text(v)[:10])


def pruefe_datum(s: str, feld: str) -> dt.date:
    try:
        return dt.date.fromisoformat(s)
    except ValueError as exc:
        raise VertriebFehler(f"{feld} '{s}' ist kein Datum (JJJJ-MM-TT)") from exc


def rel(ws: Path, p: Path) -> str:
    try:
        return p.relative_to(ws).as_posix()
    except ValueError:
        return p.resolve().relative_to(ws.resolve()).as_posix()


def monatsende(monat: str) -> dt.date:
    j, m = map(int, monat.split("-"))
    return dt.date(j + m // 12, m % 12 + 1, 1) - dt.timedelta(days=1)


def plus_monate(d: dt.date, n: int) -> dt.date:
    m0 = d.month - 1 + n
    j, m = d.year + m0 // 12, m0 % 12 + 1
    return dt.date(j, m, min(d.day, monatsende(f"{j:04d}-{m:02d}").day))


def monate_bis(ende: str, anzahl: int) -> list[str]:
    j, m = map(int, ende.split("-"))
    out = []
    for _ in range(anzahl):
        out.append(f"{j:04d}-{m:02d}")
        j, m = (j, m - 1) if m > 1 else (j - 1, 12)
    return out[::-1]


def ordnername(kunde: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', "", kunde).strip()


def datenordner(ws: Path) -> Path:
    return Path(datenquelle(ws)["ordner"])


def neuester_monat(ws: Path, vorlage: str) -> str:
    monate = []
    for p in (datenordner(ws) / "07_Daten").glob(f"{vorlage}_*.csv"):
        m = re.fullmatch(rf"{vorlage}_(\d{{4}}-\d{{2}})", p.stem)
        if m:
            monate.append(m.group(1))
    if not monate:
        raise VertriebFehler(f"Keine Daten '{vorlage}' in 07_Daten/ – bitte den Export zuerst mit daten-pruefen "
                             "übernehmen.")
    return max(monate)


def ort(z: dict) -> str:
    return f"{z['_datei']} Zeile {z['_zeile']}"


def quelle(zeilen: list[dict]) -> list[str]:
    nach: dict[str, list[int]] = {}
    for z in zeilen:
        nach.setdefault(z["_datei"], []).append(int(z["_zeile"]))
    return [f"{d} Zeile {n[0]}" if len(n) == 1 else f"{d} Zeilen {min(n)}–{max(n)}" for d, n in sorted(nach.items())]


def summe_oder_null(zeilen: list[dict], spalte: str, name: str, leer: str) -> dict:
    return summe(zeilen, spalte, name) if zeilen else wert(name, 0.0, [leer])


def addiere(name: str, teile: list[dict], formel: str) -> dict:
    """Sum of wert objects; with no parts it is 0 and says so (kennzahlen.wert refuses an empty source list)."""
    quellen = sorted({q for t in teile for q in t["quelle"]}) or ["keine bewerteten Einzelwerte – Summe 0"]
    return wert(name, round(sum(t["betrag"] for t in teile), 2), quellen, formel=formel)


def lade_zeitraum(ws: Path, vorlage: str, monate: list[str]) -> tuple[list[dict], list[str]]:
    """Rows of each month; a missing month is named, never filled in (domain default 13)."""
    zeilen, meldungen = [], []
    for m in monate:
        try:
            zeilen += lade(ws, vorlage, [m])
        except KennzahlFehler:
            meldungen.append(f"{BEZEICHNUNG[vorlage]} {monatsname(m)} fehlen in 07_Daten/ – Summen ohne diesen "
                             "Monat, nicht hochgerechnet.")
    return zeilen, meldungen


def zeitraum(ws: Path) -> list[str]:
    return monate_bis(neuester_monat(ws, "auftraege"), 12)


def ergebnis(ws: Path, felder: dict, zeilen: list[str]) -> dict:
    beispiel = bool(datenquelle(ws)["beispiel"])
    return {"ok": True, "beispiel": beispiel, **felder,
            "zusammenfassung": ([BEISPIEL_HINWEIS] if beispiel else []) + zeilen}


def heute_von(heute: str | None) -> dt.date:
    return pruefe_datum(heute, "--heute") if heute else dt.date.today()


# --- company files ------------------------------------------------------------------------------------------------

def vertragsdaten(ws: Path, kunde: str) -> dict:
    """Flat front matter of 06_Kunden/<Kunde>/vertrag.md (scalars only), plus `_datei`; {} if there is none."""
    p = datenordner(ws) / "06_Kunden" / ordnername(kunde) / "vertrag.md"
    if not p.is_file():
        return {}
    m = FRONT.match(p.read_text(encoding="utf-8-sig").replace("\r\n", "\n"))
    out = {"_datei": rel(ws, p)}
    for zeile in m.group(1).splitlines() if m else []:
        k, sep, v = zeile.partition(":")
        if sep and k.strip() and not zeile[0].isspace() and not zeile.startswith(("-", "#")):
            out[k.strip()] = v.split(" #")[0].strip().strip("\"'")
    return out


# --- installed base -----------------------------------------------------------------------------------------------

VERGLEICH = ("Maschinentyp", "Baujahr", "Vertrag", "Vertragsende")


def installed_base(ws: Path) -> tuple[list[dict], dict]:
    """Machines of the newest export; identical duplicates count once, contradicting duplicates are left out."""
    monat = neuester_monat(ws, "installed_base")
    erste: dict[tuple[str, str], dict] = {}
    widerspruch, meldungen = set(), []
    for z in lade(ws, "installed_base", [monat]):
        key = (text(z.get("Kunde")), text(z.get("Anlage")))
        if key not in erste:
            erste[key] = z
            continue
        alt = erste[key]
        stelle = f"{z['_datei']} (Zeilen {alt['_zeile']} und {z['_zeile']})"
        if all(text(alt.get(s)) == text(z.get(s)) for s in VERGLEICH):
            meldungen.append(f"{key[0]} / {key[1]} steht doppelt in {stelle} – einmal gezählt.")
        else:
            widerspruch.add(key)
            meldungen.append(f"Widerspruch: {key[0]} / {key[1]} steht mit unterschiedlichen Angaben in {stelle} – "
                             "nicht bewertet, bitte im Export klären.")
    anlagen = [z for k, z in erste.items() if k not in widerspruch]
    return anlagen, {"monat": monat, "stichtag": monatsende(monat), "meldungen": meldungen}


# --- verlaengerungs-radar -----------------------------------------------------------------------------------------

def vertragswert(ws: Path, kunde: str, ende: dt.date, meldungen: list[str]) -> tuple[dict | None, bool]:
    v = vertragsdaten(ws, kunde)
    stelle = v.get("_datei", f"06_Kunden/{ordnername(kunde)}/vertrag.md")
    widerspruch = False
    try:
        ende_md = datum(v.get("vertragsende"))
    except ValueError:
        ende_md = None
        meldungen.append(f"{kunde}: Vertragsende in {stelle} ist kein Datum – nicht abgeglichen.")
    if ende_md and ende_md != ende:
        widerspruch = True
        meldungen.append(f"Widerspruch Vertragsende {kunde}: Installed Base {ende:%d.%m.%Y}, {stelle} "
                         f"{ende_md:%d.%m.%Y} – bitte klären; die Liste nutzt die Installed Base.")
    try:
        jw = betrag(v.get("jahreswert_eur"))
    except ValueError:
        jw = 0.0
    if jw:
        return wert(f"Vertragsjahreswert {kunde}", jw, [stelle]), widerspruch
    meldungen.append(f"{kunde}: Vertragsjahreswert fehlt ({stelle}, Feld jahreswert_eur) – Wert im Risiko ohne "
                     "Vertragswert.")
    return None, widerspruch


def verlaengerungs_radar(ws: Path, monate: int | None = None, stichtag: str | None = None,
                         heute: str | None = None) -> dict:
    annahmen = []
    if monate is None:
        monate = HORIZONT_MONATE
        annahmen.append(f"Horizont {monate} Monate, dringend bis {DRINGEND_MONATE} Monate ({STANDARD})")
    if not 1 <= monate <= 36:
        raise VertriebFehler("--monate muss zwischen 1 und 36 liegen")
    tag_heute = heute_von(heute)
    anlagen, ib = installed_base(ws)
    tag = pruefe_datum(stichtag, "--stichtag") if stichtag else ib["stichtag"]
    grenze, dringend_bis = plus_monate(tag, monate), plus_monate(tag, DRINGEND_MONATE)
    auf, luecken = lade_zeitraum(ws, "auftraege", zeitraum(ws))
    meldungen = ib["meldungen"] + luecken
    gruppen: dict[tuple[str, dt.date], list[dict]] = {}
    ohne_ende = []
    for a in anlagen:
        if not gleich(a.get("Vertrag"), "ja"):
            continue
        try:
            ende = datum(a.get("Vertragsende"))
        except ValueError:
            ende = None
        if ende is None:
            ohne_ende.append(f"{text(a['Kunde'])} / {text(a['Anlage'])}: Vertrag ja, aber kein Vertragsende "
                             f"({ort(a)})")
        elif ende <= grenze:
            gruppen.setdefault((text(a["Kunde"]), ende), []).append(a)
    vertraege = []
    for (kunde, ende), liste in sorted(gruppen.items(), key=lambda x: (x[0][1], x[0][0])):
        namen = sorted({text(a["Anlage"]) for a in liste})
        zeilen = [z for z in auf if text(z.get("Kunde")) == kunde and text(z.get("Anlage")) in namen]
        umsatz = summe_oder_null(zeilen, "Umsatz_EUR", f"Serviceumsatz 12 Monate {kunde} ({', '.join(namen)})",
                                 "Aufträge der letzten 12 Monate: keine Zeilen auf diesen Anlagen")
        jw, widerspruch = vertragswert(ws, kunde, ende, meldungen)
        teile = [jw, umsatz] if jw else [umsatz]
        formel = ("Vertragsjahreswert + Serviceumsatz der Vertragsanlagen (12 Monate)" if jw
                  else "Serviceumsatz der Vertragsanlagen (12 Monate); Vertragsjahreswert fehlt")
        vertraege.append({
            "kunde": kunde, "vertragsende": ende.isoformat(), "anlagen": namen, "teile": teile,
            "status": "abgelaufen" if ende < tag else "dringend" if ende <= dringend_bis else "planen",
            "wert_im_risiko": addiere(f"Wert im Risiko {kunde}", teile, formel), "widerspruch": widerspruch,
            "faellig_vorschlag": max(ende - dt.timedelta(days=VORLAUF_TAGE),
                                     tag_heute + dt.timedelta(days=7)).isoformat()})
    gesamt = addiere("Wert im Risiko gesamt", [v["wert_im_risiko"] for v in vertraege], "Σ Wert im Risiko je Vertrag")
    zeilen_text = [f"Stichtag {tag:%d.%m.%Y}, Horizont {monate} Monate (bis {grenze:%d.%m.%Y}): {len(vertraege)} "
                   f"Verträge, Wert im Risiko {deutsch(gesamt['betrag'])} EUR."]
    zeilen_text += [f"{v['kunde']}: Vertragsende {datum(v['vertragsende']):%d.%m.%Y} ({v['status']}), "
                    f"{len(v['anlagen'])} Anlage(n), Wert im Risiko {deutsch(v['wert_im_risiko']['betrag'])} EUR"
                    for v in vertraege]
    return ergebnis(ws, {
        "stichtag": tag.isoformat(), "horizont_monate": monate, "vertraege": vertraege, "ohne_vertragsende": ohne_ende,
        "summe": gesamt, "annahmen": annahmen, "meldungen": meldungen,
        "ziel": f"03_Berichte/{tag_heute.isoformat()}_verlaengerungs-radar.xlsx"}, zeilen_text)


# --- price list ---------------------------------------------------------------------------------------------------

POSITION = ("position", "leistung", "artikel", "bezeichnung")
PREIS = ("preis_eur", "preis", "listenpreis", "netto_eur", "preis (eur)")


def preisliste(ws: Path) -> dict:
    """Rows (position, price) of the newest 04_Angebote/preisliste_*.xlsx; the header is searched in rows 1–10."""
    dateien = sorted((datenordner(ws) / "04_Angebote").glob("preisliste_*.xlsx"))
    if not dateien:
        return {"datei": None, "zeilen": []}
    p = dateien[-1]
    name, zeilen = rel(ws, p), []
    wb = load_workbook(p, read_only=True, data_only=True)
    try:
        for blatt in wb.worksheets:
            rows = list(blatt.iter_rows(values_only=True))
            for i, r in enumerate(rows[:10]):
                kopf = [text(c).casefold() for c in r]
                ip = next((kopf.index(s) for s in POSITION if s in kopf), None)
                ipr = next((kopf.index(s) for s in PREIS if s in kopf), None)
                if ip is None or ipr is None:
                    continue
                for n, z in enumerate(rows[i + 1:], start=i + 2):
                    if len(z) <= max(ip, ipr) or not text(z[ip]) or not text(z[ipr]):
                        continue
                    try:
                        zeilen.append({"position": text(z[ip]), "preis": zahl(z[ipr]),
                                       "quelle": f"{name} Blatt {blatt.title} Zeile {n}"})
                    except ValueError:
                        continue
                break
    finally:
        wb.close()
    return {"datei": name, "zeilen": zeilen}


def preis_zeile(liste: dict, *woerter: str) -> dict | None:
    """First row containing every word (case-insensitive); with several, the shortest name (base variant)."""
    w = [x.casefold() for x in woerter]
    treffer = [z for z in liste["zeilen"] if all(x in z["position"].casefold() for x in w)]
    return min(treffer, key=lambda z: (len(z["position"]), z["position"])) if treffer else None


# --- installed-base-potenziale ------------------------------------------------------------------------------------

def baujahr(a: dict) -> int | None:
    """Year of build as int; the validated CSV may hold '2012' or '2012.0'; empty or unreadable → None."""
    try:
        return None if text(a.get("Baujahr")) == "" else int(zahl(a.get("Baujahr")))
    except ValueError:
        return None


def jahrtext(a: dict) -> str:
    """Year of build for documents: '2012', not the float '2012.0' that kennzahlen.lade returns; '' if missing."""
    bj = baujahr(a)
    return "" if bj is None else str(bj)


def retrofit_grenze(typ: str) -> int:
    for prefix, jahre in RETROFIT_ALTER.items():
        if typ.casefold().startswith(prefix.casefold()):
            return jahre
    return RETROFIT_ALTER_SONST


def installed_base_potenziale(ws: Path, kunde: str | None = None, stichtag: str | None = None,
                              heute: str | None = None) -> dict:
    tag_heute = heute_von(heute)
    anlagen, ib = installed_base(ws)
    tag = pruefe_datum(stichtag, "--stichtag") if stichtag else ib["stichtag"]
    preise = preisliste(ws)
    meldungen = list(ib["meldungen"])
    if not preise["zeilen"]:
        meldungen.append("Keine Preisliste (04_Angebote/preisliste_*.xlsx) – Potenziale nur gezählt, nicht bewertet.")
    kunden: dict[str, dict] = {}
    for a in anlagen:
        k, typ, anlage = text(a.get("Kunde")), text(a.get("Maschinentyp")), text(a.get("Anlage"))
        if kunde and not gleich(k, kunde):
            continue
        e = kunden.setdefault(k, {"kunde": k, "ohne_vertrag": [], "retrofit": [], "nicht_bewertet": []})
        if not gleich(a.get("Vertrag"), "ja"):
            p = preis_zeile(preise, "wartungsvertrag", typ)
            e["ohne_vertrag"].append({"anlage": anlage, "typ": typ, "potenzial": wert(
                f"Wartungsvertrag {typ} je Jahr, Basisstufe ({k} / {anlage})", p["preis"], [ort(a), p["quelle"]])
                if p else None})
            if not p:
                e["nicht_bewertet"].append(f"{anlage} ({typ}): ohne Vertrag, kein Preis 'Wartungsvertrag {typ}' in "
                                           "der Preisliste")
        if "retrofit" in typ.casefold():
            continue
        bj = baujahr(a)
        if bj is None:
            e["nicht_bewertet"].append(f"{anlage} ({typ}): Baujahr fehlt – Retrofit nicht beurteilt ({ort(a)})")
            continue
        if bj > tag.year:
            meldungen.append(f"{k} / {anlage}: Baujahr {bj} liegt nach dem Stichtag – unplausibel, nicht beurteilt "
                             f"({ort(a)}).")
            continue
        alter, grenze = tag.year - bj, retrofit_grenze(typ)
        if alter < grenze:
            continue
        p = preis_zeile(preise, "retrofit", typ)
        e["retrofit"].append({"anlage": anlage, "typ": typ, "baujahr": bj, "alter": alter, "grenze": grenze,
                              "potenzial": wert(f"Retrofit {typ} ({k} / {anlage})", p["preis"], [ort(a), p["quelle"]])
                              if p else None})
        if not p:
            e["nicht_bewertet"].append(f"{anlage} ({typ}, {alter} Jahre): Retrofit-Kandidat, kein Preis "
                                       f"'Retrofit {typ}' in der Preisliste")
    for e in kunden.values():
        e["vertragspotenzial"] = addiere(f"Vertragspotenzial je Jahr {e['kunde']}",
                                         [x["potenzial"] for x in e["ohne_vertrag"] if x["potenzial"]],
                                         "Σ Listenpreis Wartungsvertrag je Anlage ohne Vertrag")
        e["retrofitpotenzial"] = addiere(f"Retrofitpotenzial einmalig {e['kunde']}",
                                         [x["potenzial"] for x in e["retrofit"] if x["potenzial"]],
                                         "Σ Listenpreis Retrofit je Retrofit-Kandidat")
    liste = sorted((e for e in kunden.values() if e["ohne_vertrag"] or e["retrofit"] or e["nicht_bewertet"]),
                   key=lambda e: (-(e["vertragspotenzial"]["betrag"] + e["retrofitpotenzial"]["betrag"]), e["kunde"]))
    s_v = addiere("Vertragspotenzial je Jahr gesamt", [e["vertragspotenzial"] for e in liste],
                  "Σ Vertragspotenzial je Kunde")
    s_r = addiere("Retrofitpotenzial einmalig gesamt", [e["retrofitpotenzial"] for e in liste],
                  "Σ Retrofitpotenzial je Kunde")
    n_o, n_r = sum(len(e["ohne_vertrag"]) for e in liste), sum(len(e["retrofit"]) for e in liste)
    annahmen = [f"Retrofit-Kandidat ab Alter: {', '.join(f'{t} {j} Jahre' for t, j in RETROFIT_ALTER.items())}, "
                f"sonst {RETROFIT_ALTER_SONST} Jahre; Typen mit 'Retrofit' im Namen gelten als modernisiert ({STANDARD})",
                f"Alter = Jahr des Stichtags {tag:%d.%m.%Y} − Baujahr",
                "Vertragspotenzial zum Listenpreis der Basisstufe (kürzeste Position 'Wartungsvertrag <Typ>' der "
                f"Preisliste) ({STANDARD})"]
    zeilen = [f"{n_o} Anlagen ohne Vertrag: Vertragspotenzial {deutsch(s_v['betrag'])} EUR je Jahr.",
              f"{n_r} Retrofit-Kandidaten: Retrofitpotenzial {deutsch(s_r['betrag'])} EUR einmalig "
              "(nicht mit dem Vertragspotenzial addiert)."]
    zeilen += [f"{e['kunde']}: Vertrag {deutsch(e['vertragspotenzial']['betrag'])} EUR/Jahr, Retrofit "
               f"{deutsch(e['retrofitpotenzial']['betrag'])} EUR einmalig" for e in liste]
    return ergebnis(ws, {
        "stichtag": tag.isoformat(), "kunden": liste, "summe_vertrag": s_v, "summe_retrofit": s_r,
        "anzahl_ohne_vertrag": n_o, "anzahl_retrofit": n_r, "annahmen": annahmen, "meldungen": meldungen,
        "ziel": f"03_Berichte/{tag_heute.isoformat()}_installed-base-potenziale.xlsx"}, zeilen)


# --- key-account-review -------------------------------------------------------------------------------------------

def monat_von(z: dict) -> str:
    return re.search(r"(\d{4}-\d{2})\.csv$", z["_datei"]).group(1)


def kundenumsatz(kunde: str, auf: list[dict], et: list[dict], monate: list[str]) -> dict:
    a = [z for z in auf if text(z.get("Kunde")) == kunde]
    e = [z for z in et if text(z.get("Kunde")) == kunde]
    service = summe_oder_null(a, "Umsatz_EUR", f"Serviceumsatz 12 Monate {kunde}", "Aufträge: keine Zeilen im Zeitraum")
    teile = wert(f"Ersatzteilumsatz 12 Monate {kunde}",
                 round(sum(betrag(z.get("Menge")) * betrag(z.get("Stueckpreis_EUR")) for z in e), 2),
                 quelle(e) or ["Ersatzteile: keine Zeilen im Zeitraum"], formel="Σ Menge × Stueckpreis_EUR")
    reihe = []
    for m in monate:
        zm = [z for z in a + e if monat_von(z) == m]
        reihe.append(wert(f"Umsatz {kunde} {monatsname(m)}",
                          round(sum(betrag(z.get("Umsatz_EUR")) if "Umsatz_EUR" in z
                                    else betrag(z.get("Menge")) * betrag(z.get("Stueckpreis_EUR")) for z in zm), 2),
                          quelle(zm) or [f"{monatsname(m)}: keine Zeilen"],
                          formel="Umsatz_EUR (Aufträge) + Menge × Stueckpreis_EUR (Ersatzteile)"))
    gesamt = addiere(f"Umsatz 12 Monate {kunde}", [service, teile], "Serviceumsatz (Aufträge) + Ersatzteilumsatz")
    return {"umsatz_12m": gesamt, "service": service, "ersatzteile": teile, "monate": reihe}


def offene_vorgaenge(ws: Path) -> tuple[list[dict], list[str]]:
    try:
        return [{"nr": m["nr"], "titel": m["titel"], "typ": m["typ"], "status": m["status"],
                 "faellig": m.get("faellig"), "kunde": text(m.get("kunde"))}
                for _, m, _ in all_cases(ws, ("offen",))], []
    except VorgangFehler as exc:
        return [], [f"Vorgänge nicht lesbar: {exc} – offene Vorgänge fehlen im Review."]


def key_account_review(ws: Path, kunde: str | None = None, top: int | None = None, heute: str | None = None) -> dict:
    tag_heute = heute_von(heute)
    monate = zeitraum(ws)
    auf, l1 = lade_zeitraum(ws, "auftraege", monate)
    et, l2 = lade_zeitraum(ws, "ersatzteile", monate)
    meldungen, annahmen = l1 + l2, [f"Umsatz = Serviceaufträge + Ersatzteile, {monatsname(monate[0])} bis "
                                     f"{monatsname(monate[-1])} ({STANDARD})"]
    try:
        anlagen, ib = installed_base(ws)
        radar = verlaengerungs_radar(ws, heute=tag_heute.isoformat())
        pot = installed_base_potenziale(ws, heute=tag_heute.isoformat())
        meldungen += [m for m in radar["meldungen"] + pot["meldungen"] if m not in meldungen]
    except VertriebFehler as exc:
        anlagen, ib, radar, pot = [], {"stichtag": None}, {"vertraege": []}, {"kunden": []}
        meldungen.append(f"{exc} Anlagen, Verträge, Verlängerungen und Potenziale fehlen im Review.")
    bekannt = sorted({text(z.get("Kunde")) for z in auf + et + anlagen} - {""})
    if kunde:
        treffer = [k for k in bekannt if gleich(k, kunde)]
        if not treffer:
            raise VertriebFehler(f"Kunde '{kunde.strip()}' kommt in Aufträgen, Ersatzteilen und Installed Base nicht "
                                 f"vor – Schreibweise prüfen. Bekannte Kunden: {', '.join(bekannt)}")
        auswahl = treffer[:1]
    else:
        if top is None:
            top = TOP_N
            annahmen.append(f"Top {top} Kunden nach Umsatz der letzten 12 Monate ({STANDARD})")
        if not 1 <= top <= 50:
            raise VertriebFehler("--top muss zwischen 1 und 50 liegen")
        auswahl = sorted(bekannt, key=lambda k: (-kundenumsatz(k, auf, et, monate)["umsatz_12m"]["betrag"], k))[:top]
    faelle, fehler = offene_vorgaenge(ws)
    meldungen += fehler
    konten = []
    for k in auswahl:
        u = kundenumsatz(k, auf, et, monate)
        eigene = [a for a in anlagen if text(a.get("Kunde")) == k]
        konten.append({
            "kunde": k, **u,
            "anlagen": [{"anlage": text(a.get("Anlage")), "typ": text(a.get("Maschinentyp")),
                         "baujahr": jahrtext(a),
                         "alter": ib["stichtag"].year - baujahr(a) if baujahr(a) is not None else None,
                         "vertrag": text(a.get("Vertrag")), "vertragsende": text(a.get("Vertragsende")),
                         "quelle": ort(a)} for a in sorted(eigene, key=lambda a: text(a.get("Anlage")))],
            "vertrag": {s: v for s, v in vertragsdaten(ws, k).items()},
            "verlaengerung": [v for v in radar["vertraege"] if v["kunde"] == k],
            "vorgaenge": [f for f in faelle if gleich(f["kunde"], k)],
            "potenzial": next((p for p in pot["kunden"] if p["kunde"] == k), None)})
    ziel = (f"06_Kunden/{ordnername(auswahl[0])}/{tag_heute.isoformat()}_key-account-review.docx" if kunde
            else f"03_Berichte/{tag_heute.isoformat()}_key-account-review.docx")
    zeilen = [f"{k['kunde']}: Umsatz 12 Monate {deutsch(k['umsatz_12m']['betrag'])} EUR (Service "
              f"{deutsch(k['service']['betrag'])}, Ersatzteile {deutsch(k['ersatzteile']['betrag'])}), "
              f"{len(k['anlagen'])} Anlagen, {len(k['vorgaenge'])} offene Vorgänge" for k in konten]
    return ergebnis(ws, {"zeitraum": [monate[0], monate[-1]], "konten": konten, "annahmen": annahmen,
                         "meldungen": meldungen, "ziel": ziel}, zeilen)


# --- grossangebot -------------------------------------------------------------------------------------------------

GLIEDERUNG = [
    "Anschreiben mit Bezug auf die Anfrage (Datum, Absender)",
    "Leistungsumfang je Anlage (Klausel „Leistungsumfang“)",
    "Anlagenliste: Kunde, Anlage, Maschinentyp, Baujahr",
    "Preise: Preis je Anlage und Jahr, Anzahl, Jahreswert vor Rabatt, Rabatt, Jahreswert, Angebotswert über die Laufzeit",
    "Vertragsbedingungen (Rahmenvertrag, siehe Klauseln)",
    "Gültigkeit des Angebots: 30 Tage",
    "Quellen aller Zahlen",
    "Vermerk: Entwurf – nicht versendet; Freigabe und Unterschrift durch die verantwortliche Person",
]
KLAUSELN = [
    {"titel": "Vertragsgegenstand", "text": "Wartung der in der Anlagenliste genannten Maschinen."},
    # Leistungsumfang, Reaktionszeit and Ersatzteile depend on the contract level: filled by stufenklauseln()
    {"titel": "Leistungsumfang", "text": ""},
    {"titel": "Reaktionszeit", "text": ""},
    {"titel": "Ersatzteile", "text": ""},
    {"titel": "Preisanpassung", "text": "Jährlich, höchstens 5 %, mit 3 Monaten Ankündigung."},
    {"titel": "Laufzeit und Kündigung", "text": "3 Jahre, danach Verlängerung um je 12 Monate; Kündigung 3 Monate vor Ablauf."},
    {"titel": "Zahlung", "text": "Jährlich im Voraus, 30 Tage netto."},
    {"titel": "Ausschlüsse", "text": "Bedienfehler, Fremdeingriffe, Verschleißteile außerhalb des Satzes, Verbrauchsmaterial."},
    {"titel": "Haftung", "text": "Gemäß unseren Allgemeinen Geschäftsbedingungen; im Entwurf keine Abweichung."},
    {"titel": "Gerichtsstand und AGB", "text": "Gemäß unseren Allgemeinen Geschäftsbedingungen."},
]


LEISTUNGEN = "Unternehmen/leistungen.md"


def leistung(ws: Path, position: str | None, typ: str) -> dict | None:
    """The row of Unternehmen/leistungen.md for the quoted level: price-list position 'Wartungsvertrag Basis MM-400'
    → row 'Wartungsvertrag Basis'. Only the level word is matched; no level word or no row → None."""
    if not position:
        return None
    stufe = re.sub(r"\s+", " ", position.casefold().replace("wartungsvertrag", "").replace(typ.casefold(), "")).strip()
    try:
        text = (ws / LEISTUNGEN).read_text(encoding="utf-8")
    except OSError:
        return None
    text = text.split("<!-- antworten -->", 1)[-1].split("<!-- /antworten -->", 1)[0]
    for zeile in text.splitlines():
        zellen = [z.strip() for z in zeile.strip().strip("|").split("|")]
        if len(zellen) < 2 or not stufe:
            continue
        name = zellen[0].casefold()
        if "wartungsvertrag" in name and re.search(rf"\b{re.escape(stufe)}\b", name):
            return {"leistung": zellen[0], "inhalt": zellen[1]}
    return None


def stufenklauseln(ws: Path, position: str | None, typ: str, meldungen: list[str]) -> list[dict]:
    """KLAUSELN with the level-dependent clauses taken word for word from leistungen.md (never fixed numbers that
    could contradict the company's own levels)."""
    zeile = leistung(ws, position, typ)
    if zeile:
        q = f"({LEISTUNGEN}, {zeile['leistung']})"
        texte = {"Leistungsumfang": f"{zeile['inhalt']} {q}.",
                 "Reaktionszeit": f"Wie im Leistungsumfang {q}.",
                 "Ersatzteile": f"Wie im Leistungsumfang {q}; ist dort nichts genannt, gibt es keinen Rabatt auf "
                                "Ersatzteile."}
    elif position is None:  # no price-list row: the price was calculated from visits per year
        texte = {"Leistungsumfang": f"{BESUCHE_JE_JAHR} planmäßige Wartungsbesuche pro Jahr (Annahme der "
                                    "Kalkulation).",
                 "Reaktionszeit": f"Laut {LEISTUNGEN} – vor dem Versand ergänzen.",
                 "Ersatzteile": f"Laut {LEISTUNGEN} – vor dem Versand ergänzen."}
    else:
        luecke = f"Laut {LEISTUNGEN} für „{position}“ – dort nicht gefunden, vor dem Versand ergänzen."
        texte = dict.fromkeys(("Leistungsumfang", "Reaktionszeit", "Ersatzteile"), luecke)
        meldungen.append(f"„{position}“ steht nicht in {LEISTUNGEN}: Leistungsumfang, Reaktionszeit und "
                         "Ersatzteil-Konditionen im Entwurf vor dem Versand ergänzen.")
    return [k | {"text": texte[k["titel"]]} if k["titel"] in texte else k for k in KLAUSELN]


def grenzen(ws: Path) -> dict:
    fg = definitionen(ws).get("freigabegrenzen") or {}
    return {"angebot_eur": fg.get("angebot_eur"), "rabatt_prozent": fg.get("rabatt_prozent"),
            "standard": bool(fg.get("standard", True))}


def pruefung(ws: Path, angebotswert: float, rabatt: float, abweichungen: list[str]) -> dict:
    """Spec §8 rule 1: above a limit (missing limit = always) → finanzen + qualitaet-recht; deviation → qualitaet-recht."""
    g, finanzen, hinweise = grenzen(ws), [], []
    if g["angebot_eur"] is None:
        finanzen.append("Keine Freigabegrenze für Angebote festgelegt (Unternehmen/freigabegrenzen.md) – Prüfung "
                        "immer nötig.")
    elif angebotswert > zahl(g["angebot_eur"]):
        finanzen.append(f"Angebotswert {deutsch(angebotswert)} EUR liegt über der Freigabegrenze "
                        f"{deutsch(zahl(g['angebot_eur']))} EUR.")
    if g["rabatt_prozent"] is None:
        finanzen.append("Keine Rabattgrenze festgelegt (Unternehmen/freigabegrenzen.md) – Prüfung immer nötig.")
    elif rabatt > zahl(g["rabatt_prozent"]):
        finanzen.append(f"Rabatt {deutsch(rabatt, 1)} % liegt über der Rabattgrenze "
                        f"{deutsch(zahl(g['rabatt_prozent']), 1)} %.")
    recht = (["Großangebot mit Prüfung durch Finanzen – Vertragsbedingungen ebenfalls prüfen."] if finanzen else [])
    recht += [f"Abweichung von den eigenen Bedingungen: {a} – Prüfung unabhängig vom Wert." for a in abweichungen]
    fachexperte = None
    if abweichungen:
        fachexperte = ((definitionen(ws).get("fachexperten") or {}).get("recht")
                       or "nicht benannt – bitte in Unternehmen/fachexperten.md eintragen")
    if g["angebot_eur"] is None:
        try:
            n = sum(1 for _, m, _ in all_cases(ws) if m.get("typ") == "angebot")
        except VorgangFehler:
            n = 0
        if n >= 20:
            hinweise.append(f"Schon {n} Angebotsvorgänge ohne Freigabegrenze – Vorschlag: Grenzen im Onboarding "
                            "festlegen.")
    return {"noetig": bool(finanzen or recht), "finanzen": finanzen, "qualitaet_recht": recht,
            "fachexperte_recht": fachexperte, "grenzen": g, "hinweise": hinweise}


def vertragsstufe(preise: dict, typ: str, stufe: str | None) -> dict | None:
    """K1 (Max, 2026-10-07): the row 'Wartungsvertrag <Typ>' for the quote. One matching row is used directly;
    several rows (Basis / Standard / Premium …) need --stufe, which picks the one row whose Position contains it."""
    treffer = [z for z in preise["zeilen"]
               if "wartungsvertrag" in z["position"].casefold() and typ.casefold() in z["position"].casefold()]
    namen = ", ".join(z["position"] for z in treffer) or "keine"
    datei = preise["datei"] or "04_Angebote/preisliste_*.xlsx fehlt"
    if stufe:
        wahl = [z for z in treffer if stufe.casefold() in z["position"].casefold()]
        if len(wahl) != 1:
            wie = "mehrere" if wahl else "keine"
            raise VertriebFehler(f"Stufe '{stufe}' passt auf {wie} Zeilen 'Wartungsvertrag {typ}' in {datei} – "
                                 f"vorhandene Zeilen: {namen}. Bitte die Stufe eindeutig angeben.")
        return wahl[0]
    if len(treffer) > 1:
        raise StufeFehlt(f"Stufe fehlt: für 'Wartungsvertrag {typ}' stehen {len(treffer)} Stufen in {datei} "
                         f"({namen}) – bitte die Vertragsstufe erfragen und mit --stufe angeben.",
                         [{"position": z["position"], "preis": z["preis"]} for z in treffer])
    return treffer[0] if treffer else None


def listenpreis(preise: dict, typ: str, annahmen: list[str], stufe: str | None = None) -> tuple[dict, str | None]:
    zeile = vertragsstufe(preise, typ, stufe)
    if zeile:
        return wert(f"Preis {zeile['position']} je Anlage und Jahr", zeile["preis"], [zeile["quelle"]]), \
            zeile["position"]
    satz, anfahrt = preis_zeile(preise, "stundensatz"), preis_zeile(preise, "anfahrt")
    if not satz:
        raise VertriebFehler(f"Kein Preis 'Wartungsvertrag {typ}' und kein Techniker-Stundensatz in der Preisliste "
                             f"({preise['datei'] or '04_Angebote/preisliste_*.xlsx fehlt'}) – bitte Preisliste "
                             "ergänzen oder den Preis nennen.")
    b, s = BESUCHE_JE_JAHR, STUNDEN_JE_BESUCH
    annahmen.append(f"Kein Listenpreis für {typ}: {b} Wartungsbesuche je Jahr à {s} Stunden ({STANDARD})")
    q = [satz["quelle"]] + ([anfahrt["quelle"]] if anfahrt else [])
    return wert(f"Preis Wartungsvertrag {typ} je Anlage und Jahr (kalkuliert)",
                round(b * s * satz["preis"] + b * (anfahrt["preis"] if anfahrt else 0.0), 2), q,
                formel=f"{b} Besuche × {s} h × Stundensatz" + (f" + {b} × Anfahrtspauschale" if anfahrt else "")), \
        None


def grossangebot(ws: Path, kunde: str, typ: str, anzahl: int, laufzeit: int | None = None,
                 rabatt: float | None = None, abweichungen: list[str] | tuple = (), heute: str | None = None,
                 stufe: str | None = None) -> dict:
    kunde, typ, abweichungen = kunde.strip(), typ.strip(), [a.strip() for a in abweichungen if a.strip()]
    stufe = (stufe or "").strip() or None
    if not 1 <= anzahl <= 500:
        raise VertriebFehler("--anzahl muss zwischen 1 und 500 liegen")
    tag_heute, annahmen, meldungen = heute_von(heute), [], []
    aufruf = f'grossangebot --kunde "{kunde}" --typ "{typ}" --anzahl {anzahl}'
    if laufzeit is None:
        laufzeit = LAUFZEIT_JAHRE
        annahmen.append(f"Laufzeit {laufzeit} Jahre ({STANDARD})")
    else:
        aufruf += f" --laufzeit-jahre {laufzeit}"
    if not 1 <= laufzeit <= 10:
        raise VertriebFehler("--laufzeit-jahre muss zwischen 1 und 10 liegen")
    if rabatt is None:
        rabatt = next(p for ab, p in RABATT_STAFFEL if anzahl >= ab)
        annahmen.append(f"Mengenrabatt {deutsch(rabatt, 1)} % bei {anzahl} Anlagen (Staffel: ab 2 Anlagen 5 %, ab 5 "
                        f"Anlagen 8 %; {STANDARD})")
    else:
        aufruf += f" --rabatt-prozent {deutsch(rabatt, 1).replace(',0', '')}"
    if not 0 <= rabatt <= 50:
        raise VertriebFehler("--rabatt-prozent muss zwischen 0 und 50 liegen")
    aufruf += f' --stufe "{stufe}"' if stufe else ""
    aufruf += "".join(f' --abweichung "{a}"' for a in abweichungen)
    je, position = listenpreis(preisliste(ws), typ, annahmen, stufe)
    brutto = wert("Jahreswert vor Rabatt", round(je["betrag"] * anzahl, 2), je["quelle"],
                  formel=f"{anzahl} Anlagen × Preis je Anlage und Jahr")
    nachlass = wert("Rabatt je Jahr", round(brutto["betrag"] * rabatt / 100, 2), je["quelle"],
                    formel=f"Jahreswert vor Rabatt × {deutsch(rabatt, 1)} %")
    jahr = wert("Jahreswert", round(brutto["betrag"] - nachlass["betrag"], 2), je["quelle"],
                formel="Jahreswert vor Rabatt − Rabatt je Jahr")
    gesamt = wert("Angebotswert", round(jahr["betrag"] * laufzeit, 2), je["quelle"],
                  formel=f"Jahreswert × {laufzeit} Jahre")
    try:
        eigene = [a for a in installed_base(ws)[0] if gleich(a.get("Kunde"), kunde) and gleich(a.get("Maschinentyp"), typ)]
    except VertriebFehler as exc:
        eigene = []
        meldungen.append(f"{exc} Anlagenliste bitte von Hand ergänzen.")
    anlagen = [{"anlage": text(a.get("Anlage")), "typ": typ, "baujahr": jahrtext(a),
                "vertrag": text(a.get("Vertrag")), "vertragsende": text(a.get("Vertragsende")), "quelle": ort(a)}
               for a in sorted(eigene, key=lambda a: text(a.get("Anlage")))]
    if len(anlagen) != anzahl:
        meldungen.append(f"In der Installed Base stehen {len(anlagen)} Anlagen {typ} von {kunde}; angefragt sind "
                         f"{anzahl} – Anlagenliste im Vertrag bitte prüfen.")
    for a in anlagen:
        if gleich(a["vertrag"], "ja"):
            ende = datum(a["vertragsende"])
            meldungen.append(f"{a['anlage']} hat schon einen Vertrag" + (f" bis {ende:%d.%m.%Y}" if ende else "")
                             + " – Angebot als Verlängerung bzw. Erweiterung formulieren.")
    pr = pruefung(ws, gesamt["betrag"], rabatt, abweichungen)
    gueltig = tag_heute + dt.timedelta(days=GUELTIG_TAGE)
    zeilen = [f"Preis je Anlage und Jahr: {deutsch(je['betrag'])} EUR; {anzahl} Anlagen: {deutsch(brutto['betrag'])} "
              f"EUR; Rabatt {deutsch(rabatt, 1)} %: {deutsch(nachlass['betrag'])} EUR; Jahreswert: "
              f"{deutsch(jahr['betrag'])} EUR.",
              f"Angebotswert über {laufzeit} Jahre: {deutsch(gesamt['betrag'])} EUR (gültig bis {gueltig:%d.%m.%Y}).",
              ("Prüfung nötig: " + " ".join(pr["finanzen"] + pr["qualitaet_recht"])) if pr["noetig"]
              else "Unter allen Freigabegrenzen, keine Abweichung – keine Prüfung nötig."]
    return ergebnis(ws, {
        "kunde": kunde, "typ": typ, "anzahl": anzahl, "stufe": position,
        "werte": {"preis_je_anlage": je, "jahreswert_vor_rabatt": brutto, "rabatt": nachlass, "jahreswert": jahr,
                  "angebotswert": gesamt},
        "rabatt_prozent": rabatt, "laufzeit_jahre": laufzeit, "gueltig_bis": gueltig.isoformat(), "anlagen": anlagen,
        "gliederung": GLIEDERUNG, "klauseln": stufenklauseln(ws, position, typ, meldungen), "pruefung": pr, "aufruf": aufruf, "annahmen": annahmen,
        "meldungen": meldungen,
        "ziel": f"04_Angebote/{ordnername(kunde)}/{tag_heute.isoformat()}_angebot-wartungsvertrag.docx"}, zeilen)


def prozent(s: str | None) -> float | None:
    if s is None:
        return None
    try:
        return zahl(s)
    except ValueError as exc:
        raise VertriebFehler(f"Rabatt '{s}' ist keine Zahl") from exc


# --- CLI ----------------------------------------------------------------------------------------------------------

def parser() -> JsonParser:
    ap = JsonParser(prog="vertrieb")
    sub = ap.add_subparsers(dest="cmd", required=True)

    def add(name: str):
        sp = sub.add_parser(name)
        sp.add_argument("--ws", required=True)
        sp.add_argument("--heute")
        return sp

    r = add("verlaengerungs-radar")
    r.add_argument("--monate", type=int)
    r.add_argument("--stichtag")
    p = add("installed-base-potenziale")
    p.add_argument("--kunde")
    p.add_argument("--stichtag")
    k = add("key-account-review")
    k.add_argument("--kunde")
    k.add_argument("--top", type=int)
    g = add("grossangebot")
    g.add_argument("--kunde", required=True)
    g.add_argument("--typ", required=True)
    g.add_argument("--anzahl", type=int, required=True)
    g.add_argument("--stufe")
    g.add_argument("--laufzeit-jahre", dest="laufzeit", type=int)
    g.add_argument("--rabatt-prozent", dest="rabatt")
    g.add_argument("--abweichung", action="append", default=[])
    return ap


BEFEHLE = {
    "verlaengerungs-radar": lambda a, ws: verlaengerungs_radar(ws, a.monate, a.stichtag, a.heute),
    "installed-base-potenziale": lambda a, ws: installed_base_potenziale(ws, a.kunde, a.stichtag, a.heute),
    "key-account-review": lambda a, ws: key_account_review(ws, a.kunde, a.top, a.heute),
    "grossangebot": lambda a, ws: grossangebot(ws, a.kunde, a.typ, a.anzahl, a.laufzeit, prozent(a.rabatt),
                                               a.abweichung, a.heute, a.stufe),
}


def _main(argv: list[str] | None) -> tuple[int, dict]:
    a = parser().parse_args(argv)
    ws = Path(a.ws)
    try:
        if not (ws / "Unternehmen").is_dir():
            raise VertriebFehler(f"{ws} ist kein Kit-Arbeitsordner (Unternehmen/ fehlt).")
        return 0, BEFEHLE[a.cmd](a, ws)
    except StufeFehlt as exc:
        return 1, {"ok": False, "fehler": [str(exc)], "meldungen": [str(exc)], "stufen": exc.stufen}
    except (VertriebFehler, KennzahlFehler) as exc:
        return 1, {"ok": False, "fehler": [str(exc)], "meldungen": [str(exc)]}


def main(argv: list[str] | None = None) -> int:
    return run(_main, argv)


if __name__ == "__main__":
    raise SystemExit(main())
