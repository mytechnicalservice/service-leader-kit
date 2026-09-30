# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl==3.1.5"]
# ///
"""Regenerates the read-only Excel view of the case ledger (spec §6)."""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from vorgang import all_cases

DATEINAME = "Vorgaenge-Uebersicht (nur Ansicht).xlsx"
HINWEIS = "Nur Ansicht – Änderungen hier werden überschrieben. Sag der Assistenz, was sich ändern soll."
SPALTEN = [("Nr", "nr"), ("Titel", "titel"), ("Typ", "typ"), ("Status", "status"), ("Kunde", "kunde"),
           ("Verantwortlich", "verantwortlich"), ("Fällig", "faellig"), ("Überfällig", None),
           ("Wartet auf", "wartet_auf"), ("Betrag (EUR)", "betrag_eur"), ("Aktualisiert", "aktualisiert")]


def ueberfaellig(meta: dict, heute: str) -> bool:
    return bool(meta.get("faellig")) and meta["faellig"] < heute and meta["status"] != "erledigt"


def fill(sheet, rows: list[dict], heute: str) -> None:
    sheet["A1"] = HINWEIS
    sheet["A1"].font = Font(bold=True, color="C00000")
    for col, (titel, _) in enumerate(SPALTEN, start=1):
        sheet.cell(row=3, column=col, value=titel).font = Font(bold=True)
    for r, meta in enumerate(rows, start=4):
        for col, (_, key) in enumerate(SPALTEN, start=1):
            sheet.cell(row=r, column=col, value=("ja" if ueberfaellig(meta, heute) else "nein") if key is None else meta.get(key))
    sheet.freeze_panes = "A4"
    sheet.protection.sheet = True


def build(ws: Path, heute: str) -> Path:
    offen = [m for _, m, _ in all_cases(ws, ("offen",))]
    offen.sort(key=lambda m: (not ueberfaellig(m, heute), m.get("faellig") or "9999-12-31", m["nr"]))
    erledigt = sorted((m for _, m, _ in all_cases(ws, ("erledigt",))), key=lambda m: m["nr"])
    wb = Workbook()
    fill(wb.active, offen, heute)
    wb.active.title = "Offen"
    fill(wb.create_sheet("Erledigt"), erledigt, heute)
    target = ws / "01_Vorgaenge" / DATEINAME
    target.parent.mkdir(parents=True, exist_ok=True)
    tmp = target.with_name(target.name + ".tmp")
    wb.save(tmp)
    os.replace(tmp, target)
    return target


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="vorgaenge_uebersicht")
    ap.add_argument("--ws", required=True)
    ap.add_argument("--heute", default=dt.date.today().isoformat())
    a = ap.parse_args(argv)
    ws = Path(a.ws)
    path = build(ws, a.heute)
    offen = [m for _, m, _ in all_cases(ws, ("offen",))]
    print(json.dumps({"ok": True, "datei": path.relative_to(ws).as_posix(), "offen": len(offen),
                      "ueberfaellig": sum(ueberfaellig(m, a.heute) for m in offen),
                      "erledigt": len(all_cases(ws, ("erledigt",)))}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
