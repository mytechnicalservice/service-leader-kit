# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl==3.1.5"]
# ///
"""Validates and imports a data export against an import template (spec §7.3)."""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

VORLAGEN = Path(__file__).resolve().parents[1] / "vorlagen" / "import"
MONATE = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober",
          "November", "Dezember"]
DE_TAUSEND = re.compile(r"-?\d{1,3}(\.\d{3})+(,\d+)?")
DE_KOMMA = re.compile(r"-?\d+,\d+")
PLAIN = re.compile(r"-?\d+(\.\d+)?")


def zahl(v) -> float:
    if isinstance(v, bool) or v is None:
        raise ValueError(v)
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("€", "").replace(" ", "").replace(" ", "")
    if DE_TAUSEND.fullmatch(s) or DE_KOMMA.fullmatch(s):
        return float(s.replace(".", "").replace(",", "."))
    if PLAIN.fullmatch(s):
        return float(s)
    raise ValueError(v)


def datum(v) -> dt.date:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(v)


def monat(v) -> str:
    if isinstance(v, (dt.date, dt.datetime)):
        return f"{v.year:04d}-{v.month:02d}"
    s = str(v).strip()
    for pat, fmt in ((r"\d{4}-\d{2}", "{0}"), (r"(\d{2})[./](\d{4})", "{2}-{1}")):
        m = re.fullmatch(pat, s)
        if m:
            return fmt.format(s, *m.groups())
    raise ValueError(v)


def ja_nein(v) -> str:
    s = str(v).strip().casefold()
    if s in {"ja", "j", "x", "1", "true", "wahr"}:
        return "ja"
    if s in {"nein", "n", "", "0", "false", "falsch", "none"}:
        return "nein"
    raise ValueError(v)


PARSER = {"zahl": zahl, "datum": datum, "monat": monat, "ja_nein": ja_nein, "text": lambda v: str(v).strip()}
TYPNAME = {"zahl": "keine Zahl", "datum": "kein Datum", "monat": "kein Monat", "ja_nein": "kein ja/nein"}


def fmt_de(x: float) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def monatsname(periode: str) -> str:
    y, m = periode.split("-")
    return f"{MONATE[int(m) - 1]} {y}"


def load_zuordnung(ws: Path) -> dict:
    p = ws / "Unternehmen" / "datenzuordnung.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}


def pruefe(ws: Path, datei: Path, vorlage: str, kontrollsumme: str | None, zuordnung: list[str],
           uebernehmen: bool, heute: str) -> dict:
    tpl = json.loads((VORLAGEN / f"{vorlage}.json").read_text(encoding="utf-8"))
    result = {"ok": False, "meldungen": [], "zeilen": 0, "duplikate": 0, "periode": None, "summe": None,
              "zuordnung": {}, "ziel": None}
    # not read_only: read-only mode keeps the file open, which blocks the later move on Windows
    rows = [list(r) for r in load_workbook(datei, data_only=True).active.iter_rows(values_only=True)]
    if not rows:
        result["meldungen"].append(f"{datei.name}: Datei ist leer")
        return result
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    gespeichert = load_zuordnung(ws).get(vorlage, {})
    manuell = dict(z.split("=", 1) for z in zuordnung)
    kanon = {s["name"].casefold(): s["name"] for s in tpl["spalten"]}
    syn = {x.casefold(): s["name"] for s in tpl["spalten"] for x in s["synonyme"]}
    mapping = {}
    for h in header:
        target = manuell.get(h) or gespeichert.get(h) or kanon.get(h.casefold()) or syn.get(h.casefold())
        if target:
            mapping[h] = target
    result["zuordnung"] = {h: t for h, t in mapping.items() if h != t}
    fehlend = [s["name"] for s in tpl["spalten"] if s["pflicht"] and s["name"] not in mapping.values()]
    result["meldungen"] += [f"Spalte '{n}' fehlt" for n in fehlend]
    if fehlend:
        return result

    spalten = {s["name"]: s for s in tpl["spalten"]}
    index = {t: header.index(h) for h, t in mapping.items()}
    norm, seen = [], set()
    for zeile, raw in enumerate(rows[1:], start=2):
        if all(c in (None, "") for c in raw):
            continue
        rec = {}
        for name, i in index.items():
            v = raw[i] if i < len(raw) else None
            typ = spalten[name]["typ"]
            if v in (None, "") and not spalten[name]["pflicht"]:
                rec[name] = None
                continue
            try:
                rec[name] = PARSER[typ](v)
            except ValueError:
                result["meldungen"].append(f"Zeile {zeile}, Spalte '{name}': '{v}' ist {TYPNAME[typ]}")
        key = json.dumps(rec, default=str, sort_keys=True)
        if key in seen:
            result["duplikate"] += 1
            continue
        seen.add(key)
        norm.append(rec)
    if result["meldungen"]:
        return result

    if tpl["periode"].get("stichtag"):
        perioden = {heute[:7]}
    else:
        col = tpl["periode"]["spalte"]
        perioden = {r[col] if isinstance(r[col], str) else f"{r[col].year:04d}-{r[col].month:02d}" for r in norm}
    if len(perioden) != 1:
        result["meldungen"].append("Datei enthält mehrere Monate: " + ", ".join(sorted(perioden)))
        return result
    periode = perioden.pop()
    result["periode"], result["zeilen"] = periode, len(norm)
    if tpl["summe"]:
        result["summe"] = round(sum(r[tpl["summe"]] or 0 for r in norm), 2)
    if kontrollsumme is not None and result["summe"] is not None:
        soll = round(zahl(kontrollsumme), 2)
        if abs(soll - result["summe"]) > 0.005:
            result["meldungen"].append(f"Summe {tpl['summe']} {fmt_de(result['summe'])} passt nicht zur "
                                       f"Kontrollsumme {fmt_de(soll)}")
            return result
    ziel = ws / "07_Daten" / f"{vorlage}_{periode}.csv"
    if ziel.exists():
        result["meldungen"].append(f"{monatsname(periode)} ist bereits importiert ({ziel.relative_to(ws).as_posix()}). "
                                   "Datei bleibt im Eingang.")
        return result
    result["ok"] = True
    if uebernehmen:
        ziel.parent.mkdir(parents=True, exist_ok=True)
        tmp = ziel.with_name(ziel.name + ".tmp")
        with tmp.open("w", encoding="utf-8", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=list(index))
            w.writeheader()
            for r in norm:
                w.writerow({k: (v.isoformat() if isinstance(v, dt.date) else v) for k, v in r.items()})
        os.replace(tmp, ziel)
        orig = ws / "07_Daten" / "original" / datei.name
        orig.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(datei), orig)
        alle = load_zuordnung(ws)
        eigene = {h: t for h, t in mapping.items() if h != t and (h in manuell or h in gespeichert)}
        if eigene:
            alle[vorlage] = gespeichert | eigene
            (ws / "Unternehmen").mkdir(exist_ok=True)
            (ws / "Unternehmen" / "datenzuordnung.json").write_text(
                json.dumps(alle, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        result["ziel"] = ziel.relative_to(ws).as_posix()
    return result


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="daten_pruefen")
    ap.add_argument("--ws", required=True)
    ap.add_argument("--datei", required=True)
    ap.add_argument("--vorlage", required=True, choices=sorted(p.stem for p in VORLAGEN.glob("*.json")))
    ap.add_argument("--kontrollsumme")
    ap.add_argument("--zuordnung", action="append", default=[])
    ap.add_argument("--uebernehmen", action="store_true")
    ap.add_argument("--heute", default=dt.date.today().isoformat())
    a = ap.parse_args(argv)
    r = pruefe(Path(a.ws), Path(a.datei), a.vorlage, a.kontrollsumme, a.zuordnung, a.uebernehmen, a.heute)
    print(json.dumps(r, ensure_ascii=False, default=str))
    return 0 if r["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
