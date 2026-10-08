# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl==3.1.5", "python-docx==1.1.2", "python-pptx==1.0.2"]
# ///
"""Personal & Qualifikation (Plan 4h): staffing, skill matrix, key-person risk, talk prep, person-data guards.
Team level only (spec §1, §9.3): nothing this script returns or writes holds a value per person."""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import math
import re
import sys
from pathlib import Path

import kennzahlen as kz
from arbeitsordner import lies_konfig
from daten_pruefen import ImportFehler, datum, leer, lese, monat
from slk_common import JsonParser, csv_sicher, run, write_atomic, zahl

HINWEIS = ("Hinweis: Personalthemen können die Mitbestimmung des Betriebsrats (§ 87 Abs. 1 Nr. 6, § 94 und § 98 "
           "BetrVG) und den Beschäftigtendatenschutz (DSGVO, BDSG) berühren. Bitte Betriebsrat und "
           "Datenschutzbeauftragte(n) einbeziehen. Das Kit gibt keine rechtliche Freigabe.")  # Max's correction K2
MARKE = "Standardannahme des Kits – bitte prüfen"
BEISPIEL = "Beispieldaten – Muster Maschinenbau GmbH"
STANDARD = {"netto_stunden": 1600.0, "auslastung": 0.75, "schwelle_einstellung": 0.5, "vollkosten_eur": 85000.0,
            "einmalkosten_eur": 15000.0, "vorlauf_monate": 5, "einarbeitung_monate": 3, "einarbeitung_quote": 0.5,
            "min_qualifiziert": 2, "min_teamgroesse": 3}  # plan 4h, domain defaults 1–15
TEXT = {"netto_stunden": "Netto-Jahresarbeitszeit je FTE (Std)", "auslastung": "Auslastungsziel",
        "schwelle_einstellung": "Einstellungsschwelle (FTE)", "vollkosten_eur": "Vollkosten je Techniker und Jahr",
        "einmalkosten_eur": "Einmalkosten je Einstellung", "vorlauf_monate": "Vorlaufzeit Einstellung (Monate)",
        "einarbeitung_monate": "Einarbeitung (Monate)", "einarbeitung_quote": "Produktivität in der Einarbeitung"}
PROZENT = {"auslastung", "einarbeitung_quote"}
NK = {"FTE": 1, "EUR/Std": 2}


def fehler(*texte: str) -> dict:
    return {"ok": False, "fehler": list(texte), "meldungen": list(texte)}


def monat_plus(m: str, n: int) -> str:
    y, mo = map(int, m.split("-"))
    i = y * 12 + mo - 1 + n
    return f"{i // 12:04d}-{i % 12 + 1:02d}"


def fenster(bis: str, monate: int) -> list[str]:
    return [monat_plus(bis, -i) for i in range(monate - 1, -1, -1)]


def bereiche(nums: list[int]) -> str:
    nums = sorted(set(nums))
    teile, start = [], nums[0]
    for a, b in zip(nums, nums[1:] + [None]):
        if b != a + 1:
            teile.append(f"{start}" if start == a else f"{start}–{a}")
            start = b
    return ", ".join(teile)


def zeilen_text(rows: list[dict]) -> list[str]:
    """'<datei> Zeilen 2–4, 7' per file, for values the script derives itself."""
    je: dict[str, list[int]] = {}
    for r in rows:
        je.setdefault(r["_datei"], []).append(int(r["_zeile"]))
    return [f"{d} Zeilen {bereiche(n)}" for d, n in sorted(je.items())]


def summe(zeilen: list[dict], spalte: str, name: str, einheit: str) -> dict:
    if not zeilen:
        return kz.wert(name, 0.0, ["keine passenden Zeilen"], None, einheit)
    w = kz.summe(zeilen, spalte, name)
    w["einheit"] = einheit
    return w


def neueste(rows: list[dict]) -> list[dict]:
    """Rows of the latest file only (stichtag templates: installed base, qualification)."""
    if not rows:
        return []
    letzte = max(r["_datei"] for r in rows)
    return [r for r in rows if r["_datei"] == letzte]


def lade_optional(ws: Path, vorlage: str, perioden: list[str] | None = None) -> list[dict]:
    try:
        return kz.lade(ws, vorlage, perioden)
    except kz.KennzahlFehler:
        return []


def letzter_monat(ordner: Path, vorlage: str) -> str:
    monate = sorted(p.stem.rsplit("_", 1)[1] for p in (ordner / "07_Daten").glob(f"{vorlage}_*.csv"))
    if not monate:
        raise kz.KennzahlFehler(f"Keine geprüften Daten '{vorlage}' in 07_Daten/ – bitte Exporte zuerst mit "
                                "daten-pruefen übernehmen.")
    return monate[-1]


def kontext(ws: Path) -> dict:
    q = kz.datenquelle(ws)
    defs = kz.definitionen(ws)
    ds = (defs.get("fachexperten") or {}).get("datenschutz")
    hinweis = HINWEIS + (f" Ansprechperson Datenschutz laut Unternehmen/fachexperten.md: {ds}." if ds else "")
    return {"defs": defs, "beispiel": bool(q["beispiel"]), "ordner": Path(q["ordner"]), "hinweis": hinweis}


def zeile(w: dict) -> str:
    s = f"{w['name']}: {kz.deutsch(w['betrag'], NK.get(w['einheit'], 0))} {w['einheit']}"
    if w["berechnet"]:
        s += f" (berechnet: {w['formel']})"
    return s + " – Quelle: " + "; ".join(w["quelle"])


def abschnitt(titel: str, eintraege: list) -> dict:
    return {"titel": titel, "zeilen": [zeile(e) if isinstance(e, dict) else e for e in eintraege]}


def kopf(k: dict) -> list[dict]:
    return [abschnitt("Kennzeichnung", [BEISPIEL])] if k["beispiel"] else []


def kleine_teams(kap: list[dict]) -> list[str]:
    g = STANDARD["min_teamgroesse"]
    return [f"Team {r['Team']} hat weniger als {g} Köpfe – Zahlen dieses Teams lassen Rückschlüsse auf Einzelne zu; "
            "nur für die Teamplanung nutzen." for r in kap if zahl(r["Techniker_Anzahl"]) < g]


def paare(eintraege: list[str]) -> dict[str, float]:
    out = {}
    for e in eintraege:
        team, _, z = e.partition("=")
        try:
            out[team.strip()] = zahl(z.strip())
        except ValueError as exc:
            raise ValueError(f"'{e}' braucht die Form Team=Zahl") from exc
    return out


def kpi_auslastung(defs: dict) -> float | None:
    """Target billable ÷ present hours: the KPI "Verrechenbarkeit" (Max's correction K1). The KPI "Auslastung
    Techniker" is Ist ÷ Soll hours (4a's kapazitaet-lage) and is deliberately not read here."""
    sek = defs.get("kpi-ziele") or {}
    if sek.get("standard", True):
        return None
    for k in sek.get("kennzahlen") or []:
        if not isinstance(k, dict):
            continue
        if "verrechenbar" in str(k.get("name", "")).casefold() and k.get("ziel") is not None:
            z = float(k["ziel"])
            return z / 100 if z > 1 else z
    return None


def parameter(a, defs: dict) -> tuple[dict, list[dict]]:
    p, herkunft = dict(STANDARD), {k: MARKE for k in STANDARD}
    ziel = kpi_auslastung(defs)
    if ziel is not None:
        p["auslastung"], herkunft["auslastung"] = ziel, "Unternehmen/kpi-ziele.md"
    firma = (defs.get("ergebnisrechnung") or {}).get("personal")
    if isinstance(firma, dict):  # Plan 3: ergebnisrechnung.md front matter `personal: {…}` (4h-G4)
        for quelle, k in (("vollkosten_techniker_eur", "vollkosten_eur"), ("netto_stunden", "netto_stunden")):
            if firma.get(quelle) not in (None, ""):
                p[k], herkunft[k] = float(zahl(firma[quelle])), "Unternehmen/ergebnisrechnung.md"
    for k in ("netto_stunden", "vollkosten_eur", "einmalkosten_eur"):
        if getattr(a, k, None) is not None:
            p[k], herkunft[k] = getattr(a, k), "Angabe im Gespräch"
    if getattr(a, "auslastung_prozent", None) is not None:
        p["auslastung"], herkunft["auslastung"] = a.auslastung_prozent / 100, "Angabe im Gespräch"
    anzeige = {k: f"{kz.deutsch(p[k] * 100)} %" if k in PROZENT else kz.deutsch(p[k], 1 if k == "schwelle_einstellung"
                                                                                  else 0) for k in TEXT}
    return p, [{"name": TEXT[k], "wert": p[k], "anzeige": anzeige[k], "herkunft": herkunft[k]} for k in TEXT]


def entscheidungsrecht(defs: dict, betrag: float) -> str:
    er = (defs.get("ergebnisrechnung") or {}).get("entscheidungsrechte") or []
    for e in er if isinstance(er, list) else [er]:
        if isinstance(e, dict) and e.get("thema") == "personal":
            grenze = e.get("allein_bis_eur")
            if grenze is not None and betrag <= float(grenze):
                return f"Im eigenen Entscheidungsrahmen (bis {kz.deutsch(float(grenze))} EUR, Unternehmen/ergebnisrechnung.md)."
            return f"Entscheidung liegt bei: {e.get('sonst') or 'nicht festgelegt'} (Unternehmen/ergebnisrechnung.md)."
    return "Entscheidungsrecht für Personal ist noch nicht festgelegt (Onboarding) – als Vorgang zur Entscheidung vorlegen."


def business_case(t, n, umsatz, std, luecke, p, std_fte, defs, heute) -> tuple[dict, list[dict]]:
    q = umsatz["quelle"] + std["quelle"]
    erl_h = round(umsatz["betrag"] / std["betrag"], 2)
    em, eq = p["einarbeitung_monate"], p["einarbeitung_quote"]
    monat_std, kosten_monat = std_fte * n / 12, p["vollkosten_eur"] * n / 12
    kum, amort = -p["einmalkosten_eur"] * n, 0
    for m in range(1, 25):
        kum += monat_std * (eq if m <= em else 1.0) * erl_h - kosten_monat
        if not amort and kum >= 0:
            amort = m
    w = [kz.wert(f"Erlös je Stunde {t}", erl_h, q, "Umsatz / Auftragsstunden des Teams", "EUR/Std"),
         kz.wert(f"Kosten Jahr 1 {t}", round((p["vollkosten_eur"] + p["einmalkosten_eur"]) * n, 2),
                 ["Annahmen dieser Planung"], f"{n} × (Vollkosten + Einmalkosten)"),
         kz.wert(f"Erlös Jahr 1 {t}", round(monat_std * (em * eq + 12 - em) * erl_h, 2), q,
                 f"verrechenbare Std je Monat × ({em} Monate × {kz.deutsch(eq * 100)} % + {12 - em} Monate) × Erlös je Stunde")]
    w += [kz.wert(f"Deckungsbeitrag Jahr 1 {t}", round(w[2]["betrag"] - w[1]["betrag"], 2), q,
                  "Erlös Jahr 1 − Kosten Jahr 1"),
          kz.wert(f"Deckungsbeitrag Folgejahr {t}", round((std_fte * erl_h - p["vollkosten_eur"]) * n, 2), q,
                  "verrechenbare Std je FTE × Erlös je Stunde − Vollkosten"),
          kz.wert(f"Amortisation Monat {t}", float(amort), q,
                  "erster Monat ab Eintritt mit kumuliert ≥ 0 (0 = nicht innerhalb von 24 Monaten)", "Monat"),
          kz.wert(f"Umsatzwert ungedeckter Bedarf {t}", round(luecke * std_fte * erl_h, 2), q,
                  "Lücke FTE × verrechenbare Std je FTE × Erlös je Stunde")]
    fall = {"team": t, "einstellungen": n, "eintritt_fruehestens": monat_plus(heute[:7], p["vorlauf_monate"]),
            "entscheidung": entscheidungsrecht(defs, w[1]["betrag"])}
    return fall, w


def antwort_block(teams: list[str], werte: list[dict], faelle: list[dict], hinweis: str) -> str:
    """The chat block the skill copies unchanged: team table, each hiring case, the note (eval run 0.2.4)."""
    w = {x["name"]: x["betrag"] for x in werte}
    z = ["| Team | Jahresbedarf Stunden | Bedarf FTE | Köpfe | Lücke | Einstellungen |", "|---|---|---|---|---|---|"]
    z += [f"| {t} | {kz.deutsch(w[f'Jahresbedarf Stunden {t}'])} | {kz.deutsch(w[f'FTE-Bedarf {t}'], 1)} | "
          f"{kz.deutsch(w[f'Köpfe {t}'])} | {kz.deutsch(w[f'Lücke FTE {t}'], 1)} | {kz.deutsch(w[f'Einstellungen {t}'])} |"
          for t in teams]
    for f in faelle:
        t, monat_ = f["team"], int(w[f"Amortisation Monat {f['team']}"])
        z += ["", f"Einstellung Team {t} ({f['einstellungen']} Servicetechniker): Kosten Jahr 1: "
                  f"{kz.deutsch(w[f'Kosten Jahr 1 {t}'])} EUR · Erlös Jahr 1: {kz.deutsch(w[f'Erlös Jahr 1 {t}'])} EUR"
                  f" · Amortisation: " + (f"Monat {monat_}" if monat_ else "nicht innerhalb von 24 Monaten")
                  + f" · Eintritt frühestens: {f['eintritt_fruehestens']}"]
    return "\n".join(z + ["", f"Hinweis: {hinweis}"])


def cmd_personalplanung(a, ws: Path) -> tuple[int, dict]:
    k = kontext(ws)
    bis = monat(a.bis) if a.bis else letzter_monat(k["ordner"], "auftraege")
    perioden = fenster(bis, a.monate)
    auf, kap = kz.lade(ws, "auftraege", perioden), kz.lade(ws, "kapazitaet", [bis])
    p, annahmen = parameter(a, k["defs"])
    zusatz, abgang = paare(a.zusatz), paare(a.abgang)
    std_fte, faktor = p["netto_stunden"] * p["auslastung"], 12 / a.monate
    werte, faelle, bc, meldungen, budget = [], [], [], kleine_teams(kap), {}
    teams = sorted({str(r["Team"]) for r in kap if not leer(r.get("Team"))})
    for t in teams:
        rows = [r for r in auf if r.get("Team") == t]
        std = summe(rows, "Stunden", f"Auftragsstunden {t}", "Std")
        umsatz = summe(rows, "Umsatz_EUR", f"Umsatz {t}", "EUR")
        koepfe = summe([r for r in kap if r.get("Team") == t], "Techniker_Anzahl", f"Köpfe {t}", "Köpfe")
        bedarf_std = std["betrag"] * faktor * (1 + a.wachstum_prozent / 100) + zusatz.get(t, 0.0)
        bedarf = kz.wert(f"Jahresbedarf Stunden {t}", round(bedarf_std, 1), std["quelle"],
                         f"Auftragsstunden × 12/{a.monate} × (1 + {kz.deutsch(a.wachstum_prozent, 1)} %) + "
                         f"{kz.deutsch(zusatz.get(t, 0.0))} Std Zusatz", "Std")
        fte = kz.wert(f"FTE-Bedarf {t}", round(bedarf_std / std_fte, 1), std["quelle"],
                      f"Jahresbedarf / ({kz.deutsch(p['netto_stunden'])} Std × {kz.deutsch(p['auslastung'] * 100)} %)",
                      "FTE")
        lb = round(fte["betrag"] - koepfe["betrag"] + abgang.get(t, 0.0), 1)
        luecke = kz.wert(f"Lücke FTE {t}", lb, std["quelle"] + koepfe["quelle"],
                         "FTE-Bedarf − Köpfe + bekannte Abgänge", "FTE")
        n = math.floor(lb + 0.5) if lb >= p["schwelle_einstellung"] else 0
        werte += [std, umsatz, koepfe, bedarf, fte, luecke,
                  kz.wert(f"Einstellungen {t}", float(n), luecke["quelle"],
                          f"Lücke ab {kz.deutsch(p['schwelle_einstellung'], 1)} FTE kaufmännisch gerundet", "Köpfe")]
        if lb < 0:
            meldungen.append(f"Team {t}: Überdeckung {kz.deutsch(-lb, 1)} FTE – keine Personalmaßnahme vorschlagen; "
                             "Einsatz für andere Teams oder zusätzliche Aufträge prüfen.")
        budget[t] = int(koepfe["betrag"] - abgang.get(t, 0.0) + n)
        if n and std["betrag"] > 0:
            fall, w = business_case(t, n, umsatz, std, lb, p, std_fte, k["defs"], a.heute)
            faelle.append(fall)
            werte += w
            bc.append(abschnitt(f"Einstellung Team {t}: {n} Servicetechniker", w + [
                f"Eintritt frühestens: {fall['eintritt_fruehestens']} (Vorlaufzeit {p['vorlauf_monate']} Monate)",
                fall["entscheidung"]]))
    gesamt = [w for w in werte if w["name"].startswith(("FTE-Bedarf", "Köpfe"))]
    werte.append(kz.wert("Lücke FTE gesamt", round(sum(w["betrag"] for w in gesamt if w["einheit"] == "FTE")
                                                    - sum(w["betrag"] for w in gesamt if w["einheit"] == "Köpfe"), 1),
                         [q for w in gesamt for q in w["quelle"]], "Summe FTE-Bedarf − Summe Köpfe", "FTE"))
    ohne = [r for r in auf if leer(r.get("Team"))]
    if ohne:
        w = summe(ohne, "Stunden", "Auftragsstunden ohne Team", "Std")
        werte.append(w)
        meldungen.append(f"{len(ohne)} Aufträge ohne Team ({kz.deutsch(w['betrag'])} Std) – keinem Team zugeordnet, "
                         "nicht verteilt.")
    meldungen += [f"Team {t} hat Aufträge, aber keine Kapazitätszeile für {bis}." for t in
                  sorted({str(r["Team"]) for r in auf if not leer(r.get("Team"))} - set(teams))]
    meldungen.append("Köpfe = Techniker_Anzahl; Teilzeit ist im Export nicht erkennbar (1 Kopf = 1 FTE).")
    mehr = sum(f["einstellungen"] for f in faelle) * p["vollkosten_eur"]
    zeitraum = f"{perioden[0]} bis {bis}"
    gl = kopf(k) + [abschnitt(f"Personalplanung je Team – Basis {zeitraum}", []),
                    abschnitt("Annahmen", [f"{x['name']}: {x['anzeige']} – {x['herkunft']}" for x in annahmen]),
                    abschnitt("Bedarf und Lücke je Team", [w for w in werte if w["einheit"] in ("Std", "FTE", "Köpfe")]),
                    *bc, abschnitt("Hinweise", meldungen + [k["hinweis"]])]
    return 0, {"ok": True, "beispiel": k["beispiel"], "zeitraum": zeitraum, "hinweis": k["hinweis"],
               "antwort": antwort_block(teams, werte, faelle, k["hinweis"]), "annahmen": annahmen, "werte": werte, "einstellungen": faelle,
               "fuer_budget": {"koepfe_plan": budget, "mehrkosten_eur": mehr}, "meldungen": meldungen,
               "gliederung": gl}


PERSON = re.compile(r"^(?:(?:vor|nach|mitarbeiter|techniker|monteur|bearbeiter)?name|mitarbeiter(?:in)?|mitarbeitende"
                    r"|techniker(?:in)?|monteur(?:in)?|bearbeiter(?:in)?|person|(?:personal|pers|ma)[\s._-]*(?:nr|nummer))\.?$")
SP = {"team": ("team", "gruppe", "region"), "monat": ("monat", "periode"),
      "soll": ("soll", "soll_stunden", "soll-stunden", "sollstunden", "soll stunden", "kapazität"),
      "ist": ("ist", "ist_stunden", "ist-stunden", "iststunden", "ist stunden", "geleistet")}


def _ist_zahl(v) -> bool:
    try:
        zahl(v)
        return True
    except ValueError:
        return False


def tabellen(datei: Path) -> list[tuple[list[str], list[list], int]]:
    out = []
    for _, rows in lese(datei):
        for i, r in enumerate(rows[:10]):
            if sum(not leer(c) for c in r) >= 2:
                out.append(([("" if leer(c) else str(c).strip()) for c in r], rows[i + 1:], i + 2))
                break
    return out


def personen_spalten(kopf: list[str], rows: list[list]) -> list[str]:
    out = []
    for i, h in enumerate(kopf):
        if not PERSON.match(h.casefold()):
            continue
        werte = [r[i] for r in rows if i < len(r) and not leer(r[i])]
        if h.casefold().startswith(("techniker", "monteur")) and werte and all(_ist_zahl(v) for v in werte):
            continue  # a head count, not a person
        out.append(h)
    return out


def personen_werte(datei: Path) -> set[str]:
    namen = set()
    for kopf, rows, _ in tabellen(datei):
        for s in personen_spalten(kopf, rows):
            i = kopf.index(s)
            namen |= {str(r[i]).strip() for r in rows if i < len(r) and not leer(r[i]) and len(str(r[i]).strip()) >= 4}
    return namen


def text_aus(datei: Path) -> list[tuple[str, str]]:
    e = datei.suffix.lower()
    if e in (".md", ".txt", ".csv"):
        return [(f"Zeile {i}", z) for i, z in enumerate(datei.read_text(encoding="utf-8-sig").splitlines(), 1)]
    if e == ".docx":
        from docx import Document
        d = Document(str(datei))
        out = [(f"Absatz {i}", p.text) for i, p in enumerate(d.paragraphs, 1)]
        return out + [(f"Tabelle {t}", c.text) for t, tab in enumerate(d.tables, 1) for r in tab.rows for c in r.cells]
    if e == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(datei, data_only=True)
        return [(f"{sh.title}!{c.coordinate}", str(c.value)) for sh in wb.worksheets for row in sh.iter_rows()
                for c in row if c.value is not None]
    if e == ".pptx":
        from pptx import Presentation
        out = []
        for i, s in enumerate(Presentation(str(datei)).slides, 1):
            for sh in s.shapes:
                if sh.has_text_frame:
                    out.append((f"Folie {i}", sh.text_frame.text))
                if getattr(sh, "has_table", False):
                    out += [(f"Folie {i}", c.text) for r in sh.table.rows for c in r.cells]
        return out
    raise ImportFehler(f"{datei.name}: nur .md, .docx, .xlsx oder .pptx können geprüft werden.")


def cmd_personenbezug(a, ws) -> tuple[int, dict]:
    spalten, zeilen = [], 0
    for kopf, rows, _ in tabellen(Path(a.datei)):
        sp = personen_spalten(kopf, rows)
        if sp:
            spalten += sp
            zeilen += sum(1 for r in rows if any(not leer(c) for c in r))
    m = ([f"Die Datei enthält Personenbezug (Spalten: {', '.join(spalten)}). Nur über team-aggregat verwenden."]
         if spalten else [])
    return 0, {"ok": True, "personenbezug": bool(spalten), "spalten": spalten, "zeilen": zeilen, "meldungen": m}


def cmd_team_aggregat(a, ws: Path) -> tuple[int, dict]:
    datei = Path(a.datei)
    tabs = [(k_, r_, s_) for k_, r_, s_ in tabellen(datei) if personen_spalten(k_, r_)]
    if not tabs:
        return 1, fehler(f"{datei.name} enthält keine Spalte mit Namen oder Personalnummern – bitte direkt mit "
                         "daten-pruefen übernehmen.")
    kopf, rows, start = tabs[0]
    low = [h.casefold() for h in kopf]
    idx = {key: next((i for i, h in enumerate(low) if h in namen), None) for key, namen in SP.items()}
    fehlt = [key for key, i in idx.items() if i is None]
    if fehlt:
        return 1, fehler(f"Spalten fehlen: {', '.join(fehlt)} (vorhanden: {', '.join(h for h in kopf if h)})")
    pers = personen_spalten(kopf, rows)
    pid = kopf.index(next((s for s in pers if re.search(r"nr|nummer", s.casefold())), pers[0]))
    teams, monate = {}, set()
    for n, r in enumerate(rows, start=start):
        if all(leer(c) for c in r):
            continue
        try:
            m, soll, ist = monat(r[idx["monat"]]), zahl(r[idx["soll"]]), zahl(r[idx["ist"]])
        except (ValueError, IndexError):
            return 1, fehler(f"Zeile {n}: Monat, Soll oder Ist ist nicht lesbar.")
        if leer(r[idx["team"]]):
            return 1, fehler(f"Zeile {n}: Team fehlt.")
        e = teams.setdefault(str(r[idx["team"]]).strip(), {"personen": set(), "soll": 0.0, "ist": 0.0, "zeilen": []})
        e["personen"].add(str(r[pid]).strip())
        e["soll"], e["ist"] = e["soll"] + soll, e["ist"] + ist
        e["zeilen"].append(n)
        monate.add(m)
    if len(monate) != 1:
        return 1, fehler("Datei enthält mehrere Monate: " + ", ".join(sorted(monate)))
    m = monate.pop()
    ziel = ws / "00_Eingang" / f"{datei.stem}_je_team.csv"
    rel = ziel.relative_to(ws).as_posix()
    if ziel.exists():
        return 1, fehler(f"{rel} gibt es schon – das Kit überschreibt nichts.")
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["Monat", "Team", "Techniker_Anzahl", "Soll_Stunden", "Ist_Stunden"])
    werte, meldungen = [], []
    for t in sorted(teams):
        e, z = teams[t], (lambda x: int(x) if float(x).is_integer() else x)
        w.writerow([m, csv_sicher(t), len(e["personen"]), z(e["soll"]), z(e["ist"])])
        q = [f"{datei.name} Zeilen {bereiche(e['zeilen'])}"]
        werte += [kz.wert(f"Köpfe {t} ({m})", float(len(e["personen"])), q, "verschiedene Personen im Team", "Köpfe"),
                  kz.wert(f"Ist-Stunden {t} ({m})", e["ist"], q, "Summe je Team", "Std")]
        if len(e["personen"]) < STANDARD["min_teamgroesse"]:
            meldungen.append(f"Team {t} hat weniger als {STANDARD['min_teamgroesse']} Köpfe – Zahlen dieses Teams "
                             "lassen Rückschlüsse auf Einzelne zu.")
    write_atomic(ziel, buf.getvalue(), encoding="utf-8-sig")
    meldungen.insert(0, f"{datei.name} enthält Namen oder Personalnummern. Übernommen werden nur Teamwerte ({rel}); "
                        "die Originaldatei bleibt in 00_Eingang – ob sie gelöscht wird, entscheidest du.")
    return 0, {"ok": True, "datei": rel, "monat": m, "werte": werte, "meldungen": meldungen,
               "hinweis": kontext(ws)["hinweis"]}


def namensquellen(ws: Path) -> list[Path]:
    """Tables in 00_Eingang/ and 07_Daten/ that hold a name or personnel-number column (read here, never typed)."""
    out = []
    for f in sorted(p for o in ("00_Eingang", "07_Daten") if (ws / o).is_dir() for p in (ws / o).rglob("*")
                    if p.suffix.lower() in (".csv", ".xlsx")):
        try:
            if any(personen_spalten(k_, r_) for k_, r_, _ in tabellen(f)):
                out.append(f)
        except (ImportFehler, ValueError, OSError):
            continue
    return out


def cmd_pruefe_ausgabe(a, ws) -> tuple[int, dict]:
    namen = {n.strip() for n in a.name if len(n.strip()) >= 4}
    for q in a.namen_aus + (namensquellen(ws) if ws else []):
        namen |= personen_werte(Path(q))
    if not namen and ws:
        return 0, {"ok": True, "treffer": 0, "fundstellen": [], "namen_geprueft": 0,
                   "meldungen": ["Im Arbeitsordner liegt keine Datei mit Namen oder Personalnummern – es gab nichts "
                                 "zum Abgleichen."]}
    if not namen:
        return 1, fehler("Keine Namen zum Prüfen – --ws, --namen-aus oder --name angeben.")
    muster = [re.compile(rf"(?<!\w){re.escape(n)}(?!\w)", re.I) for n in namen]
    funde = [stelle for stelle, text in text_aus(Path(a.datei)) if any(m.search(text) for m in muster)]
    m = [] if not funde else [f"Das Dokument nennt Personen ({len(funde)} Stellen: {', '.join(funde)}). Bitte dort nur "
                              "Teamwerte schreiben und erneut prüfen."]
    return (1 if funde else 0), {"ok": not funde, "treffer": len(funde), "fundstellen": funde,
                                 "namen_geprueft": len(namen), "meldungen": m}


STUFEN = ((3, "A – robust (3 und mehr)"), (2, "B – abgedeckt (2)"), (1, "C – Einzelwissen (1)"),
          (0, "D – nicht abgedeckt (0)"))
FELDER = {"team", "maschinentyp", "auftragsart", "qualifiziert", "in_schulung", "ausbilder", "stufe"}
WISSENSTRANSFER = ["Anlagenwissen je Kunde dokumentieren (Besonderheiten, Störungshistorie, Einstellwerte)",
                   "Begleitung bei Einsätzen: Nachfolge oder Schulungsteilnehmende fahren mit",
                   "Kundenkontakte übergeben (Ansprechpartner, offene Zusagen)",
                   "Zugänge, Werkzeuge, Software und Messmittel übergeben"]


def stufe(q: int) -> str:
    return next(t for g, t in STUFEN if q >= g)


def anzahl(r: dict | None, spalte: str) -> int:
    return 0 if r is None or leer(r.get(spalte)) else int(zahl(r[spalte]))


def matrixdaten(ws: Path, perioden: list[str]):
    auf = kz.lade(ws, "auftraege", perioden)
    ib, qual = neueste(lade_optional(ws, "installed_base")), neueste(lade_optional(ws, "qualifikation"))
    typ = {(r["Kunde"], r.get("Anlage")): r["Maschinentyp"] for r in ib}
    zellen, ohne = {}, []
    for r in auf:
        mt = typ.get((r["Kunde"], r.get("Anlage")))
        if leer(r.get("Team")) or mt is None:
            ohne.append(r)
        else:
            zellen.setdefault((str(r["Team"]), mt, str(r["Auftragsart"])), []).append(r)
    return zellen, ohne, ib, qual


def quali(qual: list[dict], team: str, mt: str, art: str) -> dict | None:
    passend = [r for r in qual if r["Team"] == team and r["Maschinentyp"] == mt]
    return (next((r for r in passend if r["Auftragsart"] == art), None)
            or next((r for r in passend if str(r["Auftragsart"]).casefold() == "alle"), None))


def ohne_matrix() -> tuple[int, dict]:
    return 1, fehler("Noch keine Qualifikationsmatrix in 07_Daten/ – bitte zuerst mit der Skill-Matrix erfassen "
                     "(Anzahlen je Team).")


def cmd_skill_matrix(a, ws: Path) -> tuple[int, dict]:
    k = kontext(ws)
    bis = monat(a.bis) if a.bis else letzter_monat(k["ordner"], "auftraege")
    zellen, ohne, ib, qual = matrixdaten(ws, fenster(bis, a.monate))
    if not qual:
        return ohne_matrix()
    mind, werte, matrix, bedarf = STANDARD["min_qualifiziert"], [], [], []
    for (t, mt, art), rows in sorted(zellen.items()):
        w = summe(rows, "Stunden", f"Auftragsstunden {t} · {mt} · {art}", "Std")
        r = quali(qual, t, mt, art)
        q = anzahl(r, "Qualifiziert_Anzahl")
        z = {"team": t, "maschinentyp": mt, "auftragsart": art, "stunden": w["betrag"], "qualifiziert": q,
             "in_schulung": anzahl(r, "In_Schulung_Anzahl"), "ausbilder": anzahl(r, "Ausbilder_Anzahl"),
             "stufe": stufe(q), "quelle": w["quelle"] + (zeilen_text([r]) if r else ["nicht in der Matrix"])}
        werte.append(w)
        matrix.append(z)
        if q < mind:
            bedarf.append(z | {"fehlend": mind - q})
    mit_auftraegen = {(z["team"], z["maschinentyp"]) for z in matrix}
    for r in qual:
        if (r["Team"], r["Maschinentyp"]) not in mit_auftraegen:
            q = anzahl(r, "Qualifiziert_Anzahl")
            matrix.append({"team": r["Team"], "maschinentyp": r["Maschinentyp"], "auftragsart": r["Auftragsart"],
                           "stunden": 0.0, "qualifiziert": q, "in_schulung": anzahl(r, "In_Schulung_Anzahl"),
                           "ausbilder": anzahl(r, "Ausbilder_Anzahl"), "stufe": stufe(q), "quelle": zeilen_text([r])})
    bedarf.sort(key=lambda z: -z["stunden"])
    gedeckt = {r["Maschinentyp"] for r in qual if anzahl(r, "Qualifiziert_Anzahl") > 0}
    ohne_abdeckung = sorted({r["Maschinentyp"] for r in ib} - gedeckt)
    werte += [kz.wert("Zellen Stufe C", float(sum(1 for z in matrix if z["stunden"] > 0 and z["qualifiziert"] == 1)),
                      zeilen_text(qual), "Zellen mit Aufträgen und genau 1 qualifizierten Person", "Zellen"),
              kz.wert("Fehlende Qualifizierungen", float(sum(z["fehlend"] for z in bedarf)), zeilen_text(qual),
                      f"Summe (mindestens {mind} − qualifiziert) über Zellen mit Aufträgen", "Anzahl")]
    meldungen = kleine_teams(lade_optional(ws, "kapazitaet", [bis]))
    if ohne:
        meldungen.append(f"{len(ohne)} Aufträge ohne Team oder ohne Maschinentyp in der Installed Base – nicht in "
                         "der Matrix.")
    meldungen += [f"Maschinentyp {m}: in keinem Team jemand qualifiziert." for m in ohne_abdeckung]
    gl = kopf(k) + [
        abschnitt(f"Skill-Matrix je Team – Aufträge {fenster(bis, a.monate)[0]} bis {bis}",
                  [f"{z['team']} · {z['maschinentyp']} · {z['auftragsart']}: {z['stufe']}, {z['in_schulung']} in "
                   f"Schulung, {z['ausbilder']} Ausbilder, {kz.deutsch(z['stunden'])} Auftragsstunden – Quelle: "
                   + "; ".join(z["quelle"]) for z in matrix]),
        abschnitt("Schulungsbedarf (unter 2 qualifiziert, nach Auftragsstunden)",
                  [f"{z['team']} · {z['maschinentyp']} · {z['auftragsart']}: {z['fehlend']} fehlt, "
                   f"{z['in_schulung']} in Schulung, {kz.deutsch(z['stunden'])} Std" for z in bedarf]),
        abschnitt("Kennzahlen", werte[-2:]), abschnitt("Hinweise", meldungen + [k["hinweis"]])]
    return 0, {"ok": True, "beispiel": k["beispiel"], "matrix": matrix, "schulungsbedarf": bedarf,
               "ohne_abdeckung": ohne_abdeckung, "werte": werte, "meldungen": meldungen, "hinweis": k["hinweis"],
               "gliederung": gl}


def eingabe(quelle: str):
    text = sys.stdin.read() if quelle == "-" else Path(quelle).read_text(encoding="utf-8")
    try:
        return json.loads(text)
    except json.JSONDecodeError as exc:
        raise ValueError("Die Eingabe ist kein gültiges JSON.") from exc


def cmd_matrix_erfassen(a, ws: Path) -> tuple[int, dict]:
    eintraege = eingabe(a.eingabe)
    if not isinstance(eintraege, list) or not eintraege:
        return 1, fehler("Die Eingabe muss eine Liste von Zeilen sein.")
    summen: dict[tuple, dict] = {}
    for nr, e in enumerate(eintraege, 1):
        fremd = sorted(set(e) - FELDER) if isinstance(e, dict) else ["?"]
        if fremd:
            return 1, fehler(f"Eintrag {nr}: Feld {', '.join(fremd)} wird nicht angenommen. Die Skill-Matrix speichert "
                             "nur Anzahlen je Team – keine Namen und keine Einzelbewertungen.")
        if any(leer(e.get(f)) for f in ("team", "maschinentyp", "auftragsart")):
            return 1, fehler(f"Eintrag {nr}: Team, Maschinentyp und Auftragsart sind Pflicht.")
        s = summen.setdefault((str(e["team"]).strip(), str(e["maschinentyp"]).strip(), str(e["auftragsart"]).strip()),
                              {"qualifiziert": 0, "in_schulung": 0, "ausbilder": 0})
        if "stufe" in e:
            st = e["stufe"]
            if isinstance(st, bool) or st not in (0, 1, 2, 3):
                return 1, fehler(f"Eintrag {nr}: Stufe muss 0, 1, 2 oder 3 sein.")
            s["qualifiziert"] += int(st >= 2)
            s["in_schulung"] += int(st == 1)
            s["ausbilder"] += int(st == 3)
            continue
        for f in ("qualifiziert", "in_schulung", "ausbilder"):
            v = e.get(f, 0)
            if not isinstance(v, int) or isinstance(v, bool) or v < 0:
                return 1, fehler(f"Eintrag {nr}: {f} muss eine ganze Zahl ab 0 sein.")
            s[f] += v
    ziel = ws / "00_Eingang" / f"qualifikation_erfasst_{a.heute}.csv"
    rel = ziel.relative_to(ws).as_posix()
    if ziel.exists():
        return 1, fehler(f"{rel} gibt es schon – das Kit überschreibt nichts. Bitte zuerst übernehmen.")
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["Team", "Maschinentyp", "Auftragsart", "Qualifiziert_Anzahl", "In_Schulung_Anzahl", "Ausbilder_Anzahl"])
    zeilen = [{"team": t, "maschinentyp": m, "auftragsart": x, **s} for (t, m, x), s in sorted(summen.items())]
    # The evaluation reads only the newest month: carry the other cells of the latest state over unchanged.
    alt = neueste(lade_optional(ws, "qualifikation"))
    bleiben = [{"team": str(r["Team"]), "maschinentyp": str(r["Maschinentyp"]), "auftragsart": str(r["Auftragsart"]),
                "qualifiziert": anzahl(r, "Qualifiziert_Anzahl"), "in_schulung": anzahl(r, "In_Schulung_Anzahl"),
                "ausbilder": anzahl(r, "Ausbilder_Anzahl")} for r in alt
               if (str(r["Team"]), str(r["Maschinentyp"]), str(r["Auftragsart"])) not in summen]
    for z in zeilen + bleiben:
        w.writerow([csv_sicher(z["team"]), csv_sicher(z["maschinentyp"]), csv_sicher(z["auftragsart"]),
                    z["qualifiziert"], z["in_schulung"], z["ausbilder"]])
    write_atomic(ziel, buf.getvalue(), encoding="utf-8-sig")
    meldungen = [f"Gespeichert sind nur Anzahlen je Team ({rel}). Jetzt mit daten-pruefen als Vorlage "
                 "'qualifikation' übernehmen."]
    if bleiben:
        meldungen.append(f"{len(bleiben)} weitere Zeilen unverändert aus {alt[0]['_datei']} mitgenommen, damit der "
                         "neue Stand vollständig ist.")
    return 0, {"ok": True, "datei": rel, "zeilen": zeilen, "uebernommen": len(bleiben), "meldungen": meldungen}


def cmd_abgang(a, ws: Path) -> tuple[int, dict]:
    k = kontext(ws)
    bis = monat(a.bis) if a.bis else letzter_monat(k["ordner"], "auftraege")
    letzter = datum(a.letzter_tag).isoformat()
    zellen, _, ib, qual = matrixdaten(ws, fenster(bis, a.monate))
    if not qual:
        return ohne_matrix()
    mind, faktor = STANDARD["min_qualifiziert"], 12 / a.monate
    ziele: set[tuple] = set()
    for eintrag in a.qualifikation:
        mt, _, art = (s.strip() for s in eintrag.partition("|"))
        art = art or "alle"
        kandidaten = {c for c in zellen if c[0] == a.team and c[1] == mt} | {
            (r["Team"], r["Maschinentyp"], r["Auftragsart"]) for r in qual
            if r["Team"] == a.team and r["Maschinentyp"] == mt and str(r["Auftragsart"]).casefold() != "alle"}
        ziele |= {c for c in kandidaten if art.casefold() == "alle" or c[2] == art} or (
            {(a.team, mt, art)} if art.casefold() != "alle" else set())
    if not ziele:
        return 1, fehler(f"Für Team {a.team} gibt es weder Aufträge noch Matrixzeilen zu {', '.join(a.qualifikation)}.")
    andere = sorted({str(r["Team"]) for r in qual} - {a.team})
    out_z, rows0, rows1, typen, schulung, abordnung, meldungen = [], [], [], set(), [], [], []
    for t, mt, art in sorted(ziele):
        r = quali(qual, t, mt, art)
        vor = anzahl(r, "Qualifiziert_Anzahl")
        nach = max(vor - 1, 0)
        if vor == 0:
            meldungen.append(f"{mt} · {art}: laut Matrix ist in Team {t} niemand qualifiziert – Matrix aktualisieren.")
        firma = vor + sum(anzahl(quali(qual, o, mt, art), "Qualifiziert_Anzahl") for o in andere)
        rows = zellen.get((t, mt, art), [])
        out_z.append({"maschinentyp": mt, "auftragsart": art, "team_vorher": vor, "team_nachher": nach,
                      "firma_nachher": firma - (vor - nach), "stufe_nachher": stufe(nach)})
        if nach < mind:
            (rows0 if nach == 0 else rows1).extend(rows)
            typen.add(mt)
        if anzahl(r, "In_Schulung_Anzahl"):
            schulung.append({"maschinentyp": mt, "auftragsart": art, "anzahl": anzahl(r, "In_Schulung_Anzahl")})
        abordnung += [{"team": o, "maschinentyp": mt, "auftragsart": art, "qualifiziert": q} for o in andere
                      if (q := anzahl(quali(qual, o, mt, art), "Qualifiziert_Anzahl")) >= mind + 1]
    werte = []
    for rows, art_text in ((rows0, "nicht abgedeckt"), (rows1, "Einzelwissen")):
        u = summe(rows, "Umsatz_EUR", f"Umsatz {art_text} (Zeitraum)", "EUR")
        s = summe(rows, "Stunden", f"Stunden {art_text} (Zeitraum)", "Std")
        werte += [u, s, kz.wert(f"Umsatzrisiko {art_text} (Jahr)", round(u["betrag"] * faktor, 2), u["quelle"],
                                f"Umsatz im Zeitraum × 12/{a.monate}"),
                  kz.wert(f"Stunden {art_text} (Jahr)", round(s["betrag"] * faktor, 1), s["quelle"],
                          f"Stunden im Zeitraum × 12/{a.monate}", "Std")]
    betroffen = [r for r in ib if r["Maschinentyp"] in typen]
    kunden = [{"kunde": kd, "anlagen": sum(r["Kunde"] == kd for r in betroffen),
               "mit_vertrag": sum(r["Kunde"] == kd and str(r.get("Vertrag")).casefold() == "ja" for r in betroffen),
               "vertragsende_frueh": min((str(r["Vertragsende"]) for r in betroffen if r["Kunde"] == kd
                                          and not leer(r.get("Vertragsende"))), default=None)}
              for kd in sorted({r["Kunde"] for r in betroffen})]
    qib = zeilen_text(betroffen) if betroffen else ["keine passenden Anlagen"]
    werte += [kz.wert("Betroffene Anlagen", float(len(betroffen)), qib, "Anlagen der betroffenen Maschinentypen", "Anlagen"),
              kz.wert("Anlagen mit Vertrag", float(sum(x["mit_vertrag"] for x in kunden)), qib, "Vertrag = ja", "Anlagen"),
              kz.wert("Betroffene Kunden", float(len(kunden)), qib, "verschiedene Kunden", "Kunden")]
    bereit = monat_plus(a.heute[:7], STANDARD["vorlauf_monate"] + STANDARD["einarbeitung_monate"])
    lm = monat(letzter[:7])
    luecke = max(0, (int(bereit[:4]) * 12 + int(bereit[5:])) - (int(lm[:4]) * 12 + int(lm[5:])) - 1)
    optionen = {"in_schulung": schulung, "abordnung": abordnung,
                "einstellung": {"bereit_ab": bereit, "luecke_monate": luecke}}
    gl = kopf(k) + [
        abschnitt(f"Abgang einer Fachkraft – Team {a.team}, letzter Arbeitstag {letzter}",
                  [f"{z['maschinentyp']} · {z['auftragsart']}: Team {z['team_vorher']} → {z['team_nachher']} "
                   f"({z['stufe_nachher']}), Firma danach {z['firma_nachher']}" for z in out_z]),
        abschnitt("Auswirkung", werte),
        abschnitt("Kunden mit betroffenen Maschinentypen",
                  [f"{x['kunde']}: {x['anlagen']} Anlagen, {x['mit_vertrag']} mit Vertrag"
                   + (f", frühestes Vertragsende {x['vertragsende_frueh']}" if x["vertragsende_frueh"] else "")
                   for x in kunden]),
        abschnitt("Optionen", [f"Schulung abschließen: {s['anzahl']} in Schulung ({s['maschinentyp']} · {s['auftragsart']})"
                               for s in schulung]
                  + [f"Abordnung aus Team {o['team']} ({o['qualifiziert']} qualifiziert)" for o in abordnung]
                  + [f"Einstellung: einsatzbereit frühestens {bereit} – {luecke} Monate ohne Abdeckung",
                     "Überbrückung durch Herstellerservice oder Fremdleistung"]),
        abschnitt(f"Wissenstransfer bis {letzter}", WISSENSTRANSFER),
        abschnitt("Hinweise", meldungen + kleine_teams(lade_optional(ws, "kapazitaet", [bis])) + [k["hinweis"]])]
    return 0, {"ok": True, "beispiel": k["beispiel"], "zellen": out_z, "kunden": kunden, "optionen": optionen,
               "werte": werte, "meldungen": meldungen, "hinweis": k["hinweis"], "gliederung": gl}


GESPRAECHE = {
    "jahresgespraech": [
        ("rueckblick", "1. Rückblick auf das Jahr", ["Welche Aufgaben und Ergebnisse waren aus deiner Sicht wichtig?",
                                                     "Was lief gut, was nicht?"]),
        ("zusammenarbeit", "2. Zusammenarbeit und Rahmenbedingungen", ["Wie läuft die Zusammenarbeit im Team und mit dir?",
                                                                        "Was behindert die Arbeit?"]),
        ("sicht_mitarbeiter", "3. Sicht der Mitarbeiterin / des Mitarbeiters",
         ["Mit welchen offenen Fragen holst du die Sicht deines Gegenübers ein?"]),
        ("entwicklung", "4. Entwicklung und Qualifizierung", ["Welche Schulung oder nächste Aufgabe ist sinnvoll?"]),
        ("ziele", "5. Ziele für das nächste Jahr", ["Welche 3–5 Ziele, messbar und mit Termin?"]),
        ("vereinbarungen", "6. Vereinbarungen und nächste Schritte", ["Was wird bis wann von wem erledigt?"])],
    "zielgespraech": [
        ("zielerreichung", "1. Zielerreichung", ["Welche Ziele waren vereinbart, wie weit sind sie aus deiner Sicht erreicht?"]),
        ("rahmen", "2. Einflüsse und Rahmenbedingungen", ["Was hat die Zielerreichung beeinflusst, das nicht in der Hand "
                                                          "der Person lag?"]),
        ("neue_ziele", "3. Neue Ziele", ["Welche 3–5 Ziele, messbar und mit Termin?"]),
        ("unterstuetzung", "4. Unterstützung", ["Was braucht es von dir oder vom Unternehmen dafür?"]),
        ("vereinbarungen", "5. Vereinbarungen und nächste Schritte", ["Was wird bis wann von wem erledigt?"])]}
EXPORT_SPUREN = re.compile(r"07_Daten|00_Eingang|_quelle_|Ist_Stunden|Soll_Stunden|Techniker_Anzahl|Umsatz_EUR"
                           r"|\bSA-\d{4,}|\.(?:csv|xlsx|xlsm)\b", re.I)
BESONDERE = re.compile(r"krank|diagnose|therapie|schwanger|behinderung|religion|gewerkschaft", re.I)
GESCHUETZT = ("00_Eingang", "01_Vorgaenge", "07_Daten", "Unternehmen")


def zielpfad(ws: Path, ziel: str) -> tuple[str, list[str], bool]:
    voll = Path(ziel) if Path(ziel).is_absolute() else ws / ziel
    probleme = []
    if voll.suffix.lower() not in (".docx", ".md"):
        probleme.append("Bitte als .docx oder .md speichern.")
    try:
        rel = voll.resolve().relative_to(ws.resolve()).as_posix()
    except ValueError:
        rel = None
    if rel and rel.split("/")[0] in GESCHUETZT:
        probleme.append(f"Nicht in {rel.split('/')[0]}/ speichern – der Ordner ist für Daten, Vorgänge oder "
                        "Firmenwissen reserviert.")
    if voll.exists():
        probleme.append(f"{rel or voll} gibt es schon – das Kit überschreibt nichts. Bitte einen neuen Namen nennen.")
    return (rel or str(voll)), probleme, rel is not None


def cmd_gespraech(a, ws: Path) -> tuple[int, dict]:
    daten = eingabe(a.eingabe)
    abschnitte = daten.get("abschnitte") if isinstance(daten, dict) else None
    if not isinstance(abschnitte, dict):
        return 1, fehler('Die Eingabe braucht die Form {"abschnitte": {...}}.')
    erlaubt = [s for s, _, _ in GESPRAECHE[a.art]]
    fremd = sorted(set(abschnitte) - set(erlaubt))
    if fremd:
        return 1, fehler(f"Unbekannte Abschnitte: {', '.join(fremd)} – erlaubt: {', '.join(erlaubt)}.")
    text = " ".join(str(v) for v in abschnitte.values())
    if EXPORT_SPUREN.search(text):
        return 1, fehler("Die Vorbereitung nutzt nur, was du selbst im Gespräch eingibst – keine Daten aus Exporten "
                         "oder Dateien zur Person (spec §9.3). Bitte die Stelle in eigenen Worten beschreiben.")
    ziel, probleme, im_ws = zielpfad(ws, a.ziel)
    if probleme:
        return 1, fehler(*probleme)
    warnungen = []
    if BESONDERE.search(text):
        warnungen.append("Gesundheits- oder andere besondere Angaben (Art. 9 DSGVO) gehören nicht in die "
                         "Vorbereitung – sie werden weggelassen; bitte mit HR klären.")
    konf, _ = lies_konfig(ws)
    if im_ws and konf and konf.get("ablage") in ("cloud", "github"):
        warnungen.append("Der Ordner liegt in einer Cloud oder auf GitHub – Personalunterlagen dort nur mit Freigabe "
                         "der IT und des Datenschutzes ablegen (spec §9.4).")
    gl = []
    for key, titel, fragen in GESPRAECHE[a.art]:
        t = str(abschnitte.get(key) or "").strip()
        gl.append({"titel": titel, "text": t, "offene_fragen": [] if t else fragen})
    return 0, {"ok": True, "art": a.art, "ziel": ziel, "gliederung": gl, "warnungen": warnungen,
               "hinweis": kontext(ws)["hinweis"], "meldungen": warnungen}


BEFEHLE = {"personalplanung": cmd_personalplanung, "personenbezug": cmd_personenbezug,
           "team-aggregat": cmd_team_aggregat, "pruefe-ausgabe": cmd_pruefe_ausgabe,
           "skill-matrix": cmd_skill_matrix, "matrix-erfassen": cmd_matrix_erfassen, "abgang": cmd_abgang,
           "gespraech": cmd_gespraech}


def parser() -> JsonParser:
    ap = JsonParser(prog="personal")
    sub = ap.add_subparsers(dest="cmd", required=True)

    def add(name: str, ws: bool = True, fenster_args: bool = False):
        sp = sub.add_parser(name)
        if ws:
            sp.add_argument("--ws", required=True)
        sp.add_argument("--heute", default=dt.date.today().isoformat())
        if fenster_args:
            sp.add_argument("--bis")
            sp.add_argument("--monate", type=int, default=12)
        return sp

    sp = add("personalplanung", fenster_args=True)
    sp.add_argument("--wachstum-prozent", dest="wachstum_prozent", type=float, default=0.0)
    sp.add_argument("--zusatz", action="append", default=[])
    sp.add_argument("--abgang", action="append", default=[])
    for f in ("netto-stunden", "auslastung-prozent", "vollkosten-eur", "einmalkosten-eur"):
        sp.add_argument(f"--{f}", dest=f.replace("-", "_"), type=float)
    sp = add("personenbezug", ws=False)
    sp.add_argument("--datei", required=True)
    sp = add("team-aggregat")
    sp.add_argument("--datei", required=True)
    sp = add("pruefe-ausgabe", ws=False)
    sp.add_argument("--ws")
    sp.add_argument("--datei", required=True)
    sp.add_argument("--name", action="append", default=[])
    sp.add_argument("--namen-aus", dest="namen_aus", action="append", default=[])
    add("skill-matrix", fenster_args=True)
    sp = add("matrix-erfassen")
    sp.add_argument("--eingabe", required=True)
    sp = add("abgang", fenster_args=True)
    sp.add_argument("--team", required=True)
    sp.add_argument("--qualifikation", action="append", required=True)
    sp.add_argument("--letzter-tag", dest="letzter_tag", required=True)
    sp = add("gespraech")
    sp.add_argument("--art", required=True, choices=sorted(GESPRAECHE))
    sp.add_argument("--ziel", required=True)
    sp.add_argument("--eingabe", required=True)
    return ap


def _main(argv: list[str] | None) -> tuple[int, dict]:
    a = parser().parse_args(argv)
    if getattr(a, "monate", 1) < 1:
        return 1, fehler("--monate muss mindestens 1 sein.")
    ws = Path(a.ws) if getattr(a, "ws", None) else None
    try:
        return BEFEHLE[a.cmd](a, ws)
    except (kz.KennzahlFehler, ImportFehler, ValueError) as exc:
        return 1, fehler(str(exc))


def main(argv: list[str] | None = None) -> int:
    return run(_main, argv)


if __name__ == "__main__":
    sys.exit(main())
