import csv
import json
import os
import re
import shutil
import subprocess

import pytest
import yaml
from docx import Document
from openpyxl import Workbook

import erwartet_finanzen as ef
import finanzen
import kennzahlen as kz
import vorgang
from conftest import ROOT, runner_env

KOPF = {
    "ergebnis": ["Monat", "Position", "Plan_EUR", "Ist_EUR"],
    "auftraege": ["Auftragsnr", "Kunde", "Anlage", "Auftragsart", "Eingang", "Abschluss", "Status", "Stunden",
                  "Umsatz_EUR", "Kosten_EUR", "Team"],
    "kapazitaet": ["Monat", "Team", "Techniker_Anzahl", "Soll_Stunden", "Ist_Stunden"],
    "installed_base": ["Kunde", "Anlage", "Maschinentyp", "Baujahr", "Vertrag", "Vertragsende"],
}
# September: Umsatz 114.000 (Plan 120.000), Material 36.000 (30.000), Personal 41.000 (40.000), Gewährl. 9.000 (2.000)
SEPT = [["2026-09", "Umsatz Service", 100000, 89000], ["2026-09", "Umsatz Ersatzteile", 20000, 25000],
        ["2026-09", "Material", -30000, -36000], ["2026-09", "Personalkosten", -40000, -41000],
        ["2026-09", "Gewährleistung", -2000, -9000]]
AUFTRAEGE = [["SA-1", "Müller GmbH", "Anlage 1", "Wartung", "2026-09-02", "2026-09-03", "abgeschlossen", 8, 1000, 600, "Nord"],
             ["SA-2", "Hansa Pack AG", "Anlage 2", "Reparatur", "2026-09-10", "2026-09-11", "abgeschlossen", 4, 800, 500, "Süd"],
             ["SA-3", "Nordmetall GmbH", "Anlage 1", "Retrofit", "2026-09-20", "2026-09-25", "abgeschlossen", 10, 2000, 1500, "Süd"]]
KAP = [["2026-09", "Nord", 2, 300, 270], ["2026-09", "Süd", 1, 150, 120]]
HEUTE = "2026-10-06"


def schreibe_csv(ordner, vorlage, periode, zeilen):
    """A validated import exactly as daten_pruefen writes it: UTF-8 with BOM, comma, source columns."""
    p = ordner / "07_Daten" / f"{vorlage}_{periode}.csv"
    p.parent.mkdir(parents=True, exist_ok=True)
    with p.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(KOPF[vorlage] + ["_quelle_datei", "_quelle_blatt", "_quelle_zeile"])
        for i, z in enumerate(zeilen, start=2):
            w.writerow([*z, f"{vorlage}_{periode}.xlsx", "Tabelle1", i])
    return p


def monat_zeilen(rows, monat):
    return [[monat, *r[1:]] for r in rows]


def rufe(*argv):
    return finanzen._main([str(a) for a in argv])


def kern(out, name):
    return next(k for k in out["kernzahlen"] if k["name"] == name)


@pytest.fixture
def fin_ws(kit_ws):
    schreibe_csv(kit_ws, "ergebnis", "2026-09", SEPT)
    schreibe_csv(kit_ws, "auftraege", "2026-09", AUFTRAEGE)
    schreibe_csv(kit_ws, "kapazitaet", "2026-09", KAP)
    return kit_ws


def test_inputs_in_german_formats():
    assert finanzen.eingabe("45.000,00 €") == 45000
    assert finanzen.eingabe("1,2 Mio.") == 1200000
    assert finanzen.eingabe("120 T€") == 120000
    assert finanzen.prozent("12,5 %") == 12.5
    with pytest.raises(finanzen.FinanzFehler):
        finanzen.eingabe("ungefähr viel")


def test_classes_and_periods():
    assert [finanzen.klasse(p) for p in ("Umsatz Ersatzteile", "Material", "Personalkosten", "Gewährleistung",
                                         "Miete Werkstatt")] == ["umsatz", "material", "personal", "gewaehrleistung",
                                                                 "sonstige"]
    assert finanzen.bereich("2026-07..2026-09") == ["2026-07", "2026-08", "2026-09"]
    assert finanzen.ytd("2026-02", 1) == ["2026-01", "2026-02"]
    assert finanzen.ytd("2026-02", 10) == ["2025-10", "2025-11", "2025-12", "2026-01", "2026-02"]


def test_report_numbers_come_from_the_data(fin_ws):
    code, out = rufe("management-report", "--ws", fin_ws, "--monat", "2026-09", "--heute", HEUTE)
    assert code == 0, out
    assert kern(out, "Umsatz Monat Ist")["betrag"] == 114000
    assert kern(out, "Umsatz Monat Plan")["betrag"] == 120000
    assert kern(out, "DB I Monat Ist")["betrag"] == 37000  # D19: 114.000 − 36.000 − 41.000
    assert kern(out, "Ergebnis Monat Ist")["betrag"] == 28000
    assert kern(out, "Ergebnis Monat Plan")["betrag"] == 48000
    assert kern(out, "DB I in % vom Umsatz Monat")["anzeige"] == f"{kz.deutsch(37000 / 114000 * 100, 1)} %"
    assert kern(out, "Auftragseingang Monat")["betrag"] == 3800
    assert kern(out, "Auslastung Monat")["anzeige"] == f"{kz.deutsch(390 / 450 * 100, 1)} %"
    assert kern(out, "Umsatz Monat Ist")["quelle"][0].startswith("07_Daten/ergebnis_2026-09.csv")
    assert kern(out, "DB I Monat Ist")["berechnet"] is True and "Material" in kern(out, "DB I Monat Ist")["formel"]
    assert [t["name"] for t in out["auslastung_teams"]] == ["Auslastung Team Nord", "Auslastung Team Süd"]


def test_page_one_margin_target_is_db1(fin_ws):
    """D19: the margin target is the DB I target (DB I after service personnel); DB II % is shown without one."""
    _, out = rufe("management-report", "--ws", fin_ws, "--monat", "2026-09", "--heute", HEUTE)
    namen = [k["name"] for k in out["kernzahlen"]]
    assert kern(out, "DB I in % vom Umsatz Monat")["anzeige"] == "32,5 %"  # 37.000 / 114.000
    assert kern(out, "Ziel DB I in %")["betrag"] == 35.0
    assert kern(out, "Ziel DB I in %")["quelle"] == [finanzen.STANDARD_HINWEIS]
    # DB II = DB I 37.000 − Gewährleistung 9.000 = 28.000
    assert kern(out, "DB II in % vom Umsatz Monat")["anzeige"] == f"{kz.deutsch(28000 / 114000 * 100, 1)} %"
    assert not any(n.startswith("Ziel DB II") for n in namen), namen  # no target line on DB II


def test_margin_target_reads_the_db1_kpi_and_ignores_db2(kit_ws):
    (kit_ws / "Unternehmen" / "kpi-ziele.md").write_text(
        "---\nkennzahlen:\n"
        '  - {"name": "DB II-Marge", "formel": "DB II / Umsatz", "quelle": "ergebnis", "ziel": 32, "einheit": "%"}\n'
        '  - {"name": "DB I-Marge", "formel": "DB I / Umsatz", "quelle": "ergebnis", "ziel": 40, "einheit": "%"}\n'
        "---\n\n# KPIs und Ziele\n", encoding="utf-8")
    assert finanzen.db1_ziel(finanzen.definitionen(kit_ws)) == (40.0, False)
    (kit_ws / "Unternehmen" / "kpi-ziele.md").write_text(
        "---\nkennzahlen:\n"
        '  - {"name": "DB II-Marge", "formel": "DB II / Umsatz", "quelle": "ergebnis", "ziel": 32, "einheit": "%"}\n'
        "---\n\n# KPIs und Ziele\n", encoding="utf-8")
    assert finanzen.db1_ziel(finanzen.definitionen(kit_ws)) == (finanzen.DB1_ZIEL_STANDARD, True)


def test_db1_target_is_the_kit_standard_until_kpis_are_set(kit_ws):
    """Labelling rule (D19): the kit's 35 % is never reported as the company's own target."""
    assert finanzen.db1_ziel(finanzen.definitionen(kit_ws)) == (35.0, True)
    (kit_ws / "Unternehmen" / "kpi-ziele.md").write_text("---\nabweichung_massnahme_eur: 8000\n---\n", encoding="utf-8")
    assert finanzen.db1_ziel(finanzen.definitionen(kit_ws)) == (35.0, True)  # thresholds set, KPI list still standard
    shutil.copy(ROOT / "plugin" / "beispiel" / "Unternehmen" / "kpi-ziele.md", kit_ws / "Unternehmen" / "kpi-ziele.md")
    assert finanzen.db1_ziel(finanzen.definitionen(kit_ws)) == (35.0, False)


def test_warranty_borne_elsewhere_is_not_deducted_in_db2(fin_ws):
    """D19 edge: with gewaehrleistung_traeger produkt or qualitaet, warranty is not service's cost – DB II = DB I."""
    for traeger in ("produkt", "qualitaet"):
        (fin_ws / "Unternehmen" / "ergebnisrechnung.md").write_text(
            f"---\ngewaehrleistung_traeger: {traeger}\n---\n# Ergebnisrechnung\n", encoding="utf-8")
        _, out = rufe("management-report", "--ws", fin_ws, "--monat", "2026-09", "--heute", HEUTE)
        zeilen = {z["zeile"]: z for z in out["guv"]["monat"]}
        assert zeilen["DB I"]["ist"]["betrag"] == zeilen["DB II"]["ist"]["betrag"] == 37000, traeger
        assert zeilen["DB II"]["plan"]["betrag"] == 50000  # 120.000 − 30.000 − 40.000, no warranty deducted
        assert traeger in zeilen["DB II"]["ist"]["formel"]
        assert kern(out, "Ergebnis Monat Ist")["betrag"] == 37000


def test_variances_follow_the_thresholds(fin_ws):
    _, out = rufe("management-report", "--ws", fin_ws, "--monat", "2026-09", "--heute", HEUTE)
    by_pos = {a["position"]: a for a in out["abweichungen"]}
    assert set(by_pos) == {"Umsatz Service", "Umsatz Ersatzteile", "Material", "Gewährleistung"}  # Personal: 2,5 %
    assert {p for p, a in by_pos.items() if a["massnahme_pruefen"]} == {"Umsatz Service", "Material", "Gewährleistung"}
    assert by_pos["Umsatz Ersatzteile"]["richtung"] == "günstig"
    assert by_pos["Umsatz Service"]["differenz"]["betrag"] == -11000 and out["massnahmen_offen"] == 3
    assert out["abweichungen"][0]["position"] == "Umsatz Service"  # largest first


def test_standard_definitions_are_flagged(fin_ws):
    _, out = rufe("management-report", "--ws", fin_ws, "--monat", "2026-09", "--heute", HEUTE)
    erste = out["definitionen"][0]
    assert erste["datei"] == "Unternehmen/ergebnisrechnung.md" and erste["hinweis"] == finanzen.STANDARD_HINWEIS
    assert any("abweichung_massnahme_eur" in d["hinweis"] for d in out["definitionen"])
    assert [g["kapitel"] for g in out["gliederung"]][0] == "Auf einen Blick"


def test_company_threshold_changes_the_action_list(fin_ws):
    (fin_ws / "Unternehmen" / "kpi-ziele.md").write_text(
        "---\nabweichung_massnahme_eur: 8000\n---\n\n# KPIs und Ziele\n", encoding="utf-8")
    _, out = rufe("management-report", "--ws", fin_ws, "--monat", "2026-09", "--heute", HEUTE)
    assert {a["position"] for a in out["abweichungen"] if a["massnahme_pruefen"]} == {"Umsatz Service"}
    assert not any("abweichung_massnahme_eur" in d["hinweis"] for d in out["definitionen"])


def test_positive_cost_sign_gives_same_result(kit_ws):
    positiv = [[m, p, abs(pl), abs(i)] for m, p, pl, i in SEPT]
    schreibe_csv(kit_ws, "ergebnis", "2026-09", positiv)
    _, out = rufe("management-report", "--ws", kit_ws, "--monat", "2026-09", "--heute", HEUTE)
    assert kern(out, "Ergebnis Monat Ist")["betrag"] == 28000
    assert {a["position"] for a in out["abweichungen"] if a["massnahme_pruefen"]} == {"Umsatz Service", "Material",
                                                                                       "Gewährleistung"}


def test_missing_plan_is_named_not_zero(kit_ws):
    rows = [r[:] for r in SEPT]
    rows[2][2] = ""  # Material without plan
    schreibe_csv(kit_ws, "ergebnis", "2026-09", rows)
    code, out = rufe("management-report", "--ws", kit_ws, "--monat", "2026-09", "--heute", HEUTE)
    assert code == 0
    assert "Material" not in {a["position"] for a in out["abweichungen"]}
    assert any("Plan fehlt für: Material" in d for d in out["datenlage"])
    assert not any(k["name"] == "DB I Monat Plan" for k in out["kernzahlen"])


def test_missing_months_drop_the_cumulated_column(fin_ws):
    _, out = rufe("management-report", "--ws", fin_ws, "--monat", "2026-09", "--heute", HEUTE)
    assert out["guv"]["kumuliert"] is None
    assert any("kumuliert nicht berechenbar" in d for d in out["datenlage"])


def test_cumulated_year_to_date(fin_ws):
    for m in range(1, 9):
        schreibe_csv(fin_ws, "ergebnis", f"2026-{m:02d}", monat_zeilen(SEPT, f"2026-{m:02d}"))
    _, out = rufe("management-report", "--ws", fin_ws, "--monat", "2026-09", "--heute", HEUTE)
    assert kern(out, "Umsatz kumuliert Ist")["betrag"] == 9 * 114000
    assert kern(out, "Ergebnis kumuliert Plan")["betrag"] == 9 * 48000


def test_missing_month_is_an_error(fin_ws):
    code, out = rufe("management-report", "--ws", fin_ws, "--monat", "2026-08", "--heute", HEUTE)
    assert code == 1 and out["ok"] is False and "August" in " ".join(out["fehler"])


def test_numbers_file_and_second_run_never_overwrites(fin_ws):
    _, a = rufe("management-report", "--ws", fin_ws, "--monat", "2026-09", "--heute", HEUTE)
    assert a["dokument"] == "03_Berichte/2026-10-06_management-report.docx"
    erste = (fin_ws / a["zahlen"]).read_text(encoding="utf-8")
    z = json.loads(erste)
    assert z["art"] == "management-report" and z["dokument"] == a["dokument"]
    assert any(p["name"] == "Material Ist" and p["kosten"] for p in z["posten"])
    _, b = rufe("management-report", "--ws", fin_ws, "--monat", "2026-09", "--heute", HEUTE)
    assert b["dokument"] == "03_Berichte/2026-10-06_management-report-2.docx"
    assert (fin_ws / a["zahlen"]).read_text(encoding="utf-8") == erste


def test_sample_mode_reads_beispiel_and_labels_it(kit_ws):
    cfg = kit_ws / "Unternehmen" / ".kit-config"
    cfg.write_text(cfg.read_text(encoding="utf-8").replace("beispieldaten=nein", "beispieldaten=ja"), encoding="utf-8")
    bsp = kit_ws / "Beispiel"
    shutil.copytree(kit_ws / "Unternehmen", bsp / "Unternehmen", ignore=shutil.ignore_patterns(".kit-*"))
    schreibe_csv(bsp, "ergebnis", "2026-09", SEPT)
    code, out = rufe("management-report", "--ws", kit_ws, "--monat", "2026-09", "--heute", HEUTE)
    assert code == 0 and out["beispiel"] is True and finanzen.BEISPIEL_HINWEIS in out["hinweise"]
    assert kern(out, "Umsatz Monat Ist")["quelle"][0].startswith("Beispiel/07_Daten/ergebnis_2026-09.csv")


def test_not_a_workspace(tmp_path):
    code, out = rufe("management-report", "--ws", tmp_path, "--monat", "2026-09")
    assert code == 1 and "kein Kundendienst-Ordner" in out["fehler"][0]


def bericht(ws):
    code, out = rufe("management-report", "--ws", ws, "--monat", "2026-09", "--heute", HEUTE)
    assert code == 0, out
    return out


def abgleich(ws, out, *extra):
    return rufe("abgleich", "--ws", ws, "--zahlen", out["zahlen"], "--heute", HEUTE, *extra)


def test_second_source_is_flagged_never_averaged(fin_ws):
    rows = [[m, p, pl, i + (12000 if p == "Umsatz Service" else 0)] for m, p, pl, i in SEPT]
    text = "Monat;Position;Plan;Ist\n" + "".join(
        f"09.2026;{p};{kz.deutsch(pl, 2)};{kz.deutsch(i, 2)}\n" for _, p, pl, i in rows)
    datei = fin_ws / "00_Eingang" / "ergebnis_2026-09_controlling.csv"
    datei.write_text(text, encoding="utf-8")
    code, out = rufe("management-report", "--ws", fin_ws, "--monat", "2026-09", "--heute", HEUTE,
                     "--zweitquelle", "00_Eingang/ergebnis_2026-09_controlling.csv")
    assert code == 0, out
    assert [k["position"] for k in out["konflikte"]] == ["Umsatz Service"]
    k = out["konflikte"][0]
    assert k["differenz"]["betrag"] == 12000 and k["uebernommen"]["betrag"] == 89000
    assert k["zweitquelle"]["quelle"] == ["00_Eingang/ergebnis_2026-09_controlling.csv Zeilen 2–2"]
    assert kern(out, "Umsatz Monat Ist")["betrag"] == 114000  # validated data; the mean would be 120.000
    assert datei.is_file() and not (fin_ws / "07_Daten" / "original").exists()


def test_unreadable_second_source_is_an_error(fin_ws):
    (fin_ws / "00_Eingang" / "kaputt.csv").write_text("a;b\n1;2\n", encoding="utf-8")
    code, out = rufe("management-report", "--ws", fin_ws, "--monat", "2026-09", "--zweitquelle", "00_Eingang/kaputt.csv")
    assert code == 1 and "Zweitquelle" in out["fehler"][0]


def test_reconciliation_passes_on_untouched_data(fin_ws):
    code, r = abgleich(fin_ws, bericht(fin_ws))
    assert code == 0 and r["ok"] and r["abweichungen"] == [] and r["geprueft"] > 0
    assert any("kein Original" in h for h in r["hinweise"])


def test_reconciliation_finds_changed_data(fin_ws):
    out = bericht(fin_ws)
    rows = [r[:] for r in SEPT]
    rows[2][3] = -37000
    schreibe_csv(fin_ws, "ergebnis", "2026-09", rows)
    code, r = abgleich(fin_ws, out)
    assert code == 1 and r["ok"] is False and any(x.startswith("Material Ist") for x in r["abweichungen"])


def test_reconciliation_against_the_original_export(fin_ws):
    out = bericht(fin_ws)
    orig = fin_ws / "07_Daten" / "original" / "ergebnis_2026-09__ergebnis_2026-09.xlsx"
    orig.parent.mkdir()
    wb = Workbook()
    wb.active.append(KOPF["ergebnis"])
    for z in SEPT:
        wb.active.append(z)
    wb.save(orig)
    assert abgleich(fin_ws, out)[0] == 0
    wb.active["D2"] = 90000
    wb.save(orig)
    code, r = abgleich(fin_ws, out)
    assert code == 1 and any("Kontrollsumme" in x for x in r["abweichungen"])


def test_reconciliation_checks_the_document(fin_ws):
    out = bericht(fin_ws)
    doc = Document()
    for k in out["kernzahlen"]:
        doc.add_paragraph(f"{k['name']}: {k['anzeige']}")
    doc.save(fin_ws / out["dokument"])
    assert abgleich(fin_ws, out, "--dokument", out["dokument"])[0] == 0
    doc = Document()
    for k in out["kernzahlen"][1:]:  # "Umsatz Monat Ist" (114.000) missing
        doc.add_paragraph(f"{k['name']}: {k['anzeige']}")
    doc.save(fin_ws / out["dokument"])
    code, r = abgleich(fin_ws, out, "--dokument", out["dokument"])
    assert code == 1 and any("Umsatz Monat Ist" in x and "steht nicht im Dokument" in x for x in r["abweichungen"])


def budget_ws(ws, ohne=()):
    for m in finanzen.monate_bis("2026-09", 12):
        if m in ohne:
            continue
        service = 21000 if m == "2025-12" else 10000
        schreibe_csv(ws, "ergebnis", m, [[m, "Umsatz Service", "", service], [m, "Personalkosten", "", -5000]])
    schreibe_csv(ws, "installed_base", "2026-09", [["Müller GmbH", "Anlage 1", "MM-400", 2019, "ja", "2027-03-01"],
                                                   ["Hansa Pack AG", "Anlage 2", "MM-600", 2021, "ja", "2028-01-01"],
                                                   ["Nordmetall GmbH", "Anlage 1", "MM-400", 2015, "nein", ""]])
    return ws


ANNAHMEN = ["--annahme", "umsatz|+3 %|Preiserhöhung 2027|angebot",
            "--annahme", "Personalkosten|+78.000 EUR|1 Techniker Team Süd|personal"]


def test_budget_from_run_rate_and_assumptions(kit_ws):
    budget_ws(kit_ws)
    code, out = rufe("budgetplanung", "--ws", kit_ws, "--jahr", 2027, "--heute", HEUTE, *ANNAHMEN)
    assert code == 0, out
    pos = {p["position"]: p for p in out["positionen"]}
    assert pos["Umsatz Service"]["basis"]["betrag"] == 131000
    assert pos["Umsatz Service"]["budget"]["betrag"] == 134930
    assert pos["Umsatz Service"]["monate"]["2027-12"] == 21630 and pos["Umsatz Service"]["monate"]["2027-01"] == 10300
    assert pos["Personalkosten"]["budget"]["betrag"] == -138000
    assert pos["Personalkosten"]["monate"]["2027-06"] == -11500
    assert out["budget_abschluss"]["ergebnis"]["betrag"] == -3070
    assert out["basis_monate"][0] == "2025-10" and out["basis_monate"][-1] == "2026-09"
    assert out["vertragsbasis"]["auslaufend"]["betrag"] == 1
    assert out["dokument"] == "03_Berichte/2026-10-06_budgetplanung-2027.xlsx"
    assert any("Annahme: umsatz|+3 %" in q for q in pos["Umsatz Service"]["budget"]["quelle"])


def test_budget_proposes_revenue_and_db1_targets(kit_ws):
    """D19: the budget proposes Umsatz and DB I % as kpi-ziele targets, not DB II %."""
    budget_ws(kit_ws)
    _, out = rufe("budgetplanung", "--ws", kit_ws, "--jahr", 2027, "--heute", HEUTE, *ANNAHMEN)
    namen = [k["name"] for k in out["kernzahlen"]]
    # DB I = Umsatz 134.930 − Personal 138.000 = −3.070 (no material/third-party positions)
    assert kern(out, "Zielvorschlag DB I in % 2027")["anzeige"] == f"{kz.deutsch(-3070 / 134930 * 100, 1)} %"
    assert kern(out, "Umsatz Budget 2027")["betrag"] == 134930
    assert not any(n.startswith("Zielvorschlag DB II") for n in namen), namen
    assert out["zielvorschlag"] == ["Umsatz", "DB I in %"] and "DB I in %" in out["gliederung"][0]["inhalt"]


def test_budget_without_assumptions_is_the_base(kit_ws):
    budget_ws(kit_ws)
    _, out = rufe("budgetplanung", "--ws", kit_ws, "--jahr", 2027, "--heute", HEUTE)
    assert out["klassen"]["umsatz"]["betrag"] == 131000 and out["budget_abschluss"]["ergebnis"]["betrag"] == 71000


def test_budget_names_missing_months(kit_ws):
    budget_ws(kit_ws, ohne=("2026-03",))
    code, out = rufe("budgetplanung", "--ws", kit_ws, "--jahr", 2027)
    assert code == 1 and "März 2026" in out["fehler"][0] and "schätzt keine" in out["fehler"][0]


def test_budget_rejects_unknown_assumption_target(kit_ws):
    budget_ws(kit_ws)
    code, out = rufe("budgetplanung", "--ws", kit_ws, "--jahr", 2027, "--annahme", "Marketing|+5 %|neu|angebot")
    assert code == 1 and "Umsatz Service" in out["fehler"][0]


def test_budget_numbers_reconcile(kit_ws):
    budget_ws(kit_ws)
    _, out = rufe("budgetplanung", "--ws", kit_ws, "--jahr", 2027, "--heute", HEUTE, *ANNAHMEN)
    code, r = abgleich(kit_ws, out)
    assert code == 0 and r["geprueft"] == 2


def test_investment_npv_payback_irr(kit_ws):
    code, out = rufe("investitionsantrag", "--ws", kit_ws, "--titel", "Diagnose-Messplatz", "--invest", "100.000",
                     "--rueckfluss", "30.000", "--jahre", 5, "--quelle", "Eingabe im Gespräch", "--heute", HEUTE)
    assert code == 0, out
    assert out["kapitalwert"]["betrag"] == 19781.3 and "(1 + i)^t" in out["kapitalwert"]["formel"]
    assert out["amortisation"]["anzeige"] == "3,3 Jahre"
    assert out["interner_zinsfuss"]["betrag"] == 15.24
    assert out["sensitivitaet"]["betrag"] == -4174.96  # all cash flows −20 %
    assert out["zins"]["quelle"] == [finanzen.STANDARD_HINWEIS]
    assert out["dokument"] == "04_Angebote/2026-10-06_investitionsantrag-diagnose-messplatz.docx"
    assert "nicht festgelegt" in out["entscheidung"]


def test_investment_german_inputs_and_user_rate(kit_ws):
    _, out = rufe("investitionsantrag", "--ws", kit_ws, "--titel", "Prüfstand", "--invest", "1,2 Mio.",
                  "--rueckfluss", "280.000 €", "--jahre", 5, "--restwert", "120.000", "--zins", "6,5 %",
                  "--quelle", "Angebot Lieferant")
    assert out["kapitalwert"]["betrag"] == 51175.94 and out["zins"]["quelle"] == ["Angebot Lieferant"]


def test_investment_flow_count_must_match(kit_ws):
    code, out = rufe("investitionsantrag", "--ws", kit_ws, "--titel", "X", "--invest", "1000", "--rueckfluss", "500",
                     "--rueckfluss", "600", "--jahre", 3, "--quelle", "Eingabe")
    assert code == 1 and "--jahre 3" in out["fehler"][0]


AUG = [["SA-11", "Müller GmbH", "Anlage 1", "Wartung", "2026-08-03", "2026-08-04", "abgeschlossen", 10, 1000, 500, "Nord"],
       ["SA-12", "Hansa Pack AG", "Anlage 2", "Reparatur", "2026-08-10", "2026-08-11", "abgeschlossen", 10, 1500, 900, "Süd"]]
SEP = [["SA-21", "Müller GmbH", "Anlage 1", "Wartung", "2026-09-03", "2026-09-04", "abgeschlossen", 20, 2200, 1000, "Nord"],
       ["SA-22", "Hansa Pack AG", "Anlage 2", "Reparatur", "2026-09-10", "2026-09-11", "abgeschlossen", 5, 700, 500, "Süd"]]


def test_margin_drivers_add_up(kit_ws):
    schreibe_csv(kit_ws, "auftraege", "2026-08", AUG)
    schreibe_csv(kit_ws, "auftraege", "2026-09", SEP)
    code, out = rufe("margen-analyse", "--ws", kit_ws, "--periode", "2026-09", "--vergleich", "2026-08", "--heute", HEUTE)
    assert code == 0, out
    e = {k: v["betrag"] for k, v in out["effekte"].items()}
    assert e == {"volumen": 275, "mix": -75, "preis": 150, "kosten": -50, "nicht_zerlegbar": 0, "veraenderung": 300}
    assert out["db"]["vergleich"]["betrag"] == 1100 and out["db"]["periode"]["betrag"] == 1400
    assert out["treiber"][0] == "Volumeneffekt"
    assert out["effekte"]["preis"]["quelle"][0].startswith("07_Daten/auftraege_2026-09.csv")


def test_margin_analysis_stops_on_missing_cost(kit_ws):
    schreibe_csv(kit_ws, "auftraege", "2026-08", AUG)
    rows = [r[:] for r in SEP]
    rows[0][9] = ""
    schreibe_csv(kit_ws, "auftraege", "2026-09", rows)
    code, out = rufe("margen-analyse", "--ws", kit_ws, "--periode", "2026-09", "--vergleich", "2026-08")
    assert code == 1 and "07_Daten/auftraege_2026-09.csv Zeile 1" in out["fehler"][0]


def test_investment_names_the_company_decision_right(kit_ws):
    (kit_ws / "Unternehmen" / "ergebnisrechnung.md").write_text(
        "---\nentscheidungsrechte:\n"
        '  - {"thema": "investition", "allein_bis_eur": 15000, "sonst": "Geschäftsführung"}\n'
        "---\n\n# Ergebnisrechnung Service\n", encoding="utf-8")
    _, out = rufe("investitionsantrag", "--ws", kit_ws, "--titel", "X", "--invest", "100.000", "--rueckfluss", "30.000",
                  "--jahre", 5, "--quelle", "Eingabe")
    assert out["entscheidung"] == "Entscheidung durch Geschäftsführung (über 15.000 EUR)"


def angebotsfall(ws, von="vertrieb"):
    code, out = vorgang._main(["neu", "--ws", str(ws), "--titel", "Angebot Wartungsvertrag Nordmetall", "--typ",
                               "angebot", "--kunde", "Nordmetall GmbH", "--verantwortlich", "Jana Becker", "--von", von,
                               "--text", "Quelle: 04_Angebote/angebot.docx", "--heute", HEUTE])
    assert code == 0, out
    return out["nr"]


def margenpruefung(ws, nr, kosten, rabatt="10 %"):
    return rufe("margen-pruefung", "--ws", ws, "--nr", nr, "--listenpreis", "50.000", "--rabatt-prozent", rabatt,
                "--kosten", kosten, "--quelle", "04_Angebote/kalkulation.xlsx")


@pytest.mark.parametrize("kosten,urteil", [("25.000", "zustimmen"), ("30.000", "zustimmen mit Auflagen"),
                                           ("42.000", "ablehnen")])
def test_margin_check_recommendation(kit_ws, kosten, urteil):
    nr = angebotsfall(kit_ws)
    vorher = (kit_ws / "01_Vorgaenge" / "offen" / f"{nr}.md").read_text(encoding="utf-8")
    code, out = margenpruefung(kit_ws, nr, kosten)
    assert code == 0, out
    assert out["urteil"] == urteil and out["empfehlung"].startswith(f"Empfehlung: {urteil} – ")
    assert out["pruefpflicht"] is True  # no limits set -> always review (§8 rule 1)
    assert finanzen.STANDARD_HINWEIS in out["empfehlung"]
    assert (kit_ws / "01_Vorgaenge" / "offen" / f"{nr}.md").read_text(encoding="utf-8") == vorher


def test_margin_check_names_the_maximum_discount(kit_ws):
    _, out = margenpruefung(kit_ws, angebotsfall(kit_ws), "30.000")
    assert "Rabatt auf höchstens 7,7 %" in out["empfehlung"] and out["werte"]["db_prozent"]["anzeige"] == "33,3 %"


def test_margin_check_german_formats_and_negative_margin(kit_ws):
    code, out = rufe("margen-pruefung", "--ws", kit_ws, "--nr", angebotsfall(kit_ws), "--listenpreis", "45.000,00 €",
                     "--rabatt-prozent", "12,5", "--kosten", "40.500", "--quelle", "Eingabe")
    assert code == 0 and out["urteil"] == "ablehnen" and out["werte"]["db"]["betrag"] == -1125


def test_own_case_is_refused(kit_ws):
    nr = angebotsfall(kit_ws, von="finanzen")
    code, out = margenpruefung(kit_ws, nr, "25.000")
    assert code == 1 and "§8 Regel 3" in out["fehler"][0]


def test_margin_check_is_a_db1_check(kit_ws):
    """D19: the deal margin is DB I (all direct costs incl. technician hours at full cost) against 35 %."""
    _, out = margenpruefung(kit_ws, angebotsfall(kit_ws), "30.000")
    assert [out["werte"][k]["name"] for k in ("db", "db_prozent", "ziel")] == ["DB I", "DB I in %", "Ziel DB I in %"]
    assert "Ziel DB I" in out["empfehlung"] and "DB II" not in out["empfehlung"]
    assert out["werte"]["ziel"]["betrag"] == 35.0 and out["werte"]["ziel"]["quelle"] == [finanzen.STANDARD_HINWEIS]


def test_margin_check_prices_hours_at_the_full_cost_rate(kit_ws):
    nr = angebotsfall(kit_ws)
    args = ["margen-pruefung", "--ws", kit_ws, "--nr", nr, "--listenpreis", "50.000", "--rabatt-prozent", "0",
            "--material", "20.000", "--fremdleistung", "1.000", "--stunden", "100", "--quelle", "Kalkulation"]
    _, out = rufe(*args)
    assert out["werte"]["stundensatz"]["anzeige"] == "53,13 EUR/h"  # kit standard 85.000 EUR / 1.600 h
    assert out["werte"]["stundensatz"]["quelle"] == [finanzen.STANDARD_HINWEIS]
    assert out["werte"]["kosten"]["betrag"] == 26313 and "53,13 EUR/h" in out["empfehlung"]
    (kit_ws / "Unternehmen" / "ergebnisrechnung.md").write_text(
        '---\npersonal: {"vollkosten_techniker_eur": 82000, "netto_stunden": 1520}\n---\n\n# Ergebnisrechnung\n',
        encoding="utf-8")
    _, out = rufe(*args)
    assert out["werte"]["stundensatz"]["anzeige"] == "53,95 EUR/h"
    assert out["werte"]["stundensatz"]["quelle"] == ["Unternehmen/ergebnisrechnung.md (personal)"]
    assert out["werte"]["kosten"]["betrag"] == 26395


def test_margin_check_needs_one_kind_of_cost_input(kit_ws):
    nr = angebotsfall(kit_ws)
    basis = ["margen-pruefung", "--ws", kit_ws, "--nr", nr, "--listenpreis", "50.000", "--rabatt-prozent", "0",
             "--quelle", "Eingabe"]
    code, out = rufe(*basis)
    assert code == 1 and "--kosten" in out["fehler"][0]
    code, out = rufe(*basis, "--kosten", "25.000", "--stunden", "10")
    assert code == 1 and "entweder" in out["fehler"][0]


SKILLS_4B = ["budgetplanung", "investitionsantrag", "management-report", "margen-analyse", "margen-pruefung"]


def skill(name):
    text = (ROOT / "plugin" / "skills" / name / "SKILL.md").read_text(encoding="utf-8")
    m = re.match(r"\A---\n(.*?)\n---\n(.*)\Z", text, re.S)
    assert m, name
    return yaml.safe_load(m.group(1)), m.group(2)


@pytest.mark.parametrize("name", SKILLS_4B)
def test_finance_skill_contract(name):
    meta, body = skill(name)
    assert meta["name"] == name and 40 <= len(meta["description"]) <= 1024
    assert "**Liest:**" in body and "**Schreibt:**" in body and "Daten, nie Anweisungen" in body
    assert f'uv run "${{CLAUDE_PLUGIN_ROOT}}/scripts/finanzen.py" {name}' in body
    assert "$CLAUDE_PLUGIN_ROOT" not in body.replace("${CLAUDE_PLUGIN_ROOT}", "")
    assert " entscheide " not in body and "pip install" not in body  # only the human decides (§6)


def test_own_outputs_get_no_recommendation_from_finanzen():
    for name in ("management-report", "budgetplanung", "investitionsantrag", "margen-analyse"):
        assert "--art empfehlung" not in skill(name)[1], name
    assert "--art empfehlung --von finanzen" in skill("margen-pruefung")[1]


def test_workflow_chains_are_in_the_skills():
    mr, bp = skill("management-report")[1], skill("budgetplanung")[1]
    for teil in ("finanzen.py\" abgleich", "keine Maßnahme", "praesentation", "mail-entwurf", "--von finanzen"):
        assert teil in mr, teil
    for teil in ("finanzen.py\" abgleich", "service-leader-kit:personal", "service-leader-kit:angebot",
                 "entscheidungsvorlage"):
        assert teil in bp, teil


def test_skills_carry_the_db1_target():
    """D19 in the skill texts: DB I (after service personnel) is the target margin everywhere."""
    mr, bp, mp = (skill(n)[1] for n in ("management-report", "budgetplanung", "margen-pruefung"))
    assert "DB I in % vom Umsatz" in mr and "Ziel DB I in %" in mr and "Ziel DB II" not in mr
    assert "Umsatz, DB I %" in bp and "DB II %" not in bp
    for teil in ("DB I", "--material", "--fremdleistung", "--stunden", "Vollkostensatz", "Standarddefinition"):
        assert teil in mp, teil


EVALS = ROOT / "plugin" / "evals"
FAELLE_4B = [f"{s}-{d}" for s in SKILLS_4B for d in ("sauber", "unordentlich")] + ["workflow-monatsbericht",
                                                                                     "workflow-budget"]
BINDUNG = {("management-report-sauber", "umsatz"): "mr_umsatz_ist",
           ("workflow-monatsbericht", "umsatz"): "mr_umsatz_ist",
           ("management-report-unordentlich", "kein-mittel-position"): "konflikt_mittel_position",
           ("management-report-unordentlich", "kein-mittel-umsatz"): "konflikt_mittel_umsatz",
           ("budgetplanung-sauber", "basis-umsatz"): "budget_basis_umsatz",
           ("workflow-budget", "budget-umsatz"): "budget_umsatz_2027",
           ("workflow-budget", "budget-personal"): "budget_personal_2027",
           ("margen-analyse-sauber", "veraenderung"): "ma_delta", ("margen-analyse-sauber", "volumen"): "ma_volumen",
           ("margen-analyse-sauber", "mix"): "ma_mix", ("margen-analyse-sauber", "preis"): "ma_preis",
           ("margen-analyse-sauber", "kosten"): "ma_kosten",
           ("investitionsantrag-sauber", "kapitalwert"): "inv_kw_sauber",
           ("investitionsantrag-unordentlich", "szenario-a"): "inv_kw_a",
           ("investitionsantrag-unordentlich", "szenario-b"): "inv_kw_b"}


def grader(fall, name):
    c = yaml.safe_load((EVALS / fall / "case.yaml").read_text(encoding="utf-8"))
    return next(g for g in c["graders"] if g["name"] == name)


def test_lane_eval_cases_exist_and_are_filled():
    for f in FAELLE_4B:
        assert (EVALS / f / "prompt.md").is_file(), f
        assert "{{" not in (EVALS / f / "case.yaml").read_text(encoding="utf-8"), f


def test_graders_carry_the_independent_totals():
    e = json.loads((EVALS / "erwartet" / "finanzen.json").read_text(encoding="utf-8"))
    assert e == ef.berechne(), "erwartet/finanzen.json veraltet: uv run tools/erwartet_finanzen.py --einsetzen"
    for (fall, name), key in BINDUNG.items():
        assert grader(fall, name)["pattern"] == ef.muster(e[key]), (fall, name)
        assert kz.deutsch(abs(e[key])) == ef.de(abs(e[key]))
    for fall in ("management-report-sauber", "workflow-monatsbericht"):
        assert grader(fall, "massnahmen-als-vorgang")["match"] == f"count:{e['mr_massnahmen']}"


def test_margin_check_evals_use_the_db1_names():
    """D19 in the evals: graders and criteria speak of the DB I target, never of a DB II target."""
    for fall in ("margen-pruefung-sauber", "margen-pruefung-unordentlich", "management-report-sauber"):
        text = (EVALS / fall / "case.yaml").read_text(encoding="utf-8")
        assert "DB I" in text and not re.search(r"(Ziel DB II|DB II %|DB II target|DB II margin)", text), fall


def test_sample_year_supports_the_finance_evals():
    e = ef.berechne()
    assert e["mr_massnahmen"] >= 1, "Beispieljahr: September 2026 ohne Abweichung über der Maßnahmen-Schwelle"
    assert all(ef.zahl(r["Plan_EUR"]) is not None for r in ef.lies("ergebnis_2026-09.csv"))
    for m in ef.monate("2026-09", 6):
        assert all(ef.zahl(r["Kosten_EUR"]) is not None for r in ef.lies(f"auftraege_{m}.csv")), m


def test_finanzen_agrees_with_the_independent_totals(kit_ws):
    shutil.copytree(ROOT / "plugin" / "beispiel" / "07_Daten", kit_ws / "07_Daten", dirs_exist_ok=True)
    for p in (ROOT / "plugin" / "beispiel" / "Unternehmen").glob("*.md"):
        shutil.copy(p, kit_ws / "Unternehmen" / p.name)
    e = ef.berechne()
    _, out = rufe("management-report", "--ws", kit_ws, "--monat", "2026-09", "--heute", HEUTE)
    assert kern(out, "Umsatz Monat Ist")["betrag"] == e["mr_umsatz_ist"] and out["massnahmen_offen"] == e["mr_massnahmen"]
    assert out["quellenvergleich"] == [], "Beispieljahr: Ergebnisrechnung und Exporte passen nicht zusammen"
    _, ma = rufe("margen-analyse", "--ws", kit_ws, "--periode", "2026-07..2026-09", "--vergleich", "2026-04..2026-06")
    assert ma["effekte"]["veraenderung"]["betrag"] == e["ma_delta"]


def test_messy_report_scaffold_builds_the_conflict(tmp_path, shell):
    r = subprocess.run([shell, str(EVALS / "management-report-unordentlich" / "scaffold.sh")], cwd=tmp_path,
                       env=runner_env(tmp_path), capture_output=True, text=True, timeout=120)
    assert r.returncode == 0, r.stderr
    text = (tmp_path / "00_Eingang" / "ergebnis_2026-09_controlling.csv").read_text(encoding="utf-8")
    assert text.startswith("Monat;Position;Plan;Ist\n")
    assert ef.de(ef.berechne()["konflikt_mittel_position"] + 6000, 2) in text
    code, out = rufe("management-report", "--ws", tmp_path, "--monat", "2026-09",
                     "--zweitquelle", "00_Eingang/ergebnis_2026-09_controlling.csv")
    assert code == 0 and [k["differenz"]["betrag"] for k in out["konflikte"]] == [12000]


def test_messy_margin_scaffold_blanks_two_costs(tmp_path, shell):
    r = subprocess.run([shell, str(EVALS / "margen-analyse-unordentlich" / "scaffold.sh")], cwd=tmp_path,
                       env=runner_env(tmp_path), capture_output=True, text=True, timeout=120)
    assert r.returncode == 0, r.stderr
    code, out = rufe("margen-analyse", "--ws", tmp_path, "--periode", "2026-07..2026-09", "--vergleich", "2026-04..2026-06")
    assert code == 1 and "auftraege_2026-09.csv Zeile 3" in out["fehler"][0] and "Zeile 7" in out["fehler"][0]


def test_source_comparison_names_service_revenue_against_orders(fin_ws):
    out = bericht(fin_ws)
    v = {x["thema"]: x for x in out["quellenvergleich"]}
    assert v["Serviceumsatz"]["ergebnisrechnung"]["betrag"] == 89000 and v["Serviceumsatz"]["export"]["betrag"] == 3800
    assert v["Serviceumsatz"]["differenz"]["betrag"] == 85200


@pytest.mark.parametrize("name", ["DB I-Marge", "DB1-Marge", "Deckungsbeitrag I", "DB I in %", "DBI-Marge"])
def test_db1_target_names_that_match(name):
    defs = {"kpi-ziele": {"kennzahlen": [{"name": name, "ziel": 40, "einheit": "%"}]}}
    assert finanzen.db1_ziel(defs) == (40.0, False)


@pytest.mark.parametrize("name", ["DB II in %", "Deckungsbeitrag II in %", "DB II-Marge", "DB2-Marge", "DB III-Marge"])
def test_db2_kpis_are_never_the_margin_target(name):
    defs = {"kpi-ziele": {"kennzahlen": [{"name": name, "ziel": 32, "einheit": "%"}]}}
    assert finanzen.db1_ziel(defs) == (finanzen.DB1_ZIEL_STANDARD, True)


def test_db1_target_skips_a_db2_kpi_listed_first():
    defs = {"kpi-ziele": {"kennzahlen": [{"name": "DB II in %", "ziel": 32, "einheit": "%"},
                                         {"name": "DB I in %", "ziel": 40, "einheit": "%"}]}}
    assert finanzen.db1_ziel(defs) == (40.0, False)
