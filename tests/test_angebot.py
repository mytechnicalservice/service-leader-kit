import json
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

import angebot
from angebot_fixtures import PREISLISTE_SAUBER, PREISLISTE_UNORDENTLICH, schreibe_preisliste
from conftest import ROOT

PLUGIN = ROOT / "plugin"
FIX = PLUGIN / "evals" / "_angebot"
NEU_SAUBER = {"SL-100": 123.0, "SL-101": 145.0, "SL-102": 100.0, "SL-200": 91.5, "SL-201": 153.0,
              "SL-300": 3010.0, "SL-301": 4990.0, "SL-302": 8210.0, "SL-400": 1300.0, "SL-500": 424.0}


def cli(*args):
    return angebot._main([str(a) for a in args])


def preis_ws(kit_ws, zeilen=PREISLISTE_SAUBER, name="preisliste_2026.xlsx"):
    schreibe_preisliste(kit_ws / "04_Angebote" / name, zeilen)
    return kit_ws


def blatt(pfad, name="Preisliste"):
    wb = load_workbook(pfad, data_only=True)
    return [list(r) for r in wb[name].iter_rows(values_only=True)]


@pytest.mark.parametrize("roh,neu", [("99.84", "100.00"), ("91.67", "91.50"), ("122.72", "123"),
                                     ("3013.10", "3010"), ("1296.875", "1300"), ("153.47", "153")])
def test_rounding_steps(roh, neu):
    from decimal import Decimal
    assert angebot.runden(Decimal(roh)) == Decimal(neu)


def test_clean_fixture_gives_the_graded_prices(kit_ws):
    code, out = cli("preisliste", "--ws", kit_ws, "--jahr", 2027, "--lohn", "4,5", "--material", "3",
                    "--allgemein", "2", "--datei", FIX / "preisliste_2026_sauber.xlsx")
    assert code == 0 and out["ok"], out
    assert {r["artikelnr"]: r["neu"] for r in out["zeilen"]} == NEU_SAUBER
    pct = {k["name"]: k["betrag"] for k in out["kategorien"]}
    assert pct == {"Preisänderung Ersatzteil": 3.0, "Preisänderung Pauschale": 3.0, "Preisänderung Schulung": 3.75,
                   "Preisänderung Stundensatz": 4.0, "Preisänderung Vertrag": 3.9}
    assert all(k["berechnet"] and k["formel"] for k in out["kategorien"])
    assert out["meldungen"] == [] and out["warnungen"] == []
    assert out["ziel"] == "04_Angebote/preisliste_2027.xlsx" and not (kit_ws / out["ziel"]).exists()
    assert {b["kategorie"]: b["aenderung_prozent"] for b in out["budget_annahmen"]}["Vertrag"] == 3.9


def test_writes_the_new_list_with_the_same_schema(kit_ws):
    preis_ws(kit_ws)
    code, out = cli("preisliste", "--ws", kit_ws, "--jahr", 2027, "--lohn", "4.5", "--material", "3",
                    "--allgemein", "2", "--schreiben", "--heute", "2026-10-06")
    assert code == 0 and out["geschrieben"] is True
    rows = blatt(kit_ws / "04_Angebote" / "preisliste_2027.xlsx")
    assert rows[0] == ["Artikelnr", "Bezeichnung", "Kategorie", "Einheit", "Preis_EUR", "Gueltig_ab", "Bemerkung"]
    assert {r[0]: r[4] for r in rows[1:]} == NEU_SAUBER
    assert rows[1][5].date().isoformat() == "2027-01-01"
    assert blatt(kit_ws / "04_Angebote" / "preisliste_2027.xlsx", "Änderungen")[1][:6] == \
        ["SL-100", "Stundensatz Servicetechniker", "Stundensatz", 118, 123, 4]
    assert out["memo_ziel"] == "04_Angebote/2026-10-06_preisliste-update.docx"


def test_second_run_writes_v2_and_keeps_the_first(kit_ws):
    preis_ws(kit_ws)
    args = ["preisliste", "--ws", kit_ws, "--jahr", 2027, "--lohn", "4,5", "--material", "3", "--allgemein", "2",
            "--schreiben"]
    cli(*args)
    erst = (kit_ws / "04_Angebote" / "preisliste_2027.xlsx").read_bytes()
    code, out = cli(*args[:-1], "--lohn", "6", "--schreiben")
    assert code == 0 and out["ziel"] == "04_Angebote/preisliste_2027_v2.xlsx"
    assert (kit_ws / "04_Angebote" / "preisliste_2027.xlsx").read_bytes() == erst


def test_dry_run_writes_nothing(kit_ws):
    preis_ws(kit_ws)
    vorher = sorted(p.name for p in (kit_ws / "04_Angebote").iterdir())
    code, out = cli("preisliste", "--ws", kit_ws, "--jahr", 2027, "--lohn", "4", "--material", "3", "--allgemein", "2")
    assert code == 0 and out["geschrieben"] is False
    assert sorted(p.name for p in (kit_ws / "04_Angebote").iterdir()) == vorher


def test_messy_price_list_is_flagged_not_guessed(kit_ws):
    preis_ws(kit_ws, PREISLISTE_UNORDENTLICH)
    code, out = cli("preisliste", "--ws", kit_ws, "--jahr", 2027, "--lohn", "5,5", "--material", "3",
                    "--allgemein", "2,2", "--markt", "Stundensatz=-2", "--markt", "Vertrag=6", "--schreiben")
    assert code == 0, out
    neu = {(r["zeile"], r["artikelnr"]): (r["status"], r["neu"]) for r in out["zeilen"]}
    assert neu[(2, "SL-100")] == ("angepasst", 121.0) and neu[(3, "SL-101")] == ("angepasst", 143.0)
    assert neu[(4, "SL-102")] == ("unverändert", None)
    assert neu[(5, "SL-200")][0] == neu[(6, "SL-200")][0] == "unverändert"
    assert neu[(7, "SL-300")] == ("angepasst", 3210.0) and neu[(8, "SL-301")] == ("angepasst", 5310.0)
    text = " ".join(out["meldungen"])
    assert "SL-102" in text and "Preis fehlt" in text
    assert "SL-200" in text and "mehrfach (Zeilen 5, 6)" in text
    assert "Zeile 8 (SL-301): Bemerkung enthält eine Anweisung" in text
    assert any("Vertrag" in w and "10,63" in w and "Preiselastizität" in w for w in out["warnungen"])
    rows = {(r[0], r[4]) for r in blatt(kit_ws / out["ziel"])[1:]}
    assert ("SL-200", 89) in rows and ("SL-200", 92) in rows and ("SL-102", None) in rows


def test_formula_text_stays_text(kit_ws):
    preis_ws(kit_ws, PREISLISTE_SAUBER[:1])
    p = kit_ws / "04_Angebote" / "preisliste_2026.xlsx"
    wb = load_workbook(p)
    wb.active["G2"].value = '=HYPERLINK("http://x.example","klick")'
    wb.active["G2"].data_type = "s"  # text starting with '=', as Excel stores it after a leading apostrophe
    wb.save(p)
    code, out = cli("preisliste", "--ws", kit_ws, "--jahr", 2027, "--lohn", "4", "--material", "0",
                    "--allgemein", "0", "--schreiben")
    wb = load_workbook(kit_ws / out["ziel"])
    assert wb["Preisliste"]["G2"].value == "'=HYPERLINK(\"http://x.example\",\"klick\")"
    assert wb["Preisliste"]["G2"].data_type == "s"


@pytest.mark.parametrize("extra,text", [
    (["--lohn", "45"], "außerhalb von −20 … +30 %"),
    (["--lohn", "viel"], "keine Prozentzahl"),
    (["--markt", "Hotline=3"], "Kategorie unbekannt"),
])
def test_bad_inputs_are_refused(kit_ws, extra, text):
    preis_ws(kit_ws)
    code, out = cli("preisliste", "--ws", kit_ws, "--jahr", 2027, "--lohn", "4", "--material", "3",
                    "--allgemein", "2", *extra)  # argparse: the last --lohn wins
    assert code == 1 and not out["ok"] and text in out["fehler"][0]


def test_missing_column_and_missing_file(kit_ws):
    code, out = cli("preisliste", "--ws", kit_ws, "--jahr", 2027, "--lohn", "4", "--material", "3", "--allgemein", "2")
    assert code == 1 and "preisliste_2026.xlsx nicht gefunden" in out["fehler"][0]
    schreibe_preisliste(kit_ws / "04_Angebote" / "preisliste_2026.xlsx", [])
    wb = load_workbook(kit_ws / "04_Angebote" / "preisliste_2026.xlsx")
    wb.active.delete_cols(5)
    wb.save(kit_ws / "04_Angebote" / "preisliste_2026.xlsx")
    code, out = cli("preisliste", "--ws", kit_ws, "--jahr", 2027, "--lohn", "4", "--material", "3", "--allgemein", "2")
    assert code == 1 and "Spalte(n) Preis_EUR fehlen" in out["fehler"][0]


def test_not_a_workspace(tmp_path):
    code, out = cli("preisliste", "--ws", tmp_path, "--jahr", 2027, "--lohn", "4", "--material", "3", "--allgemein", "2")
    assert code == 1 and "kein Kundendienst-Ordner" in out["fehler"][0]


def test_freier_name_is_case_insensitive(kit_ws):
    (kit_ws / "03_Berichte" / "Bericht.DOCX").write_text("x", encoding="utf-8")
    assert angebot.freier_name(kit_ws, "03_Berichte", "bericht", ".docx") == "03_Berichte/bericht_v2.docx"


def test_committed_fixtures_match_the_lists():
    for name, zeilen in (("sauber", PREISLISTE_SAUBER), ("unordentlich", PREISLISTE_UNORDENTLICH)):
        rows = blatt(FIX / f"preisliste_2026_{name}.xlsx")[1:]
        assert [r[0] for r in rows] == [z[0] for z in zeilen]
        assert [r[4] for r in rows] == [z[4] for z in zeilen]


def schreibe_kpi_ziele(kit_ws, kennzahlen):
    zeilen = "\n".join(f"  - {json.dumps(k, ensure_ascii=False)}" for k in kennzahlen)
    (kit_ws / "Unternehmen" / "kpi-ziele.md").write_text(f"---\nkennzahlen:\n{zeilen}\n---\n# KPIs\n", encoding="utf-8")


def test_target_margin_is_the_db2_kpi_not_db1(kit_ws):
    """K1 (Max, 2026-10-07): target margin = the DB II target, never DB I or the first 'Marge'."""
    schreibe_kpi_ziele(kit_ws, [{"name": "DB I-Marge", "ziel": 70}, {"name": "DB II-Marge", "ziel": 30}])
    assert angebot.zielmarge(kit_ws) == (30.0, "Unternehmen/kpi-ziele.md: DB II-Marge")


@pytest.mark.parametrize("name,treffer", [("DB II-Marge", True), ("DB2-Marge", True), ("Deckungsbeitrag II", True),
                                          ("DB I-Marge", False), ("Marge", False), ("DB I Ist-Marge", False)])
def test_db2_name_matching(name, treffer):
    assert angebot.ist_db2(name) is treffer


def test_plain_margin_kpi_falls_back_to_the_standard(kit_ws):
    schreibe_kpi_ziele(kit_ws, [{"name": "Marge", "ziel": 50}, {"name": "DB I-Marge", "ziel": 70}])
    assert angebot.zielmarge(kit_ws) == (angebot.ZIELMARGE_STANDARD, angebot.STANDARD_HINWEIS)
