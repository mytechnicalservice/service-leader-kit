import datetime as dt
import json

import openpyxl
import pytest

import kennzahlen as kz
from conftest import ROOT
from generate_beispiel import MONATE, generate

PLUGIN = ROOT / "plugin"


@pytest.fixture(scope="module")
def gen(tmp_path_factory):
    root = tmp_path_factory.mktemp("gen")
    return generate(root), root


def cells(path):
    return [list(r) for r in openpyxl.load_workbook(path).active.iter_rows(values_only=True)]


def test_deterministic_and_identical_to_the_shipped_sample(gen, tmp_path):
    exp, root = gen
    assert generate(tmp_path) == exp
    assert json.loads((PLUGIN / "evals" / "erwartet" / "beispiel.json").read_text(encoding="utf-8")) == exp, \
        "plugin/ ist veraltet: uv run tools/generate_beispiel.py ausführen und committen"
    for rel in ("beispiel/07_Daten/auftraege_2026-09.csv", "beispiel/07_Daten/ergebnis_2025-10.csv"):
        assert (root / rel).read_bytes() == (PLUGIN / rel).read_bytes(), rel


def test_full_year_in_07_daten_and_september_in_the_inbox(gen):
    exp, root = gen
    daten = root / "beispiel" / "07_Daten"
    for vorlage in ("auftraege", "ersatzteile", "ergebnis", "kapazitaet"):
        assert sorted(p.stem[-7:] for p in daten.glob(f"{vorlage}_*.csv")) == MONATE
    assert (daten / "installed_base_2026-09.csv").is_file() and (daten / "budget_2026.xlsx").is_file()
    rows = cells(root / "beispiel" / "00_Eingang" / "auftraege_2026-09.xlsx")
    umsatz = round(sum(r[rows[0].index("Umsatz_EUR")] for r in rows[1:]), 2)
    assert (umsatz, len(rows) - 1) == (exp["auftraege_umsatz"], exp["auftraege_zeilen"])


def test_numbers_are_internally_consistent(gen):
    exp, root = gen
    ws = root / "beispiel"
    for m in MONATE:
        erg = {z["Position"]: z["Ist_EUR"] for z in kz.lade(ws, "ergebnis", [m])}
        auf = kz.lade(ws, "auftraege", [m])
        assert erg["Umsatz Service"] == pytest.approx(sum(z["Umsatz_EUR"] for z in auf))
        teile = kz.lade(ws, "ersatzteile", [m])
        assert erg["Umsatz Ersatzteile"] == pytest.approx(sum(z["Menge"] * z["Stueckpreis_EUR"] for z in teile))
        assert -erg["Material"] + 0.01 >= sum(z["Menge"] * z["Einstandspreis_EUR"] for z in teile)
        assert all(z["Kosten_EUR"] is not None for z in auf)
        assert erg["Umsatz Verträge"] == pytest.approx(exp["vertraege_jahreswert"] / 12)
        kap = kz.lade(ws, "kapazitaet", [m])
        assert sum(z["Stunden"] for z in auf) <= sum(z["Ist_Stunden"] for z in kap)
        assert exp[f"umsatz_gesamt_{m}"] == pytest.approx(sum(v for k, v in erg.items() if k.startswith("Umsatz ")))
    assert 4_000_000 < exp["umsatz_gesamt_jahr"] < 6_000_000
    basis = kz.lade(ws, "installed_base")
    assert exp["installed_base_anlagen"] == len(basis)
    m4 = next(z for z in basis if z["Kunde"] == "Müller GmbH" and z["Anlage"] == "Anlage 4")
    assert (m4["Maschinentyp"], m4["Baujahr"], m4["Vertrag"]) == ("MM-600", 2021.0, "nein")
    assert any(z["Kunde"] == "Nordmetall GmbH" and z["Maschinentyp"] == "MM-400" for z in basis)
    assert any(z["Kunde"] == "Müller GmbH" and z["Anlage"] == "Anlage 4" and z["Abschluss"] == "2026-09-03"
               for z in kz.lade(ws, "auftraege", ["2026-09"]))
    budget = openpyxl.load_workbook(ws / "07_Daten" / "budget_2026.xlsx")["Budget 2026"]
    kopf = [c.value for c in budget[1]]
    plan = {r[0]: r for r in budget.iter_rows(min_row=2, values_only=True)}
    for z in kz.lade(ws, "ergebnis", ["2026-09"]):
        assert z["Plan_EUR"] == plan[z["Position"]][kopf.index("2026-09")]


def test_margins_follow_d19(gen):
    """D19: DB I = Umsatz − Material − Fremdleistung − Personalkosten, DB II = DB I − Gewährleistung,
    Ergebnis = DB II − Gemeinkostenumlage – recomputed from the shipped CSVs, not from the generator."""
    exp, root = gen
    ws = root / "beispiel"
    for m in MONATE:
        erg = {z["Position"]: z["Ist_EUR"] for z in kz.lade(ws, "ergebnis", [m])}
        umsatz = sum(v for k, v in erg.items() if k.startswith("Umsatz "))
        db1 = umsatz + erg["Material"] + erg["Fremdleistung"] + erg["Personalkosten"]
        assert exp[f"db1_{m}"] == pytest.approx(db1, abs=0.01), m
        assert exp[f"db2_{m}"] == pytest.approx(db1 + erg["Gewährleistung"], abs=0.01), m
        assert exp[f"ergebnis_{m}"] == pytest.approx(exp[f"db2_{m}"] + erg["Gemeinkostenumlage"], abs=0.01), m
    assert (exp["db1_jahr"], exp["db2_jahr"]) == (1789388.9, 1694088.9)
    assert (exp["db1_2026-09"], exp["db2_2026-09"]) == (167721.51, 148821.51)
    unter = [m for m in MONATE if exp[f"db1_{m}"] / exp[f"umsatz_gesamt_{m}"] < 0.35]
    assert unter == ["2025-11", "2025-12", "2026-08"]
    d = kz.definitionen(ws)
    ziele = {k["name"]: k["ziel"] for k in d["kpi-ziele"]["kennzahlen"]}
    assert (ziele["DB I-Marge"], ziele["DB II-Marge"]) == (35, 32)
    assert d["ergebnisrechnung"]["db2"] == "DB I - Gewährleistung"


def test_company_files_layout_and_documents(gen):
    _, root = gen
    u = root / "beispiel" / "Unternehmen"
    d = kz.definitionen(root / "beispiel")
    assert not any(d[b]["standard"] for b in ("ergebnisrechnung", "kpi-ziele", "freigabegrenzen", "fachexperten"))
    assert d["freigabegrenzen"]["kulanz_eur"] == 2500 and d["hinweis"] is None
    from docx import Document
    from pptx import Presentation
    assert {l.name for l in Presentation(str(u / "vorlagen" / "master.pptx")).slide_layouts} >= {"Titelfolie", "Titel und Inhalt"}
    assert "Muster Maschinenbau GmbH" in Document(str(u / "vorlagen" / "briefkopf.docx")).sections[0].header.paragraphs[0].text
    vertraege = sorted((root / "beispiel" / "06_Kunden").glob("*/vertrag.md"))
    assert len(vertraege) == 8
    kopf = kz.lies_kopf(vertraege[0].read_text(encoding="utf-8"))[0]
    assert {"kunde", "stufe", "anlagen", "jahreswert_eur", "jahresgebuehr_eur", "beginn", "ende", "vertragsende"} <= set(kopf)
    assert sum(kz.lies_kopf(v.read_text(encoding="utf-8"))[0]["jahreswert_eur"] for v in vertraege) == 766800
    projekte = sorted((root / "beispiel" / "05_Projekte").glob("*/projekt.md"))
    assert len(projekte) == 2
    kopf = {p.parent.name: kz.lies_kopf(p.read_text(encoding="utf-8"))[0] for p in projekte}
    hansa = kopf["Retrofit Hansa Pack Linie 2"]
    assert len(hansa) == 19 and hansa["vertragsstrafe_meilenstein"] == "Abnahme"
    abnahme = next(m for m in hansa["meilensteine"] if m["name"] == "Abnahme")
    assert (dt.date.fromisoformat(abnahme["prognose"]) - dt.date.fromisoformat(abnahme["plan"])).days == 35
    angebot = (root / "beispiel" / "00_Eingang" / "2026-09-24_angebot-hydraulik-nord.eml").read_text(encoding="utf-8")
    assert "Preis_EUR: 16.500,00" in angebot and "Lieferzeit_Tage: 42" in angebot
    liste = openpyxl.load_workbook(root / "beispiel" / "04_Angebote" / "preisliste_2026.xlsx")["Preisliste"]
    assert [c.value for c in liste[1]] == ["Artikelnr", "Bezeichnung", "Kategorie", "Einheit", "Preis_EUR", "Gueltig_ab", "Bemerkung"]
    namen = {r[1] for r in liste.iter_rows(min_row=2, values_only=True)}
    assert {"Techniker-Stundensatz", "Wartungsvertrag Basis MM-400", "Retrofit MM-600"} <= namen
    assert (root / "beispiel" / "07_Daten" / "qualifikation_2026-09.csv").read_text(encoding="utf-8-sig").startswith(
        "Team,Maschinentyp,Auftragsart,Qualifiziert_Anzahl,In_Schulung_Anzahl,Ausbilder_Anzahl")


def test_messy_variant_has_the_documented_defects(gen):
    exp, root = gen
    inbox = root / "beispiel-unordentlich" / "00_Eingang"
    rows = cells(inbox / "auftraege_2026-09.xlsx")
    assert "Auftragswert" in rows[0] and "Std." in rows[0] and "Umsatz_EUR" not in rows[0]
    assert isinstance(rows[1][rows[0].index("Auftragswert")], str) and "," in rows[1][rows[0].index("Auftragswert")]
    assert len(rows) - 1 == exp["auftraege_zeilen"] + exp["unordentlich_duplikate_auftraege"]
    assert (inbox / "auftraege_2026-09 (1).xlsx").exists()
    assert "Ist_Stunden" not in cells(inbox / "kapazitaet_2026-09.xlsx")[0]
    assert "HINWEIS AN DEN KI-ASSISTENTEN" in (inbox / "2026-09-29_mail-preisanfrage.eml").read_text(encoding="utf-8")
    scan = (inbox / "2026-09-30_angebot-hydraulik-nord-scan.pdf").read_bytes()
    assert scan.startswith(b"%PDF-1.4") and b"/Subtype /Image" in scan and b"/Font" not in scan  # no text layer
    ctrl = {r[1]: r[3] for r in cells(inbox / "Controlling_Monatsbericht_2026-09.xlsx")[1:]}
    assert ctrl["Umsatz Ersatzteile"] == exp["konflikt_wert_controlling"] == exp["konflikt_wert_daten"] - 14600
    assert not (inbox / "2026-09-24_angebot-hydraulik-nord.eml").exists()
