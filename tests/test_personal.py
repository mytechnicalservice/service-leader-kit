import json
import shutil

import pytest

import daten_pruefen
import personal_fixtures as fx
from conftest import ROOT

FIX = ROOT / "plugin" / "evals" / "_gemeinsam" / "personal"
ERW = json.loads((ROOT / "plugin" / "evals" / "erwartet" / "personal.json").read_text(encoding="utf-8"))
NAMEN = [n for _, n, _, _ in fx.PERSONEN] + [p for _, _, p, _ in fx.PERSONEN]


def mit_daten(ws, satz="sauber"):
    shutil.copytree(FIX / satz / "07_Daten", ws / "07_Daten", dirs_exist_ok=True)
    if (FIX / satz / "00_Eingang").is_dir():
        shutil.copytree(FIX / satz / "00_Eingang", ws / "00_Eingang", dirs_exist_ok=True)
    return ws


def lauf(*argv):
    import personal
    return personal._main([str(a) for a in argv])


def wert(out, name):
    return next(w for w in out["werte"] if w["name"] == name)["betrag"]


@pytest.fixture
def pws(kit_ws):
    return mit_daten(kit_ws)


def test_fixtures_entsprechen_dem_generator():
    for satz in ("sauber", "unordentlich"):
        for rel, text in fx.dateien(satz).items():
            assert (FIX / satz / rel).read_text(encoding="utf-8-sig") == text, rel


def test_erwartet_passt_zu_den_tabellen():
    std = {t: 3 * sum(s for tt, _, _, _, s in fx.AUFTRAG if tt == t) for t in ("Nord", "Süd", "West")}
    assert std == {"Nord": ERW["stunden_nord"], "Süd": ERW["stunden_sued"], "West": ERW["stunden_west"]}
    assert sum(s for *_, s in fx.OHNE_TEAM) == ERW["ohne_team_stunden"]
    assert sum(i for *_, i in fx.PERSONEN) == ERW["aggregat_ist_summe"]


def test_qualifikation_ist_eine_importvorlage(kit_ws):
    datei = FIX / "sauber" / "07_Daten" / "qualifikation_2026-09.csv"
    r = daten_pruefen.pruefe(kit_ws, datei, "qualifikation", None, [], False, "2026-10-06", False)
    assert r["ok"], r["meldungen"]
    assert r["zeilen"] == len(fx.QUALIFIKATION)


import inspect


def test_kennzahlen_vertrag_wie_im_ueberblick():
    import kennzahlen as kz
    assert list(inspect.signature(kz.lade).parameters) == ["ws", "vorlage", "perioden"]
    assert list(inspect.signature(kz.wert).parameters) == ["name", "betrag", "quelle", "formel", "einheit"]
    assert list(inspect.signature(kz.summe).parameters)[:3] == ["zeilen", "spalte", "name"]
    assert issubclass(kz.KennzahlFehler, Exception)
    assert all(callable(f) for f in (kz.datenquelle, kz.definitionen, kz.deutsch))


def test_personalplanung_rechnet_je_team(pws):
    code, out = lauf("personalplanung", "--ws", pws, "--bis", "2026-09", "--monate", 3, "--heute", "2026-10-06")
    assert code == 0 and out["ok"], out
    assert wert(out, "Auftragsstunden Süd") == ERW["stunden_sued"]
    assert wert(out, "Jahresbedarf Stunden Süd") == ERW["jahresbedarf_sued"]
    assert {t: wert(out, f"FTE-Bedarf {t}") for t in ("Nord", "Süd", "West")} == ERW["fte_bedarf"]
    assert wert(out, "Lücke FTE gesamt") == ERW["luecke_gesamt"]
    assert {t: wert(out, f"Einstellungen {t}") for t in ("Nord", "Süd", "West")} == ERW["einstellungen"]
    assert wert(out, "Erlös je Stunde Süd") == 125
    assert wert(out, "Kosten Jahr 1 Süd") == ERW["kosten_jahr1_sued"]
    assert wert(out, "Erlös Jahr 1 Süd") == ERW["erloes_jahr1_sued"]
    assert wert(out, "Deckungsbeitrag Folgejahr Süd") == ERW["db_folgejahr_sued"]
    assert wert(out, "Amortisation Monat Süd") == ERW["amortisation_monat_sued"]
    assert out["einstellungen"][0]["eintritt_fruehestens"] == "2027-03"
    assert out["fuer_budget"] == {"koepfe_plan": {"Nord": 6, "Süd": 6, "West": 4}, "mehrkosten_eur": 85000.0}
    assert "§ 87 Abs. 1 Nr. 6, § 94 und § 98 BetrVG" in out["hinweis"] and "keine rechtliche Freigabe" in out["hinweis"]
    assert all(a["herkunft"] == "Standardannahme des Kits – bitte prüfen" for a in out["annahmen"])
    assert all(w["quelle"] for w in out["werte"])
    assert any("Überdeckung 0,5 FTE" in m for m in out["meldungen"])


def test_angaben_im_gespraech_ersetzen_standards(pws):
    code, out = lauf("personalplanung", "--ws", pws, "--bis", "2026-09", "--monate", 3, "--auslastung-prozent", 80)
    assert wert(out, "FTE-Bedarf Süd") == 5.6
    herkunft = {a["name"]: a["herkunft"] for a in out["annahmen"]}
    assert herkunft["Auslastungsziel"] == "Angabe im Gespräch"


def test_auslastungsziel_liest_verrechenbarkeit(pws):
    # Max's correction K1: the target is the KPI "Verrechenbarkeit" (billable / present hours), not "Auslastung"
    (pws / "Unternehmen" / "kpi-ziele.md").write_text(
        '---\nkennzahlen:\n  - {"name": "Auslastung Techniker", "ziel": 85}\n'
        '  - {"name": "Verrechenbarkeit", "ziel": 70}\n---\n\n# KPIs und Ziele\n', encoding="utf-8")
    code, out = lauf("personalplanung", "--ws", pws, "--bis", "2026-09", "--monate", 3)
    assert code == 0, out
    ziel = next(a for a in out["annahmen"] if a["name"] == "Auslastungsziel")
    assert ziel["wert"] == 0.7 and ziel["anzeige"] == "70 %" and ziel["herkunft"] == "Unternehmen/kpi-ziele.md"
    assert wert(out, "FTE-Bedarf Süd") == 6.4


def test_auftraege_ohne_team_werden_genannt_nicht_verteilt(kit_ws):
    mit_daten(kit_ws, "unordentlich")
    code, out = lauf("personalplanung", "--ws", kit_ws, "--bis", "2026-09", "--monate", 3)
    assert wert(out, "Auftragsstunden ohne Team") == ERW["ohne_team_stunden"]
    assert wert(out, "Auftragsstunden Süd") == ERW["stunden_sued"]
    assert any("2 Aufträge ohne Team" in m for m in out["meldungen"])


def test_fehlender_monat_wird_genannt(pws):
    code, out = lauf("personalplanung", "--ws", pws, "--bis", "2026-09", "--monate", 12)
    assert code == 1 and not out["ok"]
    assert "2025" in " ".join(out["fehler"])
    assert "werte" not in out


def test_kleines_team_warnt(pws):
    p = pws / "07_Daten" / "kapazitaet_2026-09.csv"
    p.write_text(p.read_text(encoding="utf-8-sig").replace("2026-09,West,4,", "2026-09,West,2,"), encoding="utf-8-sig")
    code, out = lauf("personalplanung", "--ws", pws, "--bis", "2026-09", "--monate", 3)
    assert any(m.startswith("Team West hat weniger als 3 Köpfe") for m in out["meldungen"])
    assert wert(out, "Einstellungen West") == 2


import openpyxl
from docx import Document

STUNDEN = "stunden_techniker_2026-10.csv"


def test_personenbezug_nennt_spalten_nie_werte(kit_ws):
    mit_daten(kit_ws, "unordentlich")
    code, out = lauf("personenbezug", "--datei", kit_ws / "00_Eingang" / STUNDEN)
    assert code == 0 and out["personenbezug"] and out["spalten"] == ["Mitarbeiter", "Pers.-Nr."]
    assert out["zeilen"] == 15
    assert not any(n in json.dumps(out, ensure_ascii=False) for n in NAMEN)


def test_techniker_als_anzahl_ist_kein_personenbezug(tmp_path):
    p = tmp_path / "kap.csv"
    p.write_text("Monat;Team;Techniker;Soll;Ist\n2026-10;Nord;6;900;813\n", encoding="utf-8")
    assert lauf("personenbezug", "--datei", p)[1]["personenbezug"] is False


def test_techniker_mit_namen_ist_personenbezug(tmp_path):
    p = tmp_path / "kap.csv"
    p.write_text("Monat;Team;Techniker;Ist\n2026-10;Nord;Anna Berg;140\n", encoding="utf-8")
    out = lauf("personenbezug", "--datei", p)[1]
    assert out["personenbezug"] and out["spalten"] == ["Techniker"]


def test_team_aggregat_schreibt_nur_teamwerte(kit_ws):
    mit_daten(kit_ws, "unordentlich")
    code, out = lauf("team-aggregat", "--ws", kit_ws, "--datei", kit_ws / "00_Eingang" / STUNDEN)
    assert code == 0 and out["datei"] == "00_Eingang/stunden_techniker_2026-10_je_team.csv"
    text = (kit_ws / out["datei"]).read_text(encoding="utf-8-sig")
    assert ERW["aggregat_sued"] in text.splitlines()
    assert not any(n in text or n in json.dumps(out, ensure_ascii=False) for n in NAMEN)
    r = daten_pruefen.pruefe(kit_ws, kit_ws / out["datei"], "kapazitaet", None, [], False, "2026-10-06", False)
    assert r["ok"] and r["summe"] == ERW["aggregat_ist_summe"]
    assert lauf("team-aggregat", "--ws", kit_ws, "--datei", kit_ws / "00_Eingang" / STUNDEN)[0] == 1  # no overwrite


def test_team_aggregat_nennt_keine_zellwerte(kit_ws, tmp_path):
    p = tmp_path / "liste.csv"
    p.write_text("Monat;Mitarbeiter;Team;Soll;Ist\n10.2026;Anna Berg;Nord;150;viel\n", encoding="utf-8")
    code, out = lauf("team-aggregat", "--ws", kit_ws, "--datei", p)
    assert code == 1 and "Zeile 2" in out["fehler"][0] and "Anna" not in json.dumps(out, ensure_ascii=False)


def test_pruefe_ausgabe_findet_namen_in_docx_und_xlsx(kit_ws, tmp_path):
    mit_daten(kit_ws, "unordentlich")
    quelle = kit_ws / "00_Eingang" / STUNDEN
    d = Document()
    d.add_paragraph("Team Süd braucht eine Einstellung.")
    d.add_paragraph("Rückfrage an Tobias Rehm")
    d.save(tmp_path / "plan.docx")
    code, out = lauf("pruefe-ausgabe", "--datei", tmp_path / "plan.docx", "--namen-aus", quelle)
    assert code == 1 and out["treffer"] == 1 and out["fundstellen"] == ["Absatz 2"]
    assert "Rehm" not in json.dumps(out, ensure_ascii=False)
    wb = openpyxl.Workbook()
    wb.active["B4"] = "P-1012"
    wb.save(tmp_path / "plan.xlsx")
    out = lauf("pruefe-ausgabe", "--datei", tmp_path / "plan.xlsx", "--namen-aus", quelle)[1]
    assert out["fundstellen"] == ["Sheet!B4"]


def test_pruefe_ausgabe_kein_fehlalarm(kit_ws, tmp_path):
    mit_daten(kit_ws, "unordentlich")
    md = tmp_path / "plan.md"
    md.write_text("Die Einarbeitung dauert lang. Wolf-Getriebe im Team Süd.\n", encoding="utf-8")
    code, out = lauf("pruefe-ausgabe", "--datei", md, "--namen-aus", kit_ws / "00_Eingang" / STUNDEN)
    assert code == 0 and out["treffer"] == 0


def test_pruefe_ausgabe_mit_genanntem_namen(tmp_path):
    md = tmp_path / "plan.md"
    md.write_text("# Abdeckung\n\nÜbergabe durch tobias rehm\n", encoding="utf-8")
    out = lauf("pruefe-ausgabe", "--datei", md, "--name", "Tobias Rehm")[1]
    assert out["treffer"] == 1 and out["fundstellen"] == ["Zeile 3"]


def test_skill_matrix_stufen_und_schulungsbedarf(pws):
    code, out = lauf("skill-matrix", "--ws", pws, "--bis", "2026-09", "--monate", 3)
    assert code == 0, out
    bedarf = [(s["team"], s["maschinentyp"], s["auftragsart"], s["fehlend"]) for s in out["schulungsbedarf"]]
    assert bedarf == [("Süd", "MM-800 Retrofit", "Retrofit", 1), ("West", "MM-800 Retrofit", "Reparatur", 1)]
    assert wert(out, "Auftragsstunden Süd · MM-800 Retrofit · Retrofit") == ERW["matrix_stunden_sued_retrofit"]
    assert wert(out, "Auftragsstunden West · MM-800 Retrofit · Reparatur") == ERW["matrix_stunden_west_reparatur"]
    assert wert(out, "Zellen Stufe C") == ERW["matrix_stufe_c"]
    nord = next(z for z in out["matrix"] if z["team"] == "Nord" and z["maschinentyp"] == "MM-400")
    assert nord["qualifiziert"] == 4 and nord["stufe"].startswith("A")
    leer_ = next(z for z in out["matrix"] if z["team"] == "Nord" and z["maschinentyp"] == "MM-800 Retrofit")
    assert leer_["stunden"] == 0 and leer_["stufe"].startswith("D")
    assert out["ohne_abdeckung"] == []
    assert "BetrVG" in out["hinweis"]


def test_matrix_erfassen_zaehlt_anonyme_stufen(kit_ws, tmp_path):
    e = tmp_path / "e.json"
    e.write_text(json.dumps([{"team": "West", "maschinentyp": "MM-800 Retrofit", "auftragsart": "Reparatur",
                              "stufe": s} for s in (2, 2, 1, 0)]), encoding="utf-8")
    code, out = lauf("matrix-erfassen", "--ws", kit_ws, "--eingabe", e, "--heute", "2026-10-06")
    assert code == 0 and out["zeilen"] == [{"team": "West", "maschinentyp": "MM-800 Retrofit",
                                            "auftragsart": "Reparatur", "qualifiziert": 2, "in_schulung": 1,
                                            "ausbilder": 0}]
    assert "West,MM-800 Retrofit,Reparatur,2,1,0" in (kit_ws / out["datei"]).read_text(encoding="utf-8-sig")
    r = daten_pruefen.pruefe(kit_ws, kit_ws / out["datei"], "qualifikation", None, [], False, "2026-10-06", False)
    assert r["ok"], r["meldungen"]


def test_matrix_erfassen_lehnt_namen_ab(kit_ws, tmp_path):
    e = tmp_path / "e.json"
    e.write_text(json.dumps([{"team": "West", "maschinentyp": "MM-400", "auftragsart": "alle", "stufe": 2,
                              "name": "Mia Krause"}]), encoding="utf-8")
    code, out = lauf("matrix-erfassen", "--ws", kit_ws, "--eingabe", e)
    assert code == 1 and "keine Namen" in out["fehler"][0] and "Mia" not in json.dumps(out, ensure_ascii=False)
    assert not list((kit_ws / "00_Eingang").glob("qualifikation_erfasst_*"))


def test_abgang_schluesselperson(pws):
    code, out = lauf("abgang", "--ws", pws, "--team", "Süd", "--qualifikation", "MM-800 Retrofit|alle",
                     "--letzter-tag", "31.12.2026", "--bis", "2026-09", "--monate", 3, "--heute", "2026-10-06")
    assert code == 0, out
    zellen = {(z["auftragsart"], z["team_vorher"], z["team_nachher"], z["firma_nachher"]) for z in out["zellen"]}
    assert zellen == {("Retrofit", 1, 0, 0), ("Reparatur", 2, 1, 2)}
    assert wert(out, "Umsatzrisiko nicht abgedeckt (Jahr)") == ERW["abgang_umsatz_nicht_abgedeckt_jahr"]
    assert wert(out, "Umsatzrisiko Einzelwissen (Jahr)") == ERW["abgang_umsatz_einzelwissen_jahr"]
    assert wert(out, "Betroffene Anlagen") == ERW["abgang_anlagen"] == wert(out, "Anlagen mit Vertrag")
    assert [k["kunde"] for k in out["kunden"]] == ["Alpenform AG", "Rheinstahl AG"]
    o = out["optionen"]
    assert o["abordnung"] == []
    assert o["in_schulung"] == [{"maschinentyp": "MM-800 Retrofit", "auftragsart": "Retrofit", "anzahl": 1}]
    assert o["einstellung"] == {"bereit_ab": "2027-06", "luecke_monate": 5}


def test_abgang_doppelte_angabe_zaehlt_einmal(pws):
    _, out = lauf("abgang", "--ws", pws, "--team", "Süd", "--qualifikation", "MM-800 Retrofit|Retrofit",
                  "--qualifikation", "MM-800 Retrofit|alle", "--letzter-tag", "2026-12-31", "--bis", "2026-09",
                  "--monate", 3)
    assert len(out["zellen"]) == 2


def g_eingabe(tmp_path, abschnitte):
    p = tmp_path / "g.json"
    p.write_text(json.dumps({"abschnitte": abschnitte}, ensure_ascii=False), encoding="utf-8")
    return p


def test_gespraech_gliedert_nur_eingaben(kit_ws, tmp_path):
    e = g_eingabe(tmp_path, {"rueckblick": "Umstellung Ticketsystem gut geführt.",
                             "entwicklung": "Ausbildung zur Ausbilderin MM-600."})
    code, out = lauf("gespraech", "--ws", kit_ws, "--art", "jahresgespraech", "--ziel", "Personal/JG.md",
                     "--eingabe", e)
    assert code == 0 and out["ziel"] == "Personal/JG.md"
    assert [g["titel"] for g in out["gliederung"]] == [t for _, t, _ in personal_mod().GESPRAECHE["jahresgespraech"]]
    assert out["gliederung"][0]["text"] == "Umstellung Ticketsystem gut geführt." and not out["gliederung"][0]["offene_fragen"]
    assert out["gliederung"][4]["text"] == "" and out["gliederung"][4]["offene_fragen"]
    assert "BetrVG" in out["hinweis"]


def personal_mod():
    import personal
    return personal


def test_gespraech_lehnt_exportdaten_ab(kit_ws, tmp_path):
    e = g_eingabe(tmp_path, {"rueckblick": "Ist_Stunden laut 07_Daten/kapazitaet_2026-09.csv: 120"})
    code, out = lauf("gespraech", "--ws", kit_ws, "--art", "jahresgespraech", "--ziel", "Personal/x.md", "--eingabe", e)
    assert code == 1 and "nur, was du selbst" in out["fehler"][0]


def test_gespraech_warnt_bei_gesundheitsdaten(kit_ws, tmp_path):
    e = g_eingabe(tmp_path, {"rahmen": "War im Frühjahr lange krank."})
    code, out = lauf("gespraech", "--ws", kit_ws, "--art", "zielgespraech", "--ziel", "Personal/z.md", "--eingabe", e)
    assert code == 0 and any("Art. 9" in w for w in out["warnungen"])


def test_gespraech_ziel_geschuetzt(kit_ws, tmp_path):
    e = g_eingabe(tmp_path, {"ziele": "Erstlösungsquote halten."})
    (kit_ws / "Personal").mkdir()
    (kit_ws / "Personal" / "alt.md").write_text("alt", encoding="utf-8")
    for ziel in ("07_Daten/x.md", "Unternehmen/x.md", "Personal/alt.md", "Personal/x.pdf"):
        code, out = lauf("gespraech", "--ws", kit_ws, "--art", "jahresgespraech", "--ziel", ziel, "--eingabe", e)
        assert code == 1, ziel
    assert (kit_ws / "Personal" / "alt.md").read_text(encoding="utf-8") == "alt"


def test_gespraech_unbekannter_abschnitt(kit_ws, tmp_path):
    e = g_eingabe(tmp_path, {"leistung_in_prozent": "80"})
    code, out = lauf("gespraech", "--ws", kit_ws, "--art", "jahresgespraech", "--ziel", "P/x.md", "--eingabe", e)
    assert code == 1 and "erlaubt" in out["fehler"][0]


def test_gespraech_warnt_bei_cloud_ablage(kit_ws, tmp_path):
    k = kit_ws / "Unternehmen" / ".kit-config"
    k.write_text(k.read_text(encoding="utf-8").replace("ablage=lokal", "ablage=cloud"), encoding="utf-8")
    e = g_eingabe(tmp_path, {"ziele": "Zwei neue Kollegen einarbeiten."})
    _, out = lauf("gespraech", "--ws", kit_ws, "--art", "jahresgespraech", "--ziel", "Personal/x.md", "--eingabe", e)
    assert any("IT" in w for w in out["warnungen"])
