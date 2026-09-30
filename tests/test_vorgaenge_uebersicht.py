import json

import openpyxl

import vorgaenge_uebersicht as vu
from vorgang import main as vorgang

HEUTE = "2026-09-30"


def neu(ws, titel, faellig=None):
    args = ["neu", "--ws", str(ws), "--titel", titel, "--typ", "aufgabe", "--kunde", "Müller GmbH",
            "--verantwortlich", "Jana", "--von", "assistenz", "--text", "x", "--heute", HEUTE]
    assert vorgang(args + (["--faellig", faellig] if faellig else [])) == 0


def test_build_writes_protected_view_with_overdue_first(capsys, ws):
    neu(ws, "Später", "2026-10-15")
    neu(ws, "Überfällig", "2026-09-01")
    neu(ws, "Fertig")
    vorgang(["schliesse", "--ws", str(ws), "--nr", "V-0003", "--heute", HEUTE])
    capsys.readouterr()
    assert vu.main(["--ws", str(ws), "--heute", HEUTE]) == 0
    out = json.loads(capsys.readouterr().out)
    assert (out["offen"], out["ueberfaellig"], out["erledigt"]) == (2, 1, 1)
    wb = openpyxl.load_workbook(ws / "01_Vorgaenge" / vu.DATEINAME)
    sh = wb["Offen"]
    assert sh["A1"].value == vu.HINWEIS
    assert [c.value for c in sh[3]][:4] == ["Nr", "Titel", "Typ", "Status"]
    assert [sh["B4"].value, sh["H4"].value, sh["B5"].value, sh["H5"].value] == ["Überfällig", "ja", "Später", "nein"]
    assert sh.protection.sheet is True
    assert wb["Erledigt"]["B4"].value == "Fertig"


def test_rebuild_overwrites_user_edits(capsys, ws):
    neu(ws, "Original")
    vu.build(ws, HEUTE)
    path = ws / "01_Vorgaenge" / vu.DATEINAME
    wb = openpyxl.load_workbook(path)
    wb["Offen"]["B4"] = "Vom Nutzer geändert"
    wb.save(path)
    vu.build(ws, HEUTE)
    assert openpyxl.load_workbook(path)["Offen"]["B4"].value == "Original"
