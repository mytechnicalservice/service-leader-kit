"""Regression tests for the Plan 2a final-review findings (C-1..C-4, I-1..I-9)."""
import csv
import json
import os
import shutil
import subprocess
import sys
import zipfile

import openpyxl
import pytest

import daten_pruefen as dp
import vorgaenge_uebersicht as vu
from conftest import ROOT
from slk_common import zahl
from vorgang import main as vorgang

HEUTE = "2026-09-30"
SCRIPTS = ROOT / "plugin" / "scripts"


def run(capsys, fn, *args):
    code = fn([*args])
    return code, json.loads(capsys.readouterr().out)


def neu(capsys, ws, titel="Reklamation Spindel", *extra):
    return run(capsys, vorgang, "neu", "--ws", str(ws), "--titel", titel, "--typ", "reklamation", "--kunde",
               "Müller GmbH", "--verantwortlich", "Jana", "--von", "betrieb", "--text", "x", "--heute", HEUTE, *extra)


def offen(ws):
    return sorted(os.listdir(ws / "01_Vorgaenge" / "offen"))


def xlsx(path, *rows, sheets=None):
    wb = openpyxl.Workbook()
    if sheets:
        wb.active.title = sheets[0][0]
        for r in sheets[0][1]:
            wb.active.append(r)
        for name, rs in sheets[1:]:
            sh = wb.create_sheet(name)
            for r in rs:
                sh.append(r)
    else:
        for r in rows:
            wb.active.append(r)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# --- vorgang.py -----------------------------------------------------------------------------------------

def test_c1_lowercase_nr_is_normalised_and_never_deletes(capsys, ws):
    neu(capsys, ws)
    code, out = run(capsys, vorgang, "eintrag", "--ws", str(ws), "--nr", "v-0001", "--art", "notiz", "--von",
                    "betrieb", "--text", "y", "--heute", HEUTE)
    assert code == 0 and out["nr"] == "V-0001"
    assert offen(ws) == ["V-0001.md"]


def test_c2_copied_case_file_is_reported_not_merged(capsys, ws):
    neu(capsys, ws)
    src = ws / "01_Vorgaenge" / "offen" / "V-0001.md"
    before = src.read_text(encoding="utf-8")
    shutil.copy(src, src.with_name("V-0005.md"))
    code, out = run(capsys, vorgang, "pruefe", "--ws", str(ws), "--heute", HEUTE)
    assert code == 1 and "passt nicht zum Dateinamen" in json.dumps(out, ensure_ascii=False)
    code, out = run(capsys, vorgang, "eintrag", "--ws", str(ws), "--nr", "V-0005", "--art", "notiz", "--von",
                    "betrieb", "--text", "y", "--heute", HEUTE)
    assert code == 1 and src.read_text(encoding="utf-8") == before and offen(ws) == ["V-0001.md", "V-0005.md"]


def test_c3_parallel_neu_never_loses_a_case(ws):
    procs = [subprocess.Popen([sys.executable, str(SCRIPTS / "vorgang.py"), "neu", "--ws", str(ws), "--titel",
                               f"Fall {i}", "--typ", "aufgabe", "--kunde", "K", "--verantwortlich", "J", "--von",
                               "assistenz", "--text", "x", "--trotzdem"], stdout=subprocess.PIPE)
             for i in range(10)]
    outs = [json.loads(p.communicate()[0]) for p in procs]
    nrs = [o["nr"] for o in outs if o["ok"]]
    assert len(nrs) == 10 == len(set(nrs)) == len(offen(ws))
    assert not [f for f in os.listdir(ws / "01_Vorgaenge" / "offen") if f.endswith(".tmp")]


def test_i1_validate_checks_types_and_folder(capsys, ws):
    neu(capsys, ws)
    p = ws / "01_Vorgaenge" / "offen" / "V-0001.md"
    text = p.read_text(encoding="utf-8")
    p.write_text(text.replace('bearbeitet_von: ["betrieb"]', 'bearbeitet_von: "betrieb"')
                 .replace("betrag_eur: null", 'betrag_eur: "8.400 EUR"').replace('kunde: "Müller GmbH"', "kunde: 42")
                 .replace('status: "offen"', 'status: "erledigt"'), encoding="utf-8")
    code, out = run(capsys, vorgang, "pruefe", "--ws", str(ws), "--heute", HEUTE)
    msgs = " | ".join(out["defekt"][0]["fehler"])
    assert code == 1
    for m in ("'bearbeitet_von' muss eine Liste", "'betrag_eur' muss eine Zahl", "'kunde' muss Text",
              "Status 'erledigt' passt nicht zum Ordner 'offen'"):
        assert m in msgs


def test_i1_damaged_other_case_is_named(capsys, ws):
    neu(capsys, ws)
    bad = ws / "01_Vorgaenge" / "offen" / "V-0009.md"
    bad.write_text("---\nnr: \"V-0009\"\n", encoding="utf-8")
    code, out = neu(capsys, ws, "Neuer Fall")
    assert code == 1 and "01_Vorgaenge/offen/V-0009.md" in out["fehler"][0]
    code, out = run(capsys, vu.main, "--ws", str(ws), "--heute", HEUTE)
    assert code == 1 and out["ok"] is False and "V-0009.md" in json.dumps(out, ensure_ascii=False)


def test_i2_i3_cli_errors_are_json_and_amounts_parse_german(capsys, ws):
    code, out = run(capsys, vorgang, "neu", "--ws", str(ws))
    assert code == 1 and out["ok"] is False and "Aufruf fehlerhaft" in out["fehler"][0]
    code, out = neu(capsys, ws, "Mit Betrag", "--betrag", "1.234,56")
    assert code == 0
    code, _ = run(capsys, vorgang, "setze", "--ws", str(ws), "--nr", "V-0001", "--feld", "betrag_eur", "--wert",
                  "1.234", "--heute", HEUTE)
    text = (ws / "01_Vorgaenge" / "offen" / "V-0001.md").read_text(encoding="utf-8")
    assert code == 0 and "betrag_eur: 1234.0" in text
    code, out = run(capsys, vorgang, "setze", "--ws", str(ws), "--nr", "V-0001", "--feld", "betrag_eur", "--wert",
                    "viel", "--heute", HEUTE)
    assert code == 1 and "keine Zahl" in out["fehler"][0]


def test_i3_zahl_rules():
    assert (zahl("0.500"), zahl("-0.250"), zahl("1.234"), zahl("1.234,56-")) == (0.5, -0.25, 1234.0, -1234.56)


def test_i8_loeschen_verweise_scan_office_and_mail_with_word_boundary(capsys, ws):
    neu(capsys, ws)
    (ws / "03_Berichte").mkdir()
    with zipfile.ZipFile(ws / "03_Berichte" / "Bericht.docx", "w") as z:
        z.writestr("word/document.xml", "<w:t>Kulanz zu V-0001 abgelehnt</w:t>")
    (ws / "02_Postausgang").mkdir()
    (ws / "02_Postausgang" / "Antwort.eml").write_text("Bezug: V-0001", encoding="utf-8")
    (ws / "03_Berichte" / "Andere.md").write_text("Siehe V-00012", encoding="utf-8")
    code, out = run(capsys, vorgang, "loeschen-markieren", "--ws", str(ws), "--nr", "V-0001", "--heute", HEUTE)
    assert code == 0 and out["verweise"] == ["02_Postausgang/Antwort.eml", "03_Berichte/Bericht.docx"]


def test_i6_cli_output_is_utf8_even_with_cp1252_stdout(ws):
    env = os.environ | {"PYTHONIOENCODING": "cp1252"}
    cmd = [sys.executable, str(SCRIPTS / "vorgang.py"), "neu", "--ws", str(ws), "--titel", "Łódź →", "--typ",
           "aufgabe", "--kunde", "Łódź Maszyny", "--verantwortlich", "J", "--von", "assistenz", "--text", "x"]
    subprocess.run(cmd, capture_output=True, env=env)
    raw = subprocess.run(cmd, capture_output=True, env=env).stdout  # duplicate → German message with "ergänzen"
    out = json.loads(raw.decode("utf-8"))
    assert out["ok"] is False and "ergänzen" in out["meldung"]


# --- vorgaenge_uebersicht.py --------------------------------------------------------------------------

def test_i5_formula_title_stays_text(capsys, ws):
    neu(capsys, ws, '=HYPERLINK("http://evil.example","Details")')
    vu.build(ws, HEUTE)
    cell = openpyxl.load_workbook(ws / "01_Vorgaenge" / vu.DATEINAME)["Offen"]["B4"]
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")


def test_i5_locked_overview_returns_json_and_leaves_no_tmp(capsys, ws, monkeypatch):
    neu(capsys, ws)

    def locked(*_):
        raise PermissionError("in use")

    monkeypatch.setattr(vu.os, "replace", locked)
    code, out = run(capsys, vu.main, "--ws", str(ws), "--heute", HEUTE)
    assert code == 1 and "schließen" in out["fehler"][0]
    assert not [f for f in os.listdir(ws / "01_Vorgaenge") if f.endswith(".tmp")]


# --- daten_pruefen.py ---------------------------------------------------------------------------------

ERG_H = ["Monat", "Position", "Ist_EUR"]


def test_c4_original_with_same_name_is_never_overwritten(ws):
    a = xlsx(ws / "00_Eingang" / "export.xlsx", ERG_H, ["2026-08", "Umsatz", 100])
    assert dp.pruefe(ws, a, "ergebnis", None, [], True, HEUTE)["ok"]
    b = xlsx(ws / "00_Eingang" / "export.xlsx", ERG_H, ["2026-09", "Umsatz", 200])
    assert dp.pruefe(ws, b, "ergebnis", None, [], True, HEUTE)["ok"]
    assert len(os.listdir(ws / "07_Daten" / "original")) == 2


def test_i4_empty_mandatory_cell_is_an_error(ws, tmp_path):
    f = xlsx(tmp_path / "ib.xlsx", ["Kunde", "Anlage", "Maschinentyp", "Vertrag"], [None, "A1", "MM", None])
    r = dp.pruefe(ws, f, "installed_base", None, [], False, HEUTE)
    assert not r["ok"] and "Zeile 2, Spalte 'Kunde': Pflichtwert fehlt" in r["meldungen"]


def test_i4_two_columns_for_one_target_are_refused(ws, tmp_path):
    f = xlsx(tmp_path / "e.xlsx", ["Monat", "Position", "Ist", "Istwert"], ["2026-09", "U", 100, 119])
    r = dp.pruefe(ws, f, "ergebnis", None, [], False, HEUTE)
    assert not r["ok"] and "passen beide zu 'Ist_EUR'" in r["meldungen"][0]


def test_i4_identical_parts_rows_need_confirmation(ws, tmp_path):
    row = ["2026-09-03", "Müller GmbH", "ET-1", 2, 50]
    f = xlsx(tmp_path / "t.xlsx", ["Datum", "Kunde", "Teilenr", "Menge", "Stueckpreis_EUR"], row, row)
    r = dp.pruefe(ws, f, "ersatzteile", None, [], False, HEUTE)
    assert not r["ok"] and "identisch" in r["meldungen"][0]
    r = dp.pruefe(ws, f, "ersatzteile", None, [], False, HEUTE, duplikate_behalten=True)
    assert r["ok"] and r["zeilen"] == 2


def test_i4_control_total_without_sum_column_is_refused(ws, daten_gen):
    f = daten_gen / "beispiel" / "00_Eingang" / "ersatzteile_2026-09.xlsx"
    r = dp.pruefe(ws, f, "ersatzteile", "999.999,00", [], False, HEUTE)
    assert not r["ok"] and "keine Summenspalte" in r["meldungen"][0]


def test_i4_sum_row_in_export_is_used_as_control_total(ws, tmp_path):
    ok = xlsx(tmp_path / "a.xlsx", ERG_H, ["2026-09", "Umsatz", 100], ["2026-09", "Kosten", -40], [None, "Summe", 60])
    r = dp.pruefe(ws, ok, "ergebnis", None, [], False, HEUTE)
    assert r["ok"] and (r["summe"], r["zeilen"]) == (60.0, 2)
    bad = xlsx(tmp_path / "b.xlsx", ERG_H, ["2026-09", "Umsatz", 100], ["2026-09", "Kosten", -40], [None, "Summe", 70])
    assert "Summenzeile" in dp.pruefe(ws, bad, "ergebnis", None, [], False, HEUTE)["meldungen"][0]


def test_i4_month_13_is_rejected(ws, tmp_path):
    f = xlsx(tmp_path / "m.xlsx", ERG_H, ["2026-13", "Umsatz", 1])
    r = dp.pruefe(ws, f, "ergebnis", None, [], False, HEUTE)
    assert not r["ok"] and "ist kein Monat" in r["meldungen"][0]


def test_i9_csv_input_with_semicolons_and_cp1252(ws, tmp_path):
    f = tmp_path / "ergebnis.csv"
    f.write_bytes("Monat;Position;Ist_EUR\n09.2026;Umsatz Gewährleistung;1.234,56\n".encode("cp1252"))
    r = dp.pruefe(ws, f, "ergebnis", None, [], False, HEUTE)
    assert r["ok"], r["meldungen"]
    assert r["summe"] == 1234.56


def test_i8_best_sheet_is_chosen_and_reported(ws, tmp_path):
    f = xlsx(tmp_path / "s.xlsx", sheets=[("Notizen", [["Bitte beachten"]]),
                                        ("Daten", [ERG_H, ["2026-09", "Umsatz", 5]])])
    r = dp.pruefe(ws, f, "ergebnis", None, [], False, HEUTE)
    assert r["ok"] and r["blatt"] == "Daten"


def test_i5_i8_csv_has_source_columns_and_neutralised_formulas(ws):
    f = xlsx(ws / "00_Eingang" / "e.xlsx", ERG_H, ["2026-09", "+Zuschlag Notdienst", 5])
    r = dp.pruefe(ws, f, "ergebnis", None, [], True, HEUTE)
    assert r["ok"], r["meldungen"]
    with (ws / r["ziel"]).open(encoding="utf-8-sig") as fh:
        row = next(csv.DictReader(fh))
    assert (row["Position"], row["_quelle_zeile"], row["_quelle_blatt"]) == ("'+Zuschlag Notdienst", "2", "Sheet")


def test_i7_failed_move_leaves_nothing_half_done(ws, monkeypatch):
    f = xlsx(ws / "00_Eingang" / "e.xlsx", ERG_H, ["2026-09", "Umsatz", 5])

    def locked(*_):
        raise PermissionError("in use")

    monkeypatch.setattr(dp.shutil, "move", locked)
    r = dp.pruefe(ws, f, "ergebnis", None, [], True, HEUTE)
    assert not r["ok"] and "schließen" in r["meldungen"][0]
    assert f.exists() and not (ws / "07_Daten" / "ergebnis_2026-09.csv").exists()


@pytest.mark.parametrize("extra,msg", [(["--kontrollsumme", "ca. 100"], "ist keine Zahl"),
                                       (["--zuordnung", "Pos"], "Spalte=Zielspalte"),
                                       (["--vorlage", "gibtsnicht"], "Aufruf fehlerhaft")])
def test_i2_bad_cli_inputs_return_json(capsys, ws, tmp_path, extra, msg):
    f = xlsx(tmp_path / "e.xlsx", ERG_H, ["2026-09", "U", 1])
    args = ["--ws", str(ws), "--datei", str(f)] + (extra if "--vorlage" in extra else ["--vorlage", "ergebnis", *extra])
    code, out = run(capsys, dp.main, *args)
    assert code == 1 and msg in json.dumps(out, ensure_ascii=False)


def test_i2_missing_file_corrupt_mapping_and_xls_return_json(capsys, ws, tmp_path):
    code, out = run(capsys, dp.main, "--ws", str(ws), "--datei", str(tmp_path / "fehlt.xlsx"), "--vorlage", "ergebnis")
    assert code == 1 and "nicht gefunden" in out["meldungen"][0]
    (ws / "Unternehmen").mkdir()
    (ws / "Unternehmen" / "datenzuordnung.json").write_text("{kaputt", encoding="utf-8")
    f = xlsx(tmp_path / "e.xlsx", ERG_H, ["2026-09", "U", 1])
    code, out = run(capsys, dp.main, "--ws", str(ws), "--datei", str(f), "--vorlage", "ergebnis")
    assert code == 1 and "datenzuordnung.json ist beschädigt" in out["meldungen"][0]
    x = tmp_path / "alt.xls"
    x.write_bytes(b"\xd0\xcf\x11\xe0")
    code, out = run(capsys, dp.main, "--ws", str(ws), "--datei", str(x), "--vorlage", "ergebnis")
    assert code == 1 and ".xls" in out["meldungen"][0]


@pytest.fixture
def daten_gen(tmp_path):
    from generate_beispiel import generate
    generate(tmp_path / "gen")
    return tmp_path / "gen"
