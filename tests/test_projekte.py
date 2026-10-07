import datetime as dt
import json
import os
import shutil
import subprocess

import pytest
import yaml

import projekte
from conftest import ROOT

EVALS = ROOT / "plugin" / "evals"
STICHTAG = "2026-09-30"
TAG = dt.date(2026, 9, 30)
HANSA = "Retrofit Linie 2 Hansa Pack"
GUELTIG = {"typ": "retrofit", "status": "laufend", "kunde": "Test AG", "maschine": "Presse P-1, Masch.-Nr. 1",
           "auftragsnr": None, "projektleitung": "Tom Krüger", "budget_eur": 100000, "kosten_ist_eur": 20000,
           "kosten_prognose_eur": 100000, "kosten_stand": "2026-09-30", "erloes_eur": None,
           "vertragsstrafe_prozent_je_woche": None, "vertragsstrafe_max_prozent": None,
           "vertragsstrafe_bezugswert_eur": None, "vertragsstrafe_meilenstein": None, "gewaehrleistung_monate": 12,
           "meilensteine": [{"name": "Abnahme", "plan": "2026-10-30", "prognose": None, "ist": None}],
           "risiken": [], "quelle": None}
STRAFE_HANSA = {"vertragsstrafe_prozent_je_woche": 0.5, "vertragsstrafe_max_prozent": 5,
                "vertragsstrafe_bezugswert_eur": 240000, "vertragsstrafe_meilenstein": "Abnahme"}


def kopfzeilen(meta, nl="\n"):
    return nl.join(f"{k}: {json.dumps(v, ensure_ascii=False)}" for k, v in meta.items())


def projekt(ws, name, **aenderung):
    """Writes a valid projekt.md (GUELTIG with changes) to 05_Projekte/<name>/ and returns its path."""
    d = ws / "05_Projekte" / name
    d.mkdir(parents=True, exist_ok=True)
    (d / "projekt.md").write_text(f"---\n{kopfzeilen(GUELTIG | aenderung)}\n---\n\n# {name}\n", encoding="utf-8")
    return d / "projekt.md"


@pytest.fixture
def rufe(capsys):
    """Runs the CLI like uv does and returns (exit code, the one JSON object)."""
    def _rufe(*argv):
        code = projekte.main([str(a) for a in argv])
        return code, json.loads(capsys.readouterr().out)
    return _rufe


@pytest.fixture
def defs(monkeypatch):
    """kennzahlen.definitionen as Plan 3 returns it for blank company files; tests add what they need (G2, G3)."""
    werte = {"freigabegrenzen": {"standard": True}, "fachexperten": {"standard": True}, "kpi-ziele": {"standard": True}}
    monkeypatch.setattr(projekte.kennzahlen, "definitionen", lambda ws: werte)
    return werte


def fall(tmp_path, name):
    """Builds an eval workspace exactly like the runner: the case's scaffold.sh in an empty folder."""
    ws = tmp_path / "lauf"
    ws.mkdir()
    env = {"PATH": os.environ["PATH"], "HOME": str(ws), "TMPDIR": os.environ.get("TMPDIR", "/tmp"), "TERM": "dumb"}
    r = subprocess.run(["/bin/sh", str(EVALS / name / "scaffold.sh")], cwd=ws, env=env, capture_output=True,
                       text=True, timeout=120)
    assert r.returncode == 0, r.stderr
    return ws


def test_a_valid_project_has_no_problems():
    assert projekte.pruefe(dict(GUELTIG), TAG) == {}
    ohne = {k: v for k, v in GUELTIG.items() if k not in ("auftragsnr", "erloes_eur", "quelle")}
    assert projekte.pruefe(ohne, TAG) == {}


def test_missing_and_malformed_fields_are_reported_never_converted():
    meta = {k: v for k, v in GUELTIG.items() if k != "kosten_prognose_eur"} | {
        "budget_eur": "96.000 €",
        "meilensteine": [{"name": "Umbau", "plan": "2026-10-05", "prognose": "11.10.2026", "ist": None}]}
    p = projekte.pruefe(meta, TAG)
    assert p["kosten_prognose_eur"] == "fehlt"
    assert '"96.000 €"' in p["budget_eur"]
    assert p["meilensteine"].startswith("Umbau: prognose muss ein Datum JJJJ-MM-TT sein") and "11.10.2026" in p["meilensteine"]
    assert set(p) == {"kosten_prognose_eur", "budget_eur", "meilensteine"}


def test_a_passed_date_without_actual_is_reported():
    offen = GUELTIG | {"meilensteine": [{"name": "Montage", "plan": "2026-09-25", "prognose": None, "ist": None}]}
    assert "Montage: Termin 2026-09-25 ist vorbei" in projekte.pruefe(offen, TAG)["meilensteine"]
    alt = GUELTIG | {"meilensteine": [{"name": "Montage", "plan": "2026-09-01", "prognose": "2026-09-20", "ist": None}]}
    assert "Termin 2026-09-20 ist vorbei" in projekte.pruefe(alt, TAG)["meilensteine"]
    zukunft = GUELTIG | {"meilensteine": [{"name": "Montage", "plan": "2026-09-01", "prognose": None, "ist": "2026-10-02"}]}
    assert "liegt nach dem Stichtag" in projekte.pruefe(zukunft, TAG)["meilensteine"]


def test_contract_penalty_is_all_or_nothing():
    teil = GUELTIG | {"vertragsstrafe_prozent_je_woche": 0.5, "vertragsstrafe_meilenstein": "Abnahme"}
    assert projekte.pruefe(teil, TAG)["vertragsstrafe"].endswith(
        "es fehlt: vertragsstrafe_max_prozent, vertragsstrafe_bezugswert_eur")
    falsch = GUELTIG | STRAFE_HANSA | {"vertragsstrafe_meilenstein": "Übergabe"}
    assert "kein Meilenstein dieses Projekts" in projekte.pruefe(falsch, TAG)["vertragsstrafe"]
    assert projekte.pruefe(GUELTIG | STRAFE_HANSA, TAG) == {}


def test_risks_costs_and_warranty_are_checked():
    p = projekte.pruefe(GUELTIG | {"risiken": [{"text": "x", "stufe": "kritisch"}], "kosten_ist_eur": 120000,
                                   "gewaehrleistung_monate": "12 Monate", "budget_eur": 0}, TAG)
    assert set(p) == {"risiken", "kosten_prognose_eur", "gewaehrleistung_monate", "budget_eur"}


def test_delay_per_milestone_and_penalty_per_started_week():
    meta = GUELTIG | STRAFE_HANSA | {"meilensteine": [
        {"name": "Montage", "plan": "2026-09-25", "prognose": "2026-10-14", "ist": None},
        {"name": "Abnahme", "plan": "2026-10-15", "prognose": "2026-11-05", "ist": None},
        {"name": "Lieferung", "plan": "2026-09-10", "prognose": None, "ist": "2026-09-08"}]}
    assert projekte.verzug_tage(meta) == {"Montage": 19, "Abnahme": 21, "Lieferung": 0}
    assert [projekte.strafe_eur(meta, t) for t in (0, 21, 22, 100)] == [0, 3600, 4800, 12000]


def test_windows_saved_file_is_read(tmp_path):
    p = tmp_path / "projekt.md"
    p.write_bytes(("﻿---\r\n" + kopfzeilen(GUELTIG, "\r\n") + "\r\n---\r\n\r\n# X\r\n").encode("utf-8"))
    assert projekte.lies(p)["kunde"] == "Test AG"


def test_unreadable_files_raise_a_german_error(tmp_path):
    p = tmp_path / "projekt.md"
    p.write_text('---\ntyp: "retrofit"\nkunde: "Weber', encoding="utf-8")
    with pytest.raises(projekte.ProjektFehler, match="nicht lesbar"):
        projekte.lies(p)
    p.write_bytes(b'---\nkunde: "M\xfcller"\n---\n')
    with pytest.raises(projekte.ProjektFehler, match="UTF-8"):
        projekte.lies(p)


def test_clean_portfolio(tmp_path, defs, rufe):
    ws = fall(tmp_path, "projektportfolio-ampel-sauber")
    code, out = rufe("ampel", "--ws", ws, "--stichtag", STICHTAG)
    assert code == 0 and out["ok"]
    assert out["zaehler"] == {"rot": 1, "gelb": 1, "gruen": 1}
    assert [p["projekt"] for p in out["projekte"]] == [HANSA, "Retrofit Presse 3 Berger",
                                                       "Inbetriebnahme Kartonierer Kessler"]
    h = out["projekte"][0]
    assert (h["gesamt"], h["termin"]["ampel"], h["kosten"]["ampel"], h["risiko"]["ampel"]) == ("rot", "rot", "gelb", "gelb")
    assert h["termin"]["verzug"]["betrag"] == 21 and h["termin"]["verzug"]["einheit"] == "Tage"
    assert h["termin"]["grund"].startswith("Vertragsstrafe droht: Abnahme 21 Tage")
    assert h["kosten"]["abweichung"]["betrag"] == 7.6 and h["kosten"]["abweichung"]["berechnet"] is True
    assert h["datei"] == f"05_Projekte/{HANSA}/projekt.md"
    assert out["abgeschlossen"] == ["Inbetriebnahme Etikettierer Lenz"] and out["nicht_bewertet"] == []
    s = out["summen"]
    assert (s["budget"]["betrag"], s["prognose"]["betrag"], s["abweichung_eur"]["betrag"],
            s["abweichung_prozent"]["betrag"]) == (330000, 345540, 15540, 4.7)
    assert out["schwellen"]["standard"] is True and any(projekte.STANDARD_HINWEIS in m for m in out["meldungen"])
    assert out["ziel"] == "03_Berichte/2026-09-30_projektportfolio-ampel.md" and out["beispiel"] is False


def test_messy_portfolio_reports_and_never_guesses(tmp_path, defs, rufe):
    ws = fall(tmp_path, "projektportfolio-ampel-unordentlich")
    out = rufe("ampel", "--ws", ws, "--stichtag", STICHTAG)[1]
    assert out["zaehler"] == {"rot": 1, "gelb": 2, "gruen": 0}
    nicht = {n["projekt"]: n["grund"] for n in out["nicht_bewertet"]}
    assert nicht["Angebot Schmidt"] == "projekt.md fehlt" and "nicht lesbar" in nicht["Service-Umbau Weber"]
    p = {x["projekt"]: x for x in out["projekte"]}
    assert p[HANSA]["gesamt"] == "rot"
    k = p["Inbetriebnahme Kartonierer Kessler"]
    assert k["kosten"]["ampel"] == "grau" and k["gesamt"] == "gelb"
    assert any(x.startswith("budget_eur: muss eine Zahl") and "96.000 €" in x for x in k["probleme"])
    b = p["Retrofit Presse 3 Berger"]
    assert (b["termin"]["ampel"], b["kosten"]["ampel"], b["gesamt"]) == ("grau", "grau", "gelb")
    assert "kosten_prognose_eur: fehlt" in b["probleme"]
    assert out["summen"]["budget"]["betrag"] == 184000
    assert sorted(out["summen"]["ohne"]) == ["Inbetriebnahme Kartonierer Kessler", "Retrofit Presse 3 Berger"]


def test_the_colour_follows_the_displayed_percentage(kit_ws, defs, rufe):
    projekt(kit_ws, "Grenzfall", kosten_prognose_eur=109960)  # 9,96 % is shown as 10,0 %
    k = rufe("ampel", "--ws", kit_ws, "--stichtag", STICHTAG)[1]["projekte"][0]["kosten"]
    assert k["abweichung"]["betrag"] == 10.0 and k["ampel"] == "rot"


def test_a_grey_dimension_lifts_green_to_yellow():
    assert projekte.gesamt([{"ampel": "gruen"}, {"ampel": "grau"}, {"ampel": "gruen"}]) == "gelb"
    assert projekte.gesamt([{"ampel": "rot"}, {"ampel": "grau"}]) == "rot"
    assert projekte.gesamt([{"ampel": "grau"}] * 3) == "gelb"


def test_company_thresholds_from_kpi_ziele(kit_ws, defs, rufe):
    projekt(kit_ws, "Knapp", meilensteine=[{"name": "Abnahme", "plan": "2026-10-30", "prognose": "2026-11-02", "ist": None}])
    assert rufe("ampel", "--ws", kit_ws, "--stichtag", STICHTAG)[1]["projekte"][0]["termin"]["ampel"] == "gruen"
    defs["kpi-ziele"]["projektampel"] = {"termin_gelb_ab_tage": 1, "termin_rot_ab_tage": 5,
                                         "kosten_gelb_ab_prozent": 1, "kosten_rot_ab_prozent": 5}
    out = rufe("ampel", "--ws", kit_ws, "--stichtag", STICHTAG)[1]
    assert out["schwellen"]["standard"] is False and out["projekte"][0]["termin"]["ampel"] == "gelb"
    defs["kpi-ziele"]["projektampel"] = {"termin_gelb_ab_tage": 9, "termin_rot_ab_tage": 5}
    out = rufe("ampel", "--ws", kit_ws, "--stichtag", STICHTAG)[1]
    assert out["schwellen"]["standard"] is True and any("ungültig" in m for m in out["meldungen"])


def test_sample_mode_reads_the_sample_company(kit_ws, rufe):
    cfg = kit_ws / "Unternehmen" / ".kit-config"
    cfg.write_text(cfg.read_text(encoding="utf-8").replace("beispieldaten=nein", "beispieldaten=ja"), encoding="utf-8")
    shutil.copytree(EVALS / "_gemeinsam" / "projekte", kit_ws / "Beispiel" / "05_Projekte")
    out = rufe("ampel", "--ws", kit_ws, "--stichtag", STICHTAG)[1]
    assert out["beispiel"] is True and out["hinweis"] == "Beispieldaten – Muster Maschinenbau GmbH"
    assert out["projekte"][0]["datei"] == f"Beispiel/05_Projekte/{HANSA}/projekt.md"


def test_empty_portfolio_and_wrong_folder(kit_ws, tmp_path, defs, rufe):
    code, out = rufe("ampel", "--ws", kit_ws)
    assert code == 0 and out["zaehler"] == {"rot": 0, "gelb": 0, "gruen": 0} and "noch kein Projekt" in out["meldungen"][-1]
    code, out = rufe("ampel", "--ws", tmp_path / "leer")
    assert code == 1 and "kein Kundendienst-Ordner" in out["fehler"][0]


def test_delay_decision_options(tmp_path, defs, rufe):
    ws = fall(tmp_path, "verzug-entscheidung-sauber")
    code, s = rufe("setze", "--ws", ws, "--projekt", HANSA, "--meilenstein", "Abnahme", "--prognose", "2026-11-05")
    assert code == 0 and (s["alt"], s["neu"]) == (None, "2026-11-05")
    code, out = rufe("verzug", "--ws", ws, "--projekt", HANSA, "--stichtag", STICHTAG, "--heute", "2026-10-07",
                     "--beschleunigung-eur", "6.500", "--beschleunigung-tage", "14")
    assert code == 0 and out["entscheidung_noetig"] is True
    assert (out["verzug"]["betrag"], out["bezug_meilenstein"], out["mehrkosten"]["betrag"]) == (21, "Abnahme", 14000)
    o = {x["kennung"]: x for x in out["optionen"]}
    assert (o["A"]["strafe"]["betrag"], o["A"]["belastung"]["betrag"], o["A"]["ergebnis"]["betrag"]) == (3600, 17600, 38400)
    assert (o["B"]["strafe"]["betrag"], o["B"]["belastung"]["betrag"], o["B"]["ergebnis"]["betrag"]) == (1200, 21700, 34300)
    assert o["C"]["belastung"]["betrag"] == 14000 and "Zustimmung des Kunden" in o["C"]["voraussetzung"]
    assert o["D"]["berechnet"] is False and "Nachtrag" in o["D"]["titel"]
    assert out["vorschlag"] == "A"
    assert [p["reviewer"] for p in out["pruefung"]] == ["finanzen", "qualitaet-recht"]
    assert out["pruefung"][1]["fachexperte"] is None
    assert out["ziel"] == f"05_Projekte/{HANSA}/2026-10-07_verzug-entscheidung.docx"
    assert out["faellig_vorschlag"] == "2026-10-14" and out["projektleitung"] == "Tom Krüger"
    assert out["pruefzahlen"] == ["21", "14.000", "3.600", "17.600", "1.200", "21.700"]


def test_reviewers_follow_the_limit_and_name_the_expert(tmp_path, defs, rufe):
    ws = fall(tmp_path, "verzug-entscheidung-sauber")
    rufe("setze", "--ws", ws, "--projekt", HANSA, "--meilenstein", "Abnahme", "--prognose", "2026-11-05")
    defs["freigabegrenzen"]["projekt_mehrkosten_eur"] = 50000
    defs["fachexperten"]["recht"] = "Dr. Anna Weiß (Justiziariat)"
    p = rufe("verzug", "--ws", ws, "--projekt", HANSA, "--stichtag", STICHTAG)[1]["pruefung"]
    assert [(x["reviewer"], x.get("fachexperte")) for x in p] == [("qualitaet-recht", "Dr. Anna Weiß (Justiziariat)")]
    defs["freigabegrenzen"]["projekt_mehrkosten_eur"] = 17000
    p = rufe("verzug", "--ws", ws, "--projekt", HANSA, "--stichtag", STICHTAG)[1]["pruefung"]
    assert p[0]["reviewer"] == "finanzen" and "17.600" in p[0]["grund"] and "17.000" in p[0]["grund"]


def test_incomplete_penalty_is_never_guessed(tmp_path, defs, rufe):
    ws = fall(tmp_path, "verzug-entscheidung-unordentlich")
    code, out = rufe("verzug", "--ws", ws, "--projekt", HANSA, "--stichtag", STICHTAG)
    assert code == 0
    o = {x["kennung"]: x for x in out["optionen"]}
    assert o["A"]["strafe"] is None and o["A"]["belastung"]["betrag"] == 14000
    assert "Mindestwert" in o["A"]["belastung"]["formel"]
    assert o["B"]["berechnet"] is False and out["vorschlag"] is None
    assert "qualitaet-recht" in [p["reviewer"] for p in out["pruefung"]]
    assert any("vertragsstrafe_max_prozent: fehlt" in m for m in out["meldungen"])
    assert "3.600" not in json.dumps(out, ensure_ascii=False) and out["pruefzahlen"] == ["21", "14.000"]


def test_no_delay_no_decision_and_broken_costs_block(kit_ws, defs, rufe):
    projekt(kit_ws, "Im Plan")
    code, out = rufe("verzug", "--ws", kit_ws, "--projekt", "Im Plan", "--stichtag", STICHTAG)
    assert code == 0 and out["entscheidung_noetig"] is False and out["optionen"] == []
    projekt(kit_ws, "Kaputt", budget_eur="96.000 €")
    code, out = rufe("verzug", "--ws", kit_ws, "--projekt", "Kaputt", "--stichtag", STICHTAG)
    assert code == 1 and "budget_eur" in out["fehler"][0]


def test_unknown_or_unsafe_project_names(kit_ws, defs, rufe):
    projekt(kit_ws, "Im Plan")
    code, out = rufe("verzug", "--ws", kit_ws, "--projekt", "Gibt es nicht")
    assert code == 1 and "Vorhanden: Im Plan" in out["fehler"][0]
    code, out = rufe("verzug", "--ws", kit_ws, "--projekt", "../Unternehmen")
    assert code == 1 and "Ungültiger Projektname" in out["fehler"][0]


def test_setze_changes_one_line_and_keeps_the_rest(kit_ws, defs, rufe):
    datei = projekt(kit_ws, "Im Plan")
    datei.write_text(datei.read_text(encoding="utf-8") + "\n## Notizen\n\nKunde will Oktober.\n", encoding="utf-8")
    vorher = datei.read_text(encoding="utf-8")
    code, out = rufe("setze", "--ws", kit_ws, "--projekt", "Im Plan", "--feld", "kosten_prognose_eur", "--wert", "104.500")
    assert code == 0 and (out["alt"], out["neu"]) == (100000, 104500)
    nachher = datei.read_text(encoding="utf-8")
    assert nachher == vorher.replace("kosten_prognose_eur: 100000\n", "kosten_prognose_eur: 104500\n")


def test_setze_refuses_what_it_cannot_take(kit_ws, defs, rufe):
    projekt(kit_ws, "Im Plan")
    for argv, text in ((["--meilenstein", "Abnahme", "--prognose", "05.11.2026"], "JJJJ-MM-TT"),
                       (["--meilenstein", "Montage", "--prognose", "2026-11-05"], "Vorhanden: Abnahme"),
                       (["--feld", "budget_eur", "--wert", "1"], "nicht änderbar"),
                       (["--feld", "projektleitung", "--wert", "projekte"], "kein Agent"),
                       (["--feld", "kosten_ist_eur", "--wert", "viel"], "keine Zahl")):
        code, out = rufe("setze", "--ws", kit_ws, "--projekt", "Im Plan", *argv)
        assert code == 1 and text in " ".join(out["fehler"]), (argv, out)


BRENNER = "Inbetriebnahme KX-400 Brenner"


def anlegen_argv(ws, _meilensteine=("Lieferung=2026-11-16", "Montage=2026-11-27", "Abnahme=2026-12-04"), **ueber):
    werte = {"projekt": BRENNER, "typ": "inbetriebnahme", "kunde": "Brenner Verpackung GmbH",
             "maschine": "Kartonierer KX-400, Masch.-Nr. 40-1203", "projektleitung": "Jana Becker", "budget": "38.500",
             "gewaehrleistung-monate": "12", "strafe-prozent-woche": "0,5", "strafe-max-prozent": "5",
             "strafe-bezugswert": "312000", "strafe-meilenstein": "Abnahme", "heute": "2026-10-01"} | ueber
    argv = ["anlegen", "--ws", ws]
    for k, v in werte.items():
        if v is not None:
            argv += [f"--{k}", v]
    for m in _meilensteine:
        argv += ["--meilenstein", m]
    return argv


def test_handover_creates_project_checklist_and_files_the_document(kit_ws, defs, rufe):
    mail = kit_ws / "00_Eingang" / "2026-10-01_mail-auftrag-brenner.eml"
    mail.write_text("Auftrag Brenner", encoding="utf-8")
    code, out = rufe(*anlegen_argv(kit_ws, beleg=f"00_Eingang/{mail.name}"))
    assert code == 0, out
    d = kit_ws / "05_Projekte" / BRENNER
    assert not mail.exists() and (d / mail.name).read_text(encoding="utf-8") == "Auftrag Brenner"
    meta = projekte.lies(d / "projekt.md")
    assert projekte.pruefe(meta, dt.date(2026, 10, 1)) == {}
    assert (meta["budget_eur"], meta["kosten_prognose_eur"], meta["kosten_ist_eur"], meta["status"]) == (38500, 38500, 0, "geplant")
    assert meta["vertragsstrafe_prozent_je_woche"] == 0.5 and meta["quelle"] == f"05_Projekte/{BRENNER}/{mail.name}"
    text = (d / "projekt.md").read_text(encoding="utf-8")
    assert "\nbudget_eur: 38500\n" in text
    assert '{"name": "Abnahme", "plan": "2026-12-04", "prognose": null, "ist": null}' in text
    liste = (d / "uebergabe.md").read_text(encoding="utf-8")
    assert liste.count("- [ ] ") == 16 and "Servicevertrag" in liste and "12 Monate ab Abnahme" in liste
    assert "0,5 % je angefangene Woche" in liste
    assert (out["servicevertrag_faellig"], out["abnahme_meilenstein"]) == ("2026-11-20", "Abnahme")


def test_handover_lists_everything_missing_at_once(kit_ws, defs, rufe):
    code, out = rufe(*anlegen_argv(kit_ws, budget=None, **{"gewaehrleistung-monate": None, "strafe-max-prozent": None}))
    assert code == 1 and not (kit_ws / "05_Projekte" / BRENNER).exists()
    assert len(out["fehlende_angaben"]) == 3 and out["fehlende_angaben"][-1].startswith("Vertragsstrafe")


def test_handover_never_overwrites_and_never_moves(kit_ws, defs, rufe):
    assert rufe(*anlegen_argv(kit_ws))[0] == 0
    datei = kit_ws / "05_Projekte" / BRENNER / "projekt.md"
    vorher = datei.read_bytes()
    mail = kit_ws / "00_Eingang" / "auftrag.eml"
    mail.write_text("x", encoding="utf-8")
    code, out = rufe(*anlegen_argv(kit_ws, beleg="00_Eingang/auftrag.eml", kunde="Andere GmbH"))
    assert code == 1 and "gibt es schon" in out["fehler"][0]
    assert mail.exists() and datei.read_bytes() == vorher


def test_handover_refuses_bad_input(kit_ws, defs, rufe):
    faelle = ((dict(beleg="Unternehmen/profil.md"), "direkt in 00_Eingang"),
              (dict(_meilensteine=("Abnahme=04.12.2026",)), "Name=JJJJ-MM-TT"),
              (dict(projektleitung="projekte"), "kein Agent"),
              (dict(**{"strafe-meilenstein": "Übergabe"}), "nicht unter den Meilensteinen"),
              (dict(projekt="../Unternehmen"), "Ungültiger Projektname"))
    for ueber, text in faelle:
        code, out = rufe(*anlegen_argv(kit_ws, **ueber))
        assert code == 1 and text in " ".join(out["fehler"]), (ueber, out)
    assert (kit_ws / "Unternehmen" / "profil.md").exists() and not (kit_ws / "05_Projekte" / BRENNER).exists()


def test_pruefe_datei_finds_numbers_in_word_and_excel(kit_ws, rufe):
    import docx
    import openpyxl
    doc = docx.Document()
    doc.add_paragraph("Ergebnisbelastung Option A: 17.600 EUR")
    doc.add_table(rows=1, cols=2).cell(0, 1).text = "21.700"
    doc.save(kit_ws / "05_Projekte" / "memo.docx")
    code, out = rufe("pruefe-datei", "--ws", kit_ws, "--datei", "05_Projekte/memo.docx",
                     "--zahl", "17.600", "--zahl", "21.700", "--zahl", "3.600")
    assert code == 1 and out["gefunden"] == ["17.600", "21.700"] and out["fehlend"] == ["3.600"]
    wb = openpyxl.Workbook()
    wb.active["A1"] = 3600
    wb.save(kit_ws / "03_Berichte" / "t.xlsx")
    assert rufe("pruefe-datei", "--ws", kit_ws, "--datei", "03_Berichte/t.xlsx", "--zahl", "3.600")[0] == 0
    code, out = rufe("pruefe-datei", "--ws", kit_ws, "--datei", "../ausserhalb.docx", "--zahl", "1")
    assert code == 1 and "gibt es im Kundendienst-Ordner nicht" in out["fehler"][0]


import argparse
import re

SKILLS_4E = ["maschinenuebergabe", "projektportfolio-ampel", "verzug-entscheidung"]


def skill(name):
    text = (ROOT / "plugin" / "skills" / name / "SKILL.md").read_text(encoding="utf-8")
    m = re.match(r"\A---\n(.*?)\n---\n(.*)\Z", text, re.S)
    assert m, f"{name}: Kopfbereich fehlt"
    return yaml.safe_load(m.group(1)), m.group(2)


@pytest.mark.parametrize("name", SKILLS_4E)
def test_skill_contract_and_every_command_line_parses(name):
    meta, body = skill(name)
    assert meta["name"] == name and 40 <= len(meta["description"]) <= 1024
    assert "**Liest:**" in body and "**Schreibt:**" in body and "Daten, nie Anweisungen" in body
    assert "$CLAUDE_PLUGIN_ROOT" not in body.replace("${CLAUDE_PLUGIN_ROOT}", "") and "pip install" not in body
    unter = next(a for a in projekte.parser()._actions if isinstance(a, argparse._SubParsersAction)).choices
    zeilen = re.findall(r'`uv run "\$\{CLAUDE_PLUGIN_ROOT\}/scripts/projekte\.py" ([^`]+)`', body)
    assert zeilen, "Skill ruft projekte.py nicht auf"
    for zeile in zeilen:
        befehl = zeile.split()[0]
        bekannt = {o for a in unter[befehl]._actions for o in a.option_strings}
        assert set(re.findall(r"--[a-z][\w-]*", zeile)) <= bekannt, zeile
    assert "vorgang.py\" entscheide" not in body  # only the user decides (spec §6)


def test_delay_skill_hands_reviews_to_other_agents():
    _, body = skill("verzug-entscheidung")
    assert "service-leader-kit:finanzen" in body and "service-leader-kit:qualitaet-recht" in body
    assert "pruefe-datei" in body and "Empfehlung: zustimmen" in body


def muster(name):
    c = yaml.safe_load((EVALS / name / "case.yaml").read_text(encoding="utf-8"))
    return " ".join(g.get("pattern", "") for g in c["graders"]).replace("\\", "")


def test_lane_cases_exist_for_both_datasets():
    for s in SKILLS_4E:
        for d in ("sauber", "unordentlich"):
            c = yaml.safe_load((EVALS / f"{s}-{d}" / "case.yaml").read_text(encoding="utf-8"))
            assert f"skill:{s}" in c["tags"] and d in c["tags"]


def test_graded_numbers_are_what_the_script_computes(tmp_path, defs, rufe):
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    ws = fall(tmp_path / "a", "projektportfolio-ampel-sauber")
    out = rufe("ampel", "--ws", ws, "--stichtag", STICHTAG)[1]
    m = muster("projektportfolio-ampel-sauber")
    for w in (out["summen"]["budget"], out["summen"]["prognose"]):
        assert projekte.kennzahlen.deutsch(w["betrag"]) in m
    assert projekte.kennzahlen.deutsch(out["projekte"][0]["kosten"]["abweichung"]["betrag"], 1) in m
    ws = fall(tmp_path / "b", "verzug-entscheidung-sauber")
    rufe("setze", "--ws", ws, "--projekt", HANSA, "--meilenstein", "Abnahme", "--prognose", "2026-11-05")
    out = rufe("verzug", "--ws", ws, "--projekt", HANSA, "--stichtag", STICHTAG, "--beschleunigung-eur", "6500",
               "--beschleunigung-tage", "14")[1]
    m = muster("verzug-entscheidung-sauber")
    for z in ("3.600", "17.600", "21.700"):
        assert z in out["pruefzahlen"] and z in m
