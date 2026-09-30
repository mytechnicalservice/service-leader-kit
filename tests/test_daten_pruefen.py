import datetime as dt
import json

import pytest

from daten_pruefen import datum, pruefe, zahl
from generate_beispiel import generate

HEUTE = "2026-09-30"


@pytest.fixture
def daten(tmp_path):
    return generate(tmp_path / "gen"), tmp_path / "gen"


@pytest.mark.parametrize("raw,exp", [("1.234,56", 1234.56), ("1234,5", 1234.5), ("12.5", 12.5), ("1.234", 1234.0),
                                     (7, 7.0), ("-1.000.000,00", -1000000.0), (" 3 € ", 3.0)])
def test_zahl_accepts_german_and_plain(raw, exp):
    assert zahl(raw) == exp


@pytest.mark.parametrize("raw", ["abc", "1,2,3", "", None, True])
def test_zahl_rejects_garbage(raw):
    with pytest.raises(ValueError):
        zahl(raw)


def test_datum_formats():
    assert datum("30.09.2026") == datum("2026-09-30") == datum(dt.datetime(2026, 9, 30, 0, 0)) == dt.date(2026, 9, 30)


def test_clean_import_writes_csv_moves_original_and_blocks_second_import(ws, daten):
    exp, gen = daten
    src = ws / "00_Eingang" / "auftraege_2026-09.xlsx"
    src.parent.mkdir()
    src.write_bytes((gen / "beispiel" / "00_Eingang" / src.name).read_bytes())
    r = pruefe(ws, src, "auftraege", None, [], True, HEUTE)
    assert r["ok"] and (r["periode"], r["summe"], r["zeilen"]) == ("2026-09", exp["auftraege_umsatz"], exp["auftraege_zeilen"])
    assert (ws / "07_Daten" / "auftraege_2026-09.csv").exists() and not src.exists()
    assert (ws / "07_Daten" / "original" / "auftraege_2026-09__auftraege_2026-09.xlsx").exists()
    dup = ws / "00_Eingang" / "auftraege_2026-09 (1).xlsx"
    dup.write_bytes((gen / "beispiel" / "00_Eingang" / "auftraege_2026-09.xlsx").read_bytes())
    r2 = pruefe(ws, dup, "auftraege", None, [], True, HEUTE)
    assert not r2["ok"] and "September 2026 ist bereits importiert" in r2["meldungen"][0] and dup.exists()


def test_messy_orders_map_synonyms_parse_german_numbers_and_drop_duplicates(ws, daten):
    exp, gen = daten
    r = pruefe(ws, gen / "beispiel-unordentlich" / "00_Eingang" / "auftraege_2026-09.xlsx", "auftraege", None, [], False, HEUTE)
    assert r["ok"], r["meldungen"]
    assert r["zuordnung"]["Auftragswert"] == "Umsatz_EUR" and r["zuordnung"]["Std."] == "Stunden"
    assert (r["duplikate"], r["zeilen"], r["summe"]) == (3, exp["auftraege_zeilen"], exp["auftraege_umsatz"])


def test_missing_mandatory_column_is_named(ws, daten):
    _, gen = daten
    r = pruefe(ws, gen / "beispiel-unordentlich" / "00_Eingang" / "kapazitaet_2026-09.xlsx", "kapazitaet", None, [], False, HEUTE)
    assert not r["ok"] and "Spalte 'Ist_Stunden' fehlt" in r["meldungen"]


def test_control_total_mismatch_reports_both_sums(ws, daten):
    exp, gen = daten
    r = pruefe(ws, gen / "beispiel" / "00_Eingang" / "auftraege_2026-09.xlsx", "auftraege", "1.000,00", [], False, HEUTE)
    assert not r["ok"] and "Kontrollsumme 1.000,00" in r["meldungen"][0]


def test_bad_cell_names_row_and_column(ws, tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["Monat", "Position", "Ist_EUR"])
    wb.active.append(["2026-09", "Umsatz Service", "viel"])
    f = tmp_path / "ergebnis_kaputt.xlsx"
    wb.save(f)
    r = pruefe(ws, f, "ergebnis", None, [], False, HEUTE)
    assert not r["ok"] and "Zeile 2, Spalte 'Ist_EUR': 'viel' ist keine Zahl" in r["meldungen"]


def test_user_mapping_is_saved_and_reused(ws, tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["Periode", "Pos.", "Betrag Ist"])
    wb.active.append(["09.2026", "Umsatz Service", "100,00"])
    f = ws / "00_Eingang" / "erg.xlsx"
    f.parent.mkdir()
    wb.save(f)
    r = pruefe(ws, f, "ergebnis", None, ["Pos.=Position", "Betrag Ist=Ist_EUR"], True, HEUTE)
    assert r["ok"], r["meldungen"]
    saved = json.loads((ws / "Unternehmen" / "datenzuordnung.json").read_text(encoding="utf-8"))
    assert saved["ergebnis"] == {"Pos.": "Position", "Betrag Ist": "Ist_EUR"}
    wb2 = openpyxl.Workbook()
    wb2.active.append(["Periode", "Pos.", "Betrag Ist"])
    wb2.active.append(["10.2026", "Umsatz Service", "50,00"])
    f2 = ws / "00_Eingang" / "erg2.xlsx"
    wb2.save(f2)
    assert pruefe(ws, f2, "ergebnis", None, [], False, HEUTE)["ok"]
