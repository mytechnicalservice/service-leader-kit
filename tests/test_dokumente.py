import json
import shutil
import zipfile

import pytest

import layout
from conftest import ROOT

BEISPIEL = ROOT / "plugin" / "beispiel"
OHNE = ROOT / "plugin" / "evals" / "_gemeinsam" / "dateien" / "Firmenvorlage.pptx"


def test_master_check_accepts_the_sample_and_rejects_a_blank_template():
    gut = layout.pruefe_master(BEISPIEL / "Unternehmen" / "vorlagen" / "master.pptx")
    assert gut["echt"] and gut["vorschlag"]["inhalt"] == "Titel und Inhalt" and gut["meldung"] is None
    schlecht = layout.pruefe_master(OHNE)
    assert not schlecht["echt"] and ".potx" in schlecht["meldung"]


def test_sample_layout_map_names_real_layouts():
    v = BEISPIEL / "Unternehmen" / "vorlagen"
    karte = layout.lies_layout_map(v / "layout-map.md")
    namen = {l["name"] for l in layout.pruefe_master(v / "master.pptx")["layouts"]}
    assert karte["bestaetigt"] and set(karte["zuordnung"]) == set(layout.FOLIENARTEN)
    assert set(karte["zuordnung"].values()) <= namen
    assert "Muster Maschinenbau GmbH" in layout.datei_text(v / "briefkopf.docx")


def test_potx_opens_like_a_pptx(tmp_path):
    potx = tmp_path / "vorlage.potx"
    with zipfile.ZipFile(BEISPIEL / "Unternehmen" / "vorlagen" / "master.pptx") as zin, \
            zipfile.ZipFile(potx, "w") as zout:
        for item in zin.infolist():
            daten = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                daten = daten.replace(b"presentation.main+xml", b"template.main+xml")
            zout.writestr(item, daten)
    assert layout.pruefe_master(potx)["echt"]


def test_firmenlayout_in_sample_mode_and_when_files_are_missing(kit_ws):
    leer = layout.firmenlayout(kit_ws)
    assert leer["master"] is None and set(leer["fehlt"]) == {"master.pptx", "briefkopf.docx", "layout-map.md"}
    cfg = kit_ws / "Unternehmen" / ".kit-config"
    cfg.write_text(cfg.read_text(encoding="utf-8").replace("beispieldaten=nein", "beispieldaten=ja"), encoding="utf-8")
    shutil.copytree(BEISPIEL, kit_ws / "Beispiel")
    f = layout.firmenlayout(kit_ws)
    assert f["master"] == "Beispiel/Unternehmen/vorlagen/master.pptx" and f["fehlt"] == [] and f["beispiel"]
    assert f["layout_map"]["zuordnung"]["titel"] == "Titelfolie"


def test_ablage_never_overwrites_and_stays_in_the_output_folders(kit_ws):
    p = layout.ablage(kit_ws, "03_Berichte", "Management-Bericht September", "pptx", "2026-10-06")
    assert p == kit_ws / "03_Berichte" / "2026-10-06_management-bericht-september.pptx"
    p.write_bytes(b"x")
    assert layout.ablage(kit_ws, "03_Berichte", "Management-Bericht September", "pptx", "2026-10-06").name.endswith("_v2.pptx")
    for ordner in ("01_Vorgaenge", "Unternehmen", "03_Berichte/../Unternehmen"):
        with pytest.raises(layout.LayoutFehler):
            layout.ablage(kit_ws, ordner, "x", "md", "2026-10-06")


def test_pruefe_datei_reads_docx_pptx_xlsx_and_finds_totals(kit_ws, capsys):
    from docx import Document
    from pptx import Presentation

    doc = Document()
    doc.add_paragraph("Entscheidungsvorlage für die Kulanz und der Empfehlung: Betrag 8.400 EUR")
    doc.save(kit_ws / "04_Angebote" / "v.docx")
    prs = Presentation(str(BEISPIEL / "Unternehmen" / "vorlagen" / "master.pptx"))
    folie = prs.slides.add_slide(prs.slide_layouts[1])
    folie.shapes.title.text = "Umsatz September"
    folie.placeholders[1].text = "Der Umsatz für den Monat und die Marge: 434.805 EUR"
    prs.save(kit_ws / "03_Berichte" / "b.pptx")
    r = layout.pruefe_datei(kit_ws, kit_ws / "04_Angebote" / "v.docx", ["8.400"])
    assert r["ok"] and r["pruefung"] == "DATEI-OK 04_Angebote/v.docx 8.400" and r["version"].startswith("04_Angebote/v.docx (sha256 ")
    assert layout.main(["pruefe-datei", "--ws", str(kit_ws), "--datei", "03_Berichte/b.pptx", "--erwarte", "434.805",
                        "--erwarte", "999.999"]) == 1
    out = json.loads(capsys.readouterr().out)
    assert out["fehlend"] == ["999.999"] and out["pruefung"].startswith("DATEI-FEHLER")
    (kit_ws / "03_Berichte" / "kaputt.docx").write_bytes(b"kein zip")
    with pytest.raises(layout.LayoutFehler, match="lässt sich nicht öffnen"):
        layout.pruefe_datei(kit_ws, kit_ws / "03_Berichte" / "kaputt.docx", [])
