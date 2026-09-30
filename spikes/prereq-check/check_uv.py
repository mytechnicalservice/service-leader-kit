# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl==3.1.5", "python-docx==1.1.2", "python-pptx==1.0.2"]
# ///
"""Prerequisite check: can uv fetch Python + the Office libraries and write all three file types?"""
import json
import platform
import sys
import tempfile
from pathlib import Path

import docx
import openpyxl
import pptx


def main() -> int:
    out = Path(tempfile.mkdtemp(prefix="slk-check-"))
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Umsatz"
    wb.active["B1"] = 1234.56
    wb.save(out / "test.xlsx")
    d = docx.Document()
    d.add_paragraph("Service-Monatsbericht – Prüfung")
    d.save(out / "test.docx")
    p = pptx.Presentation()
    p.slides.add_slide(p.slide_layouts[0]).shapes.title.text = "Prüfung"
    p.save(out / "test.pptx")
    files = sorted(f.name for f in out.iterdir())
    print(json.dumps({"ok": files == ["test.docx", "test.pptx", "test.xlsx"], "python": sys.version.split()[0],
                      "system": platform.platform(), "ordner": str(out), "dateien": files}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
