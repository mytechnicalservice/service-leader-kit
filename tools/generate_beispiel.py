# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl==3.1.5", "python-docx==1.1.2", "python-pptx==1.0.2"]
# ///
"""Generates the fictional sample company 'Muster Maschinenbau GmbH' (decision D14): one business year
2025-10 … 2026-09 as validated imports in 07_Daten/ (imported through daten_pruefen, exactly like a user's data),
the September exports in 00_Eingang/, filled Unternehmen/ files, layout files, contracts, price list, projects,
budget, mails; the messy variant; and the expected values for the eval graders."""
from __future__ import annotations

import datetime as dt
import json
import random
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "plugin" / "scripts"))
sys.path.insert(0, str(ROOT / "tools"))

import daten_pruefen  # noqa: E402
from beispiel_dateien import (MAIL_ANFRAGE, MAIL_ESKALATION, MAIL_FLUIDTEC, MAIL_HYDRAULIK,  # noqa: E402
                              MAIL_INJECTION, MAIL_REKLAMATION, briefkopf, master_ohne_layouts, master_pptx,
                              projekte, scan_pdf, vertrag_md)

QUELLE = ROOT / "tools" / "beispiel_quelle"
MONATE = ["2025-10", "2025-11", "2025-12", "2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06",
          "2026-07", "2026-08", "2026-09"]
SEPT = "2026-09"
# name, team, contract level (None = no contract), machines, machines under contract, contract end
KUNDEN = [("Hansa Pack AG", "Nord", "Premium", 16, 14, "2026-12-31"),
          ("Alpen Getränke AG", "Süd", "Premium", 11, 10, "2027-06-30"),
          ("Müller GmbH", "Süd", "Standard", 14, 12, "2027-03-31"),
          ("Nordmetall GmbH", "Nord", "Standard", 9, 8, "2026-12-31"),
          ("Rhein Chemie Service", "West", "Premium", 10, 9, "2027-09-30"),
          ("Weber Kunststofftechnik", "West", "Basis", 12, 11, "2027-01-31"),
          ("Bäckerei Kraus KG", "Süd", "Basis", 7, 6, "2026-11-30"),
          ("Fränkische Molkerei eG", "Süd", "Standard", 13, 12, "2027-04-30"),
          ("Schreiner & Söhne", "West", None, 5, 0, None),
          ("Ostsee Fisch GmbH", "Nord", None, 8, 0, None),
          ("Bergmann Pharma GmbH", "West", None, 9, 0, None),
          ("Elbe Papier AG", "Nord", None, 10, 0, None),
          ("Kessler Feinkost", "Süd", None, 6, 0, None),
          ("Lippe Verpackung GmbH", "Nord", None, 9, 0, None)]
STUFEN = {"Basis": 4800, "Standard": 7800, "Premium": 13200}  # EUR per machine and year
TYPEN = ["MM-400", "MM-600", "MM-800"]
TEAMS = {"Nord": 6, "Süd": 5, "West": 4}  # technicians; Süd loses one from 2026-07
FAHRT = {"Nord": 110, "Süd": 140, "West": 95}
SATZ = {2025: 122, 2026: 128}  # EUR per technician hour
IBN_SATZ = 95  # internal transfer price for commissioning (Vertrieb pays)
SAISON = {"10": 1.0, "11": 1.02, "12": 0.82, "01": 0.95, "02": 0.98, "03": 1.08, "04": 1.0, "05": 1.03,
          "06": 1.05, "07": 0.97, "08": 0.85, "09": 1.06}
VERFUEGBAR = {"12": 0.8, "08": 0.72}  # holiday months: share of the normal technician hours
# Plan 2a position names (Umsatz Service = all service orders incl. retrofits and commissioning; costs negative)
POSITIONEN = ["Umsatz Ersatzteile", "Umsatz Service", "Umsatz Verträge", "Umsatz Schulung", "Material",
              "Fremdleistung", "Personalkosten", "Gewährleistung", "Gemeinkostenumlage"]
LIEFERANTEN = {"Hydraulikaggregat": "Hydraulik Nord GmbH", "Ventilblock": "Hydraulik Nord GmbH",
               "Pneumatikzylinder": "Hydraulik Nord GmbH", "Spindellager": "Kugeltec GmbH",
               "Zahnriemen": "Kugeltec GmbH", "Förderband": "Kugeltec GmbH", "Servomotor": "Elektro Brandt KG",
               "Sensor": "Elektro Brandt KG", "Steuerplatine": "Elektro Brandt KG", "Messerwelle": "Filtra AG",
               "Dichtungssatz": "Filtra AG", "Filterelement": "Filtra AG"}
HA200 = "ET-10074"  # Hydraulikaggregat HA-200: the part of the two supplier offers
SCHULUNG = {"10": 8700, "11": 5800, "12": 0, "01": 4350, "02": 11600, "03": 8700, "04": 5800, "05": 2900,
            "06": 7250, "07": 0, "08": 0, "09": 10150}
GEWAEHR = {"08": 13400, "09": 18900}  # MM-600 spindle failures (Müller GmbH)
H_AUF = ["Auftragsnr", "Kunde", "Anlage", "Auftragsart", "Eingang", "Abschluss", "Status", "Stunden", "Umsatz_EUR",
         "Kosten_EUR", "Team"]
H_BASE = ["Kunde", "Anlage", "Maschinentyp", "Baujahr", "Vertrag", "Vertragsende"]
H_TEILE = ["Datum", "Kunde", "Teilenr", "Menge", "Stueckpreis_EUR", "Lieferbar", "Einstandspreis_EUR", "Lieferant"]
H_QUAL = ["Team", "Maschinentyp", "Auftragsart", "Qualifiziert_Anzahl", "In_Schulung_Anzahl", "Ausbilder_Anzahl"]
H_ERG = ["Monat", "Position", "Plan_EUR", "Ist_EUR"]
H_KAP = ["Monat", "Team", "Techniker_Anzahl", "Soll_Stunden", "Ist_Stunden"]


def r2(x: float) -> float:
    return round(x + 0.0, 2)


def de_zahl(x: float) -> str:
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def save(path: Path, header: list[str], rows: list[list], blatt: str | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    if blatt:
        wb.active.title = blatt
    wb.active.append(header)
    for r in rows:
        wb.active.append(r)
    wb.save(path)


def teile_katalog(rng: random.Random) -> list[tuple[str, str, float]]:
    namen = ["Spindellager", "Dichtungssatz", "Hydraulikaggregat HA-200", "Servomotor", "Zahnriemen", "Sensor",
             "Ventilblock", "Steuerplatine", "Messerwelle", "Förderband", "Pneumatikzylinder", "Filterelement"]
    out = []
    for i in range(120):
        stufe = rng.random()
        preis = rng.uniform(25, 180) if stufe < 0.5 else rng.uniform(180, 900) if stufe < 0.85 else rng.uniform(900, 6500)
        out.append((f"ET-{10000 + i * 37}", f"{namen[i % len(namen)]} {TYPEN[i % 3]}", r2(preis)))
    out[2] = (HA200, "Hydraulikaggregat HA-200", 740.0)
    return out


def einstand(teil: tuple, monat: str) -> float:
    """Purchase price: 52 % of the 2026 list price; Hydraulik Nord raised its prices by 7.5 % from August 2026."""
    if teil[0] == HA200:
        return 412.5 if monat >= "2026-08" else 383.72
    faktor = 1.075 if LIEFERANTEN[teil[1].split()[0]] == "Hydraulik Nord GmbH" and monat >= "2026-08" else 1.0
    return r2(teil[2] * 0.52 * faktor * (1.0 if monat >= "2026-01" else 1 / 1.045))


def anlagen(rng: random.Random) -> list[list]:
    rows = []
    for kunde, _team, stufe, n, im_vertrag, ende in KUNDEN:
        for i in range(1, n + 1):
            typ, baujahr = rng.choice(TYPEN), rng.randrange(2008, 2026)
            vertrag = i <= im_vertrag
            if kunde == "Müller GmbH":  # Anlage 4: the MM-600 of the complaint mail, built 2021, no contract
                vertrag = i != 4 and i <= im_vertrag + 1
                typ, baujahr = ("MM-600", 2021) if i == 4 else (typ, baujahr)
            if kunde == "Nordmetall GmbH" and i == 1:
                typ = "MM-400"
            rows.append([kunde, f"Anlage {i}", typ, baujahr, "ja" if vertrag else "nein",
                         dt.date.fromisoformat(ende) if vertrag else None])
    return rows


def kapazitaet(monat: str, rng: random.Random) -> list[list]:
    rows = []
    for team, n in TEAMS.items():
        if team == "Süd" and monat >= "2026-07":
            n -= 1
        soll = round(n * 152 * VERFUEGBAR.get(monat[5:], 1.0))
        faktor = {"Nord": 0.97, "Süd": 1.12 if monat >= "2026-07" else 1.04, "West": 0.95}[team]
        rows.append([monat, team, n, soll, round(soll * (faktor + rng.uniform(-0.02, 0.02)))])
    return rows


def auftraege(monat: str, kap: list[list], basis: list[list], rng: random.Random, nr: list[int]) -> tuple[list, list]:
    """Orders until each team's order hours reach ~72 % of its worked hours; retrofits come with a material package."""
    jahr, mon = int(monat[:4]), int(monat[5:])
    satz = SATZ[jahr]
    tage = (dt.date(jahr + mon // 12, mon % 12 + 1, 1) - dt.timedelta(days=1)).day
    rows, pakete = [], []
    retrofits = {"Nord": 1 if mon % 2 else 0, "Süd": 1 if mon in (3, 6, 9, 11) else 0, "West": 1 if mon in (1, 5, 10) else 0}
    for _m, team, _n, _soll, ist in kap:
        ziel = round(ist * (0.70 if monat == SEPT else 0.72))
        maschinen = [b for b in basis if dict((k[0], k[1]) for k in KUNDEN)[b[0]] == team]
        stunden = 0
        while stunden < ziel:
            b = rng.choice(maschinen)
            eingang = dt.date(jahr, mon, rng.randrange(1, tage - 4))
            if retrofits[team]:
                retrofits[team] -= 1
                h = rng.choice([60, 80, 100, 120, 140])
                paket = r2(rng.uniform(12000, 30000))
                art, umsatz, kosten = "Retrofit", r2(h * satz + paket), r2(h * 68 + paket * 0.78)
                pakete.append(paket)
            else:
                art = rng.choices(["Wartung", "Reparatur", "Inbetriebnahme"], [35, 55, 10])[0]
                h = rng.choice({"Wartung": [4, 6, 8], "Reparatur": [3, 4, 6, 8, 12, 16], "Inbetriebnahme": [16, 24, 32]}[art])
                if art == "Inbetriebnahme":
                    umsatz = r2(h * IBN_SATZ)
                elif art == "Wartung" and b[4] == "ja":
                    umsatz = 0.0  # covered by the maintenance contract
                else:
                    umsatz = r2(h * satz + FAHRT[team])
                kosten = r2(h * 68 + FAHRT[team] * 0.8)
            stunden += h
            nr[0] += 1
            rows.append([f"SA-{monat[2:4]}{nr[0]:05d}", b[0], b[1], art, eingang,
                         eingang + dt.timedelta(days=rng.randrange(0, 4)), "abgeschlossen", h, umsatz, kosten, team])
    if monat == SEPT:  # the repair the complaint mail refers to (Müller GmbH, Anlage 4, 3 September)
        nr[0] += 1
        rows.append([f"SA-26{nr[0]:05d}", "Müller GmbH", "Anlage 4", "Reparatur", dt.date(2026, 9, 2),
                     dt.date(2026, 9, 3), "abgeschlossen", 8, r2(8 * satz + FAHRT["Süd"]), r2(8 * 68 + FAHRT["Süd"] * 0.8), "Süd"])
    rows.sort(key=lambda r: (r[4], r[0]))
    if monat == SEPT:  # backlog: September orders received from the 24th are still open (no Abschluss)
        for r in rows:
            if r[4] >= dt.date(2026, 9, 24) and r[3] != "Retrofit":
                r[5], r[6] = None, "offen"
    return rows, pakete


def ersatzteile(monat: str, katalog: list, rng: random.Random) -> list[list]:
    jahr, mon = int(monat[:4]), int(monat[5:])
    preisfaktor = 1.0 if jahr == 2026 else 1 / 1.045
    quote = 0.79 if monat in ("2026-08", "2026-09") else 0.88
    rows, gesehen = [], set()
    while len(rows) < round(96 * SAISON[monat[5:]]):
        teil = rng.choice(katalog)
        zeile = [dt.date(jahr, mon, rng.randrange(1, 29)), rng.choice(KUNDEN)[0], teil[0],
                 rng.choices([1, 2, 3, 4, 6], [45, 25, 15, 10, 5])[0], r2(teil[2] * preisfaktor),
                 "ja" if rng.random() < quote else "nein", einstand(teil, monat), LIEFERANTEN[teil[1].split()[0]]]
        if tuple(zeile[:4]) not in gesehen:  # identical rows would be refused by daten_pruefen as duplicates
            gesehen.add(tuple(zeile[:4]))
            rows.append(zeile)
    rows.sort(key=lambda r: (r[0], r[2]))
    return rows


def plan(monat: str) -> dict[str, float]:
    """Plan values: the 2026 budget for 2026 months, the (lower) 2025 plan before."""
    s, f = SAISON[monat[5:]], (1.0 if monat >= "2026-01" else 0.96)
    et, service, sch = 168000 * s * f, (138000 * s + 45000) * f, 6000 * f
    v = 66000 if monat >= "2026-01" else 62000
    personal = 24 * 6400 * (1.5 if monat[5:] == "11" else 1.0)
    werte = {"Umsatz Ersatzteile": et, "Umsatz Service": service, "Umsatz Verträge": v, "Umsatz Schulung": sch,
             "Material": -(0.52 * et + 0.55 * 45000 * f), "Fremdleistung": -14000, "Personalkosten": -personal,
             "Gewährleistung": -6000, "Gemeinkostenumlage": -36500}
    return {k: float(round(v_)) for k, v_ in werte.items()}


def ergebnis(monat: str, auf: list, teile: list, pakete: list, kap: list, rng: random.Random) -> list[list]:
    et = r2(sum(r[3] * r[4] for r in teile))
    vertraege = r2(sum(STUFEN[k[2]] * k[4] for k in KUNDEN if k[2]) / 12)
    material = -r2(sum(r[3] * r[6] for r in teile) + sum(p * 0.78 for p in pakete))  # parts at purchase price + retrofit packages
    fte = sum(r[2] for r in kap) + 9
    personal = -float(round(fte * 6250 * (1.5 if monat[5:] == "11" else 1.0)))
    ist = {"Umsatz Ersatzteile": et, "Umsatz Service": r2(sum(r[8] for r in auf)), "Umsatz Verträge": vertraege,
           "Umsatz Schulung": float(SCHULUNG[monat[5:]]), "Material": material,
           "Fremdleistung": -float(rng.randrange(11, 20) * 1000), "Personalkosten": personal,
           "Gewährleistung": -float(GEWAEHR.get(monat[5:], rng.randrange(5, 9) * 1000)),
           "Gemeinkostenumlage": -36500.0}
    p = plan(monat)
    return [[monat, pos, p[pos], ist[pos]] for pos in POSITIONEN]


def importiere(ws: Path, datei: Path, vorlage: str, heute: str) -> None:
    r = daten_pruefen.pruefe(ws, datei, vorlage, None, [], True, heute)
    if not r["ok"]:
        raise SystemExit(f"Import von {datei.name} fehlgeschlagen: {r['meldungen']}")


def kennzahlen_monat(m: str, d: dict) -> dict:
    erg = {r[1]: r[3] for r in d["ergebnis"][m]}
    planw = {r[1]: r[2] for r in d["ergebnis"][m]}
    umsatz = r2(sum(v for k, v in erg.items() if k.startswith("Umsatz ")))
    db1 = r2(umsatz + erg["Material"] + erg["Fremdleistung"])
    db2 = r2(db1 + erg["Personalkosten"])
    plan_umsatz = r2(sum(v for k, v in planw.items() if k.startswith("Umsatz ")))
    plan_db2 = plan_umsatz + planw["Material"] + planw["Fremdleistung"] + planw["Personalkosten"]
    teile, auf, kap = d["ersatzteile"][m], d["auftraege"][m], d["kapazitaet"][m]
    out = {"umsatz_gesamt": umsatz, "plan_umsatz_gesamt": plan_umsatz, "db1": db1, "db2": db2,
           "ergebnis": r2(db2 + erg["Gewährleistung"] + erg["Gemeinkostenumlage"]),
           "plan_ergebnis": r2(plan_db2 + planw["Gewährleistung"] + planw["Gemeinkostenumlage"]),
           "umsatz_ersatzteile": erg["Umsatz Ersatzteile"], "umsatz_service": erg["Umsatz Service"],
           "umsatz_vertraege": erg["Umsatz Verträge"], "umsatz_schulung": erg["Umsatz Schulung"],
           "auftraege_retrofit_umsatz": r2(sum(r[8] for r in auf if r[3] == "Retrofit")),
           "material": erg["Material"], "fremdleistung": erg["Fremdleistung"],
           "personal": erg["Personalkosten"], "gewaehrleistung": erg["Gewährleistung"],
           "gemeinkosten": erg["Gemeinkostenumlage"],
           "auftraege_umsatz": r2(sum(r[8] for r in auf)), "auftraege_zeilen": len(auf),
           "auftraege_stunden": sum(r[7] for r in auf), "auftraege_offen": sum(r[6] == "offen" for r in auf),
           "ersatzteile_umsatz": r2(sum(r[3] * r[4] for r in teile)),
           "ersatzteile_zeilen": len(teile),
           "ersatzteile_lieferquote": round(100 * sum(r[5] == "ja" for r in teile) / len(teile), 1),
           "ersatzteile_einstand": r2(sum(r[3] * r[6] for r in teile)),
           "kapazitaet_soll": sum(r[3] for r in kap), "kapazitaet_ist": sum(r[4] for r in kap)}
    return out


def daten() -> dict:
    rng = random.Random(2026)
    katalog = teile_katalog(rng)
    basis = anlagen(rng)
    d = {"katalog": katalog, "basis": basis, "auftraege": {}, "ersatzteile": {}, "kapazitaet": {}, "ergebnis": {}}
    nr = [0]
    for m in MONATE:
        kap = kapazitaet(m, rng)
        auf, pakete = auftraege(m, kap, basis, rng, nr)
        teile = ersatzteile(m, katalog, rng)
        d["kapazitaet"][m], d["auftraege"][m], d["ersatzteile"][m] = kap, auf, teile
        d["ergebnis"][m] = ergebnis(m, auf, teile, pakete, kap, rng)
    return d


def schreibe_jahr(beispiel: Path, d: dict) -> None:
    """Every month goes through daten_pruefen, so 07_Daten/ looks exactly like a user's validated imports."""
    with tempfile.TemporaryDirectory(prefix="slk-beispiel-") as tmp:
        ws = Path(tmp)
        for m in MONATE:
            for vorlage, h, rows in (("auftraege", H_AUF, d["auftraege"][m]), ("ersatzteile", H_TEILE, d["ersatzteile"][m]),
                                     ("ergebnis", H_ERG, d["ergebnis"][m]), ("kapazitaet", H_KAP, d["kapazitaet"][m])):
                save(ws / "00_Eingang" / f"{vorlage}_{m}.xlsx", h, rows)
                importiere(ws, ws / "00_Eingang" / f"{vorlage}_{m}.xlsx", vorlage, f"{m}-28")
        save(ws / "00_Eingang" / "installed_base_2026-09.xlsx", H_BASE, d["basis"])
        importiere(ws, ws / "00_Eingang" / "installed_base_2026-09.xlsx", "installed_base", "2026-09-30")
        shutil.copytree(ws / "07_Daten", beispiel / "07_Daten")


def budget(path: Path) -> None:
    wb = Workbook()
    sh = wb.active
    sh.title = "Budget 2026"
    monate = [f"2026-{i:02d}" for i in range(1, 13)]
    sh.append(["Position", *monate, "Summe_2026"])
    for pos in POSITIONEN:
        werte = [plan(m)[pos] for m in monate]
        sh.append([pos, *werte, float(sum(werte))])
    a = wb.create_sheet("Annahmen")
    for z in (["Annahme", "Wert"], ["Stundensatz Techniker 2026 (EUR)", SATZ[2026]], ["Techniker gesamt", 15],
              ["Mitarbeitende Service gesamt (inkl. Innendienst, Ersatzteile, Leitung)", 24],
              ["Personalkosten je Kopf und Monat (EUR)", 6400], ["Materialquote Ersatzteile", 0.52],
              ["Vertragsumsatz je Monat (EUR)", 66000], ["Gemeinkostenumlage je Monat (EUR)", 36500]):
        a.append(z)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def preisliste(path: Path, katalog: list) -> None:
    """One sheet 'Preisliste' (Plan 4f G1 schema); Bezeichnung + Preis_EUR also serve Plan 4d's word search."""
    wb = Workbook()
    sh = wb.active
    sh.title = "Preisliste"
    sh.append(["Artikelnr", "Bezeichnung", "Kategorie", "Einheit", "Preis_EUR", "Gueltig_ab", "Bemerkung"])
    ab = dt.date(2026, 1, 1)
    zeilen = [["SL-100", "Techniker-Stundensatz", "Stundensatz", "Stunde", SATZ[2026], "Zuschläge: Überstunden +25 %, Wochenende +50 %"],
              ["SL-102", "Ferndiagnose", "Stundensatz", "Stunde", 98, ""],
              *[[f"SL-20{i}", f"Anfahrtspauschale Zone {z}", "Pauschale", "Einsatz", FAHRT[z], ""] for i, z in enumerate(FAHRT)],
              *[[f"SL-3{i}{j}", f"Wartungsvertrag {stufe} {typ}", "Vertrag", "Anlage und Jahr", preis, "Preis je Stufe, gleich für alle Typen"]
                for i, (stufe, preis) in enumerate(STUFEN.items()) for j, typ in enumerate(TYPEN)],
              ["SL-400", "Schulung vor Ort", "Schulung", "Tag", 1450, ""],
              *[[f"SL-50{j}", f"Retrofit {typ}", "Sonstiges", "Anlage", p, "Richtpreis Steuerungs-Retrofit, Festpreis je Angebot"]
                for j, (typ, p) in enumerate(zip(TYPEN, (48000, 62000, 85000)))]]
    for z in zeilen:
        sh.append(z[:5] + [ab, z[5]])
    for t in katalog:
        sh.append([t[0], t[1], "Ersatzteil", "Stück", t[2], ab, ""])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def qualifikation(path: Path) -> None:
    """Team qualification per machine type (Plan 4h G1 schema), written directly: its import template ships with
    Plan 4h, so this file is not routed through daten_pruefen."""
    zahlen = {"Nord": (6, 5, 3), "Süd": (4, 4, 1), "West": (4, 2, 2)}  # MM-400, MM-600, MM-800 qualified heads
    import csv
    import io
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(H_QUAL + ["_quelle_datei", "_quelle_blatt", "_quelle_zeile"])
    zeile = 2
    for team, quali in zahlen.items():
        for typ, n in zip(TYPEN, quali):
            w.writerow([team, typ, "alle", n, 1 if n < 3 else 0, 1 if n >= 4 else 0, "qualifikation_2026-09.xlsx",
                        "Sheet", zeile])
            zeile += 1
    path.write_text("\ufeff" + buf.getvalue(), encoding="utf-8")


def eingang_september(ziel: Path, d: dict) -> None:
    save(ziel / "auftraege_2026-09.xlsx", H_AUF, d["auftraege"][SEPT])
    save(ziel / "installed_base_2026-09.xlsx", H_BASE, d["basis"])
    save(ziel / "ersatzteile_2026-09.xlsx", H_TEILE, d["ersatzteile"][SEPT])
    save(ziel / "ergebnis_2026-09.xlsx", H_ERG, d["ergebnis"][SEPT])
    save(ziel / "kapazitaet_2026-09.xlsx", H_KAP, d["kapazitaet"][SEPT])


def unordentlich(ziel: Path, d: dict) -> dict:
    """Plan 2c defects (German formats, renamed columns, 3 duplicate rows, duplicate file, missing column) plus
    Plan 3: the scanned supplier offer and the controlling report that contradicts 07_Daten for September."""
    auf = d["auftraege"][SEPT]
    h = [{"Umsatz_EUR": "Auftragswert", "Stunden": "Std."}.get(c, c) for c in H_AUF]
    rows = [r[:4] + [r[4].strftime("%d.%m.%Y"), r[5].strftime("%d.%m.%Y") if r[5] else None, r[6], str(r[7]), de_zahl(r[8]),
                     de_zahl(r[9]), r[10]] for r in auf]
    rows += [rows[5], rows[17], rows[33]]
    save(ziel / "auftraege_2026-09.xlsx", h, rows)
    shutil.copy(ziel / "auftraege_2026-09.xlsx", ziel / "auftraege_2026-09 (1).xlsx")
    save(ziel / "installed_base_2026-09.xlsx", H_BASE, d["basis"])
    save(ziel / "ersatzteile_2026-09.xlsx", H_TEILE, [[r[0].strftime("%d.%m.%Y"), r[1], r[2], r[3], de_zahl(r[4]), r[5],
                                                       de_zahl(r[6]), r[7]] for r in d["ersatzteile"][SEPT]])
    erg = [r[:] for r in d["ergebnis"][SEPT]]
    erg[1][3] = r2(erg[1][3] + 12000)  # "Umsatz Service" does not match the orders export
    save(ziel / "ergebnis_2026-09.xlsx", H_ERG, [[r[0], r[1], de_zahl(r[2]), de_zahl(r[3])] for r in erg])
    save(ziel / "kapazitaet_2026-09.xlsx", H_KAP[:4], [r[:4] for r in d["kapazitaet"][SEPT]])
    ctrl = [r[:] for r in d["ergebnis"][SEPT]]
    ctrl[0][3] = r2(ctrl[0][3] - 14600)  # credit note Nordmetall booked by controlling, not in the export
    save(ziel / "Controlling_Monatsbericht_2026-09.xlsx", ["Monat", "Position", "Plan", "Ist"],
         [[r[0], r[1], r[2], r[3]] for r in ctrl], blatt="Ergebnis Service")
    (ziel / "2026-09-30_angebot-hydraulik-nord-scan.pdf").write_bytes(scan_pdf())
    return {"konflikt_monat": SEPT, "konflikt_position": "Umsatz Ersatzteile",
            "konflikt_wert_daten": d["ergebnis"][SEPT][0][3], "konflikt_wert_controlling": ctrl[0][3],
            "konflikt_differenz": 14600.0, "unordentlich_ergebnis_abweichung": 12000.0}


def generate(target_root: Path) -> dict:
    d = daten()
    clean, messy = target_root / "beispiel", target_root / "beispiel-unordentlich"
    for p in (clean, messy):
        shutil.rmtree(p, ignore_errors=True)
    schreibe_jahr(clean, d)
    budget(clean / "07_Daten" / "budget_2026.xlsx")
    qualifikation(clean / "07_Daten" / "qualifikation_2026-09.csv")
    eingang_september(clean / "00_Eingang", d)
    for name, text in (("2026-09-28_mail-reklamation.eml", MAIL_REKLAMATION),
                       ("2026-09-29_mail-eskalation.eml", MAIL_ESKALATION),
                       ("2026-10-01_mail-anfrage-rahmenvertrag.eml", MAIL_ANFRAGE),
                       ("2026-09-24_angebot-hydraulik-nord.eml", MAIL_HYDRAULIK),
                       ("2026-09-25_angebot-fluidtec.eml", MAIL_FLUIDTEC)):
        (clean / "00_Eingang" / name).write_text(text, encoding="utf-8")
    shutil.copytree(QUELLE / "Unternehmen", clean / "Unternehmen")
    master_pptx(clean / "Unternehmen" / "vorlagen" / "master.pptx")
    briefkopf(clean / "Unternehmen" / "vorlagen" / "briefkopf.docx")
    for k in KUNDEN:
        if k[2]:
            p = clean / "06_Kunden" / k[0] / "vertrag.md"
            p.parent.mkdir(parents=True, exist_ok=True)
            p.write_text(vertrag_md(k, STUFEN[k[2]]), encoding="utf-8")
    preisliste(clean / "04_Angebote" / "preisliste_2026.xlsx", d["katalog"])
    for name, text in projekte().items():
        p = clean / "05_Projekte" / name / "projekt.md"
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(text, encoding="utf-8")

    (messy / "00_Eingang").mkdir(parents=True)
    extra = unordentlich(messy / "00_Eingang", d)
    for name, text in (("2026-09-28_mail-reklamation.eml", MAIL_REKLAMATION),
                       ("2026-09-29_mail-eskalation.eml", MAIL_ESKALATION),
                       ("2026-09-29_mail-preisanfrage.eml", MAIL_INJECTION),
                       ("2026-10-01_mail-anfrage-rahmenvertrag.eml", MAIL_ANFRAGE),
                       ("2026-09-25_angebot-fluidtec.eml", MAIL_FLUIDTEC)):
        (messy / "00_Eingang" / name).write_text(text, encoding="utf-8")
    dateien = target_root / "evals" / "_gemeinsam" / "dateien"
    dateien.mkdir(parents=True, exist_ok=True)
    master_ohne_layouts(dateien / "Firmenvorlage.pptx")

    m = {mon: kennzahlen_monat(mon, d) for mon in MONATE}
    sep = m[SEPT]
    erwartet: dict = {
        # Plan 2c keys (September 2026; daten-pruefen evals and tests)
        "auftraege_umsatz": sep["auftraege_umsatz"], "auftraege_zeilen": sep["auftraege_zeilen"],
        "ersatzteile_umsatz": sep["ersatzteile_umsatz"], "ergebnis_ist_umsatz": sep["umsatz_gesamt"],
        "kapazitaet_ist_stunden": sep["kapazitaet_ist"], "installed_base_anlagen": len(d["basis"]),
        "unordentlich_duplikate_auftraege": 3,
        # Plan 3: one key per number a grader checks; <name>_<JJJJ-MM> per month, <name>_jahr for 2025-10 … 2026-09
        "installed_base_vertraege": sum(b[4] == "ja" for b in d["basis"]),
        "vertraege_jahreswert": float(sum(STUFEN[k[2]] * k[4] for k in KUNDEN if k[2])),
        "budget_umsatz_gesamt_2026": float(sum(plan(f"2026-{i:02d}")[p] for i in range(1, 13)
                                               for p in POSITIONEN if p.startswith("Umsatz "))),
        "angebot_hydraulik_nord_eur": 16500.0, "angebot_fluidtec_eur": 15560.0,
        "projekt_verzug_tage": 35, **extra}
    for mon, werte in m.items():
        erwartet |= {f"{k}_{mon}": v for k, v in werte.items()}
    for k in m[SEPT]:
        if k.endswith("lieferquote"):
            continue
        erwartet[f"{k}_jahr"] = r2(sum(m[mon][k] for mon in MONATE))
    out = target_root / "evals" / "erwartet" / "beispiel.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(erwartet, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return erwartet


if __name__ == "__main__":
    root = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "plugin"
    e = generate(root)
    print(json.dumps({k: e[k] for k in ("umsatz_gesamt_jahr", "ergebnis_jahr", "auftraege_umsatz", "auftraege_zeilen")},
                     ensure_ascii=False))
